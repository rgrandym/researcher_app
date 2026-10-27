"""Editable Office masters shared by the app and document.py.

The Markdown contains a fenced office JSON model: typed Excel cells or PowerPoint
text frames and tables. Native objects and design live in an immutable source copy.
Only changed values are applied to that copy; exports never overwrite the source.
"""

import hashlib
import json
import math
import re
import shutil
import tempfile
import warnings
from zipfile import BadZipFile, is_zipfile
from copy import copy
from datetime import date, datetime, time
from difflib import SequenceMatcher
from pathlib import Path

MARKER = "<!-- office-master:1 -->"
FORMATS = {"xlsx", "pptx"}
MAX_CELLS = 100_000


def digest(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def encode(data):
    return f"{MARKER}\n\n# {data['format'].upper()} document\n\n```office\n" + json.dumps(
        data, ensure_ascii=False, indent=2, allow_nan=False
    ) + "\n```\n"


def decode(text):
    match = re.fullmatch(
        r"<!-- office-master:1 -->\s*# [^\n]+\s+```office\n(.*)\n```\s*", text, re.S
    )
    if not match:
        raise ValueError("Keep the Office master header and its single office block.")
    try:
        data = json.loads(match[1])
    except ValueError as exc:
        raise ValueError(f"The Office content is not valid JSON: {exc}") from exc
    if not isinstance(data, dict) or data.get("version") != 1 or data.get("format") not in FORMATS:
        raise ValueError("Unsupported Office master version or format.")
    if data.get("design") not in ("preserve", "new"):
        raise ValueError("Choose preserve or new for the design.")
    return data


def cell_value(cell):
    value = cell.value
    if value is None:
        return None
    if isinstance(value, (datetime, date, time)):
        return {"kind": type(value).__name__, "value": value.isoformat()}
    if not isinstance(value, (str, int, float, bool)):
        # Array/data-table formulas are retained natively and cannot be flattened.
        return {"kind": "readonly", "value": str(getattr(value, "text", "Special formula"))}
    kind = {"f": "formula", "e": "error", "b": "boolean", "n": "number"}.get(cell.data_type, "text")
    return {"kind": kind, "value": value}


def walk_shapes(shapes, prefix=""):
    for shape in shapes:
        key = f"{prefix}{shape.shape_id}"
        yield key, shape
        if shape.shape_type == 6:  # group
            yield from walk_shapes(shape.shapes, key + "/")


# ---------------------------------------------------------------- slide design

A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
THEME_COLORS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4",
                "accent5", "accent6", "hlink", "folHlink")
STYLE_KEYS = {"font", "size", "color", "bold", "italic", "fill", "align"}
ALIGN = {"left": 1, "center": 2, "right": 3, "justify": 4}
EMU_PER_PT = 12700


def _hexcolor(value):
    return isinstance(value, str) and re.fullmatch(r"[0-9A-F]{6}", value) is not None


def _box(shape):
    if None in (shape.left, shape.top, shape.width, shape.height):
        return None
    return {k: round(v / EMU_PER_PT, 2) for k, v in
            (("x", shape.left), ("y", shape.top), ("w", shape.width), ("h", shape.height))}


def _rgb(color):
    try:
        return str(color.rgb) if color.type == 1 else None  # MSO_COLOR_TYPE.RGB
    except AttributeError:
        return None


def _style(shape):
    """The look of a text shape's first run, as far as it is set on the shape itself."""
    style = {}
    first = next((r for p in shape.text_frame.paragraphs for r in p.runs), None)
    if first is not None:
        font = first.font
        if font.name:
            style["font"] = font.name
        if font.size:
            style["size"] = round(font.size.pt, 2)
        if font.bold is not None:
            style["bold"] = font.bold
        if font.italic is not None:
            style["italic"] = font.italic
        if (color := _rgb(font.color)):
            style["color"] = color
    alignment = shape.text_frame.paragraphs[0].alignment
    names = {v: k for k, v in ALIGN.items()}
    if alignment is not None and int(alignment) in names:
        style["align"] = names[int(alignment)]
    try:
        if shape.fill.type == 1 and (fill := _rgb(shape.fill.fore_color)):  # MSO_FILL.SOLID
            style["fill"] = fill
    except (AttributeError, TypeError):
        pass
    return style


def _theme_parts(deck):
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    parts = []
    for master in deck.slide_masters:
        part = master.part.part_related_by(RT.THEME)
        if part not in parts:
            parts.append(part)
    return parts


def _read_theme(deck):
    from lxml import etree

    root = etree.fromstring(_theme_parts(deck)[0].blob)
    colors = {}
    scheme = root.find(f".//{A}clrScheme")
    for name in THEME_COLORS:
        el = scheme.find(A + name) if scheme is not None else None
        if el is not None and len(el):
            value = el[0].get("val") if el[0].tag == A + "srgbClr" else el[0].get("lastClr")
            if value and _hexcolor(value.upper()):
                colors[name] = value.upper()
    fonts = {}
    for key, tag in (("major", "majorFont"), ("minor", "minorFont")):
        latin = root.find(f".//{A}{tag}/{A}latin")
        if latin is not None and latin.get("typeface"):
            fonts[key] = latin.get("typeface")
    return {"colors": colors, "fonts": fonts}


def _write_theme(deck, theme, prior):
    """Set the theme colours and fonts that changed, in every slide master's theme."""
    from lxml import etree

    colors = {k: v for k, v in theme.get("colors", {}).items() if prior.get("colors", {}).get(k) != v}
    fonts = {k: v for k, v in theme.get("fonts", {}).items() if prior.get("fonts", {}).get(k) != v}
    if not colors and not fonts:
        return
    for part in _theme_parts(deck):
        root = etree.fromstring(part.blob)
        scheme = root.find(f".//{A}clrScheme")
        for name, value in colors.items():
            el = scheme.find(A + name)
            for child in list(el):
                el.remove(child)
            etree.SubElement(el, A + "srgbClr").set("val", value)
        for key, value in fonts.items():
            latin = root.find(f".//{A}{'majorFont' if key == 'major' else 'minorFont'}/{A}latin")
            if latin is not None:
                latin.set("typeface", value)
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def _runs(shape):
    frames = [shape.text_frame] if shape.has_text_frame else []
    if getattr(shape, "has_table", False) and shape.has_table:
        frames += [cell.text_frame for row in shape.table.rows for cell in row.cells]
    for frame in frames:
        for paragraph in frame.paragraphs:
            yield from paragraph.runs


def _apply_style(shape, style, prior):
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    changed = {k: v for k, v in style.items() if prior.get(k) != v}
    for run in _runs(shape) if changed.keys() - {"align", "fill"} else ():
        if "font" in changed:
            run.font.name = changed["font"]
        if "size" in changed:
            run.font.size = Pt(changed["size"])
        if "bold" in changed:
            run.font.bold = changed["bold"]
        if "italic" in changed:
            run.font.italic = changed["italic"]
        if "color" in changed:
            run.font.color.rgb = RGBColor.from_string(changed["color"])
    if "align" in changed:
        for paragraph in shape.text_frame.paragraphs:
            paragraph.alignment = ALIGN[changed["align"]]
    if "fill" in changed:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(changed["fill"])


def _swap(deck, theme):
    """Replace a colour or a font wherever the deck sets it on text or a shape's fill."""
    from pptx.dml.color import RGBColor

    colors = theme.get("swap_colors") or {}
    fonts = theme.get("swap_fonts") or {}
    if not colors and not fonts:
        return
    for slide in deck.slides:
        for _, shape in walk_shapes(slide.shapes):
            for run in _runs(shape):
                if run.font.name in fonts:
                    run.font.name = fonts[run.font.name]
                if (color := _rgb(run.font.color)) in colors:
                    run.font.color.rgb = RGBColor.from_string(colors[color])
            try:
                if shape.fill.type == 1 and (fill := _rgb(shape.fill.fore_color)) in colors:
                    shape.fill.fore_color.rgb = RGBColor.from_string(colors[fill])
            except (AttributeError, TypeError):
                pass


def palette(path):
    """The colours and fonts a deck sets on its text and shapes, most used first."""
    from collections import Counter
    from pptx import Presentation

    deck = Presentation(path)
    colors, fonts = Counter(), Counter()
    for slide in deck.slides:
        for _, shape in walk_shapes(slide.shapes):
            for run in _runs(shape):
                if run.font.name:
                    fonts[run.font.name] += 1
                if (color := _rgb(run.font.color)):
                    colors[color] += 1
            try:
                if shape.fill.type == 1 and (fill := _rgb(shape.fill.fore_color)):
                    colors[fill] += 1
            except (AttributeError, TypeError):
                pass
    return {"colors": [c for c, _ in colors.most_common(24)], "fonts": [f for f, _ in fonts.most_common(12)]}


def _check_box(box):
    if not isinstance(box, dict) or set(box) != {"x", "y", "w", "h"} or any(
            type(v) not in (int, float) or not math.isfinite(v) or abs(v) > 20000 for v in box.values()):
        raise ValueError("A box needs x, y, w and h in points.")
    if box["w"] < 0 or box["h"] < 0:  # lines and connectors have no height or width
        raise ValueError("A box cannot have a negative width or height.")


def _check_style(style, prior):
    if not isinstance(style, dict) or not set(style) <= STYLE_KEYS:
        raise ValueError("A text style takes font, size, color, bold, italic, fill and align.")
    if not set(prior) <= set(style):
        raise ValueError("A style setting can be changed but not removed.")
    for key, value in style.items():
        ok = {"font": lambda v: isinstance(v, str) and 0 < len(v) <= 64,
              "size": lambda v: type(v) in (int, float) and 1 <= v <= 400,
              "color": _hexcolor, "fill": _hexcolor,
              "bold": lambda v: type(v) is bool, "italic": lambda v: type(v) is bool,
              "align": lambda v: v in ALIGN}[key](value)
        if not ok:
            raise ValueError(f"Invalid text style {key}: {value!r}.")


def _check_theme(theme, prior):
    if not isinstance(theme, dict) or not set(theme) <= {"colors", "fonts", "swap_colors", "swap_fonts"}:
        raise ValueError("A theme takes colors, fonts, swap_colors and swap_fonts.")
    colors, fonts = theme.get("colors", {}), theme.get("fonts", {})
    if not isinstance(colors, dict) or not set(colors) <= set(THEME_COLORS) or not all(map(_hexcolor, colors.values())):
        raise ValueError("Theme colours are six-digit hex values for dk1 … folHlink.")
    if not isinstance(fonts, dict) or not set(fonts) <= {"major", "minor"} or not all(
            isinstance(v, str) and 0 < len(v) <= 64 for v in fonts.values()):
        raise ValueError("Theme fonts are major (headings) and minor (body) typefaces.")
    if not set(prior.get("colors", {})) <= set(colors) or not set(prior.get("fonts", {})) <= set(fonts):
        raise ValueError("Theme colours and fonts can be changed but not removed.")
    for key, check in (("swap_colors", _hexcolor), ("swap_fonts", lambda v: isinstance(v, str) and 0 < len(v) <= 64)):
        swap = theme.get(key, {})
        if not isinstance(swap, dict) or not all(check(k) and check(v) for k, v in swap.items()):
            raise ValueError(f"{key} maps an old value to a new one.")


def _like(actual, expected):
    """Make the read-back deck comparable with the master: leave out what an older master
    does not record, carry swaps into the styles they reach (swaps are not read back), and
    skip styles under the new design, which rewrites every font."""
    theme = expected.get("theme", {})
    colors, fonts = theme.pop("swap_colors", None) or {}, theme.pop("swap_fonts", None) or {}
    if "theme" not in expected:
        actual.pop("theme", None)
    kinds = {s["kind"] for item in expected.get("slides", []) for s in item["shapes"]}
    for got, want in zip(actual.get("slides", []), expected.get("slides", [])):
        if "object" not in kinds:
            got["shapes"] = [s for s in got["shapes"] if s["kind"] != "object"]
        for a, e in zip(got["shapes"], want["shapes"]):
            style = e.get("style")
            if style:
                for key in ("color", "fill"):
                    if style.get(key) in colors:
                        style[key] = colors[style[key]]
                if style.get("font") in fonts:
                    style["font"] = fonts[style["font"]]
            for key in ("box", "style"):
                if key not in e or (key == "style" and expected.get("design") == "new"):
                    a.pop(key, None)
                    e.pop(key, None)


def read_native(path):
    if not is_zipfile(path):
        raise ValueError("The supplied file is not a readable Office archive.")
    try:
        return _read_native(path)
    except (BadZipFile, KeyError, TypeError) as exc:
        raise ValueError(f"The supplied Office file could not be read: {exc}") from exc


def _workbook(path):
    from openpyxl import load_workbook

    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter("always")
        book = load_workbook(path, rich_text=True)
    unsupported = [str(w.message) for w in caught if "not supported" in str(w.message) or "will be removed" in str(w.message)]
    if unsupported:
        book.close()
        raise ValueError("This workbook uses features that cannot be retained by this editor: " + "; ".join(unsupported))
    return book


def _read_native(path):
    fmt = Path(path).suffix[1:]
    data = {"version": 1, "format": fmt, "design": "preserve"}
    if fmt == "xlsx":
        book = _workbook(path)
        sheets = []
        count = 0
        for sheet in book:
            cells = {}
            # Sparse iteration avoids materialising a million empty formatted rows.
            for cell in sheet._cells.values():
                value = cell_value(cell)
                if value is not None:
                    cells[cell.coordinate] = value
            count += len(cells)
            if count > MAX_CELLS:
                raise ValueError(f"The browser editor supports up to {MAX_CELLS:,} populated cells.")
            sheets.append({"id": sheet.title, "name": sheet.title, "cells": cells})
        data["sheets"] = sheets
        book.close()
    elif fmt == "pptx":
        from pptx import Presentation

        deck = Presentation(path)
        slides = []
        for slide in deck.slides:
            shapes = []
            for key, shape in walk_shapes(slide.shapes):
                item = {"id": key}
                box = _box(shape) if "/" not in key else None
                if shape.has_text_frame:
                    item.update(kind="text", paragraphs=[p.text for p in shape.text_frame.paragraphs])
                    style = _style(shape)
                    if style:
                        item["style"] = style
                elif shape.has_table:
                    item.update(kind="table", rows=[[c.text for c in row.cells] for row in shape.table.rows])
                elif box:
                    # Pictures, charts, videos and groups: moved and resized, never rewritten.
                    item.update(kind="object", name=shape.name)
                else:
                    continue
                if box:
                    item["box"] = box
                shapes.append(item)
            slides.append({"id": str(slide.slide_id), "shapes": shapes})
        data["theme"] = _read_theme(deck)
        data["slides"] = slides
    else:
        raise ValueError("Choose an .xlsx workbook or .pptx presentation.")
    return data


def paths(run, slug):
    run = Path(run).resolve()
    if not re.fullmatch(r"[\w-]+", slug):
        raise ValueError("Invalid document name.")
    folder = run / "subjects" / slug
    if not folder.resolve().is_relative_to(run):
        raise ValueError("Document folder is outside the project.")
    master = run / f"{slug}.md"
    if not master.resolve().is_relative_to(run):
        raise ValueError("Document master is outside the project.")
    return folder, master


def baseline(run, slug):
    folder, _ = paths(run, slug)
    meta = json.loads((folder / "office-source.json").read_text())
    if meta.get("format") not in FORMATS:
        raise ValueError("Invalid Office source format.")
    source = folder / f"office-original.{meta['format']}"
    if not source.resolve().is_relative_to(Path(run).resolve()) or digest(source) != meta["sha256"]:
        raise ValueError("The saved Office source changed. Restore it before editing or exporting.")
    original = folder / "document.md"
    if not original.resolve().is_relative_to(Path(run).resolve()) or digest(original) != meta["model_sha256"]:
        raise ValueError("The saved original content changed. Restore it before editing or exporting.")
    return source, decode(original.read_text()), meta


def start(run, slug, source):
    import document

    folder, master = paths(run, slug)
    if master.exists():
        baseline(run, slug)
        validate(decode(master.read_text()), baseline(run, slug)[1])
        return master  # Retry opens the same master, never a second copy.
    source = Path(source).resolve()
    if not source.is_relative_to(Path(run).resolve()):
        raise ValueError("The supplied document must be inside this project.")
    data = read_native(source)
    folder.mkdir(parents=True, exist_ok=True)
    original = folder / f"office-original.{data['format']}"
    if original.exists() or (folder / "office-source.json").exists():
        raise ValueError("An Office source already exists without its master; restore the master first.")
    shutil.copy2(source, original)
    # Preserve the older extraction/change-plan files, if present.
    previous = folder / "document.md"
    if previous.exists():
        shutil.copy2(previous, folder / "document-before-office.md")
    previous.write_text(encode(data))
    (folder / "office-source.json").write_text(json.dumps({
        "format": data["format"], "sha256": digest(original), "model_sha256": digest(previous),
        "source": str(source.relative_to(Path(run).resolve())), "filename": source.name,
    }, indent=2) + "\n")
    master.write_text(encode(data))
    document.register(run, slug, master)
    return master


def _text(value, soft_breaks=False):
    pattern = r"[\x00-\x08\x0c\x0e-\x1f]" if soft_breaks else r"[\x00-\x08\x0b\x0c\x0e-\x1f]"
    if not isinstance(value, str) or len(value) > 32767 or re.search(pattern, value):
        raise ValueError("Text must be at most 32,767 characters and contain no control characters.")


def _cell(value):
    if not isinstance(value, dict) or set(value) != {"kind", "value"}:
        raise ValueError("Each cell needs a kind and value.")
    kind, v = value["kind"], value["value"]
    if kind == "number":
        if type(v) not in (int, float) or not (-1.7976931348623157e308 <= v <= 1.7976931348623157e308) or not math.isfinite(v):
            raise ValueError("A number cell needs a finite number.")
    elif kind == "boolean":
        if type(v) is not bool:
            raise ValueError("A boolean cell needs true or false.")
    elif kind in ("text", "formula", "error", "readonly"):
        _text(v)
        if kind == "formula" and not v.startswith("="):
            raise ValueError("A formula must start with =.")
        if kind == "error" and v not in {"#NULL!", "#DIV/0!", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "#N/A"}:
            raise ValueError("Invalid Excel error value.")
    elif kind in ("datetime", "date", "time"):
        _text(v)
        parsed = {"datetime": datetime, "date": date, "time": time}[kind].fromisoformat(v)
        if getattr(parsed, "tzinfo", None):
            raise ValueError("Excel dates and times must have no timezone.")
    else:
        raise ValueError("Unknown cell type.")


def validate(data, original):
    """Reject ambiguous or unsupported edits before a save or native export."""
    if data.get("format") != original["format"] or data.get("version") != 1:
        raise ValueError("The document format and version cannot change.")
    if data.get("design") not in ("preserve", "new"):
        raise ValueError("Choose preserve or new for the design.")
    group = "sheets" if data["format"] == "xlsx" else "slides"
    optional = {"theme"} if group == "slides" else set()
    if not {"version", "format", "design", group} <= set(data) <= {"version", "format", "design", group} | optional:
        raise ValueError("Unexpected Office document fields.")
    if "theme" in original and "theme" not in data:
        raise ValueError("Keep the presentation's theme; change its values instead.")
    if "theme" in data:
        _check_theme(data["theme"], original.get("theme", {}))
    items = data.get(group)
    if not isinstance(items, list) or not items or len(items) > 1000:
        raise ValueError("Keep between 1 and 1,000 sheets or slides.")
    known = {i["id"]: i for i in original[group]}
    seen = set()
    total = 0
    names = set()
    for item in items:
        if not isinstance(item, dict) or not isinstance(item.get("id"), str):
            raise ValueError("Each sheet or slide needs an ID.")
        key = item["id"]
        if key in seen or (key not in known and not re.fullmatch(r"new-[\w-]+", key)):
            raise ValueError("Use unique IDs; new items start with new-.")
        seen.add(key)
        old = known.get(key)
        if group == "sheets":
            from openpyxl.utils.cell import coordinate_to_tuple

            if set(item) != {"id", "name", "cells"}:
                raise ValueError("A sheet needs id, name, and cells.")
            name = item["name"]
            if not isinstance(name, str) or not name or len(name) > 31 or re.search(r"[\\/*?:\[\]]", name):
                raise ValueError("Use an Excel sheet name of 1–31 characters without \\ / * ? : [ ].")
            if name.lower() in names or (old and name != old["name"]):
                raise ValueError("Sheet names must be unique. Existing names stay fixed to preserve formula references.")
            names.add(name.lower())
            if not isinstance(item["cells"], dict):
                raise ValueError("Cells must be keyed by their Excel address.")
            total += len(item["cells"])
            if total > MAX_CELLS:
                raise ValueError("Too many populated cells for the browser editor.")
            for address, value in item["cells"].items():
                if not re.fullmatch(r"[A-Z]{1,3}[1-9][0-9]{0,6}", address):
                    raise ValueError(f"Invalid cell address: {address}")
                row, col = coordinate_to_tuple(address)
                if row > 1048576 or col > 16384:
                    raise ValueError(f"Cell {address} is outside Excel's limits.")
                _cell(value)
                previous = old and old["cells"].get(address)
                if value["kind"] == "readonly" and value != previous:
                    raise ValueError("Special formulas and rich text are read-only.")
            for address, value in (old or {}).get("cells", {}).items():
                if value["kind"] == "readonly" and item["cells"].get(address) != value:
                    raise ValueError("Special formulas and rich text must remain intact.")
        else:
            if set(item) != {"id", "shapes"} or not isinstance(item["shapes"], list):
                raise ValueError("A slide needs id and shapes.")
            old_shapes = {s["id"]: s for s in (old or {}).get("shapes", [])}
            ids = set()
            for shape in item["shapes"]:
                if not isinstance(shape, dict) or not isinstance(shape.get("id"), str):
                    raise ValueError("Each shape needs an ID.")
                sid = shape["id"]
                if sid in ids or (sid not in old_shapes and not re.fullmatch(r"new-[\w-]+", sid)):
                    raise ValueError("Use unique shape IDs; new shapes start with new-.")
                ids.add(sid)
                prior = old_shapes.get(sid)
                if prior and shape.get("kind") != prior["kind"]:
                    raise ValueError("An existing shape cannot change type.")
                if "box" in shape:
                    _check_box(shape["box"])
                elif prior and "box" in prior:
                    raise ValueError("Keep a shape's box; change its values to move it.")
                if shape.get("kind") == "text":
                    if not {"id", "kind", "paragraphs"} <= set(shape) <= {"id", "kind", "paragraphs", "box", "style"} or not isinstance(shape["paragraphs"], list) or not shape["paragraphs"]:
                        raise ValueError("A text shape needs at least one paragraph.")
                    for value in shape["paragraphs"]:
                        _text(value, soft_breaks=True)
                    if "style" in shape or (prior and "style" in prior):
                        _check_style(shape.get("style"), (prior or {}).get("style", {}))
                elif shape.get("kind") == "object":
                    if not prior or not {"id", "kind", "name"} <= set(shape) <= {"id", "kind", "name", "box"} or shape["name"] != prior["name"]:
                        raise ValueError("Pictures, charts and videos keep their name; only their box can change.")
                elif shape.get("kind") == "table":
                    rows = shape.get("rows")
                    if not {"id", "kind", "rows"} <= set(shape) <= {"id", "kind", "rows", "box"} or not isinstance(rows, list) or not rows:
                        raise ValueError("A table needs rows.")
                    if not isinstance(rows[0], list):
                        raise ValueError("Each table row must be a list of cells.")
                    width = len(rows[0])
                    if not width or any(not isinstance(r, list) or len(r) != width for r in rows):
                        raise ValueError("Each table row needs the same number of cells.")
                    if prior and (len(rows) != len(prior["rows"]) or width != len(prior["rows"][0])):
                        raise ValueError("Keep an existing table's dimensions to preserve its merges and design.")
                    for row in rows:
                        for value in row:
                            _text(value, soft_breaks=True)
                else:
                    raise ValueError("Editable shapes must be text, table or object.")
            if not set(old_shapes).issubset(ids):
                raise ValueError("Keep existing shapes; clear a text box's text to leave it blank.")
    if group == "sheets" and not set(known).issubset(seen):
        raise ValueError("Keep existing sheets to preserve formula references.")


def _set_paragraph(paragraph, text):
    """Keep existing run styles and hyperlinks around edits, including mixed runs."""
    if paragraph.text == text:
        return
    runs = paragraph.runs
    if not runs or paragraph.text != "".join(r.text for r in runs):
        # Soft breaks/fields need a whole-paragraph replacement, retaining pPr.
        style = copy(runs[0]._r.rPr) if runs and runs[0]._r.rPr is not None else None
        paragraph.text = text
        run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
        if style is not None and run._r.rPr is None:
            run._r.insert(0, style)
        return
    old = "".join(r.text for r in runs)
    owners = [i for i, r in enumerate(runs) for _ in r.text]
    values = [""] * len(runs)
    for tag, a, b, c, d in SequenceMatcher(None, old, text, autojunk=False).get_opcodes():
        if tag == "equal":
            for index in range(a, b):
                values[owners[index]] += old[index]
        elif tag in ("insert", "replace"):
            owner = owners[min(a, len(owners) - 1)] if owners else 0
            values[owner] += text[c:d]
    for run, value in zip(runs, values):
        run.text = value


def _set_frame(frame, paragraphs):
    for index, value in enumerate(paragraphs):
        paragraph = frame.paragraphs[index] if index < len(frame.paragraphs) else frame.add_paragraph()
        _set_paragraph(paragraph, value)
    for paragraph in list(frame.paragraphs)[len(paragraphs):]:
        frame._txBody.remove(paragraph._p)


def build_native(source, data, original, target):
    validate(data, original)
    if data["format"] == "xlsx":
        from openpyxl.cell.cell import MergedCell
        from openpyxl.styles import Font, PatternFill
        from openpyxl.workbook.properties import CalcProperties

        book = _workbook(source)
        known = {s["id"]: s for s in original["sheets"]}
        for item in data["sheets"]:
            old = known.get(item["id"])
            sheet = book[old["name"]] if old else book.create_sheet(item["name"])
            previous = old["cells"] if old else {}
            for address in previous.keys() | item["cells"].keys():
                value = item["cells"].get(address)
                if value == previous.get(address):
                    continue
                cell = sheet[address]
                if isinstance(cell, MergedCell):
                    raise ValueError(f"{sheet.title}!{address} belongs to a merged cell; edit its top-left cell.")
                if sheet.protection.sheet and cell.protection.locked:
                    raise ValueError(f"{sheet.title}!{address} is protected in the supplied workbook.")
                v = value["value"] if value else None
                kind = value["kind"] if value else "text"
                if kind in ("date", "datetime", "time"):
                    v = {"date": date, "datetime": datetime, "time": time}[kind].fromisoformat(v)
                cell.value = v
                if kind == "text" and v is not None:
                    cell.data_type = "s"  # A literal =… stays text, never an injected formula.
            if data["design"] == "new":
                for cell in sheet._cells.values():
                    if isinstance(cell, MergedCell):
                        continue
                    cell.font = Font(name="Aptos", size=11, bold=cell.row == 1,
                                     color="FFFFFF" if cell.row == 1 else "172B4D")
                    cell.fill = PatternFill("solid", fgColor="17365D" if cell.row == 1 else
                                            ("EFF4FA" if cell.row % 2 == 0 else "FFFFFF"))
                sheet.sheet_view.showGridLines = False
                sheet.freeze_panes = "A2"
                for column in sheet.column_dimensions.values():
                    column.width = 20
                sheet.sheet_properties.pageSetUpPr.fitToPage = True
                sheet.page_setup.fitToWidth = 1
                sheet.page_setup.fitToHeight = 0
        book._sheets = [book[s["name"]] for s in data["sheets"]]
        book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
        book.save(target)
        book.close()
    else:
        from pptx import Presentation
        from pptx.dml.color import RGBColor
        from pptx.util import Inches, Pt

        deck = Presentation(source)
        if "theme" in data:
            _write_theme(deck, data["theme"], original.get("theme", {}))
        known = {i["id"]: {s["id"]: s for s in i["shapes"]} for i in original["slides"]}
        existing = {str(s.slide_id): s for s in deck.slides}
        order = []
        for item in data["slides"]:
            slide = existing.get(item["id"])
            if slide is None:
                layout = min(deck.slide_layouts, key=lambda x: len(x.placeholders))
                slide = deck.slides.add_slide(layout)
                for placeholder in list(slide.placeholders):
                    placeholder._element.getparent().remove(placeholder._element)
            order.append(slide.slide_id)
            shapes = dict(walk_shapes(slide.shapes))
            for index, item_shape in enumerate(item["shapes"]):
                shape = shapes.get(item_shape["id"])
                if shape is None:
                    height = (deck.slide_height - Inches(1)) // max(1, len(item["shapes"]))
                    left, top = Inches(.6), Inches(.5) + index * height
                    width = deck.slide_width - Inches(1.2)
                    if item_shape["kind"] == "text":
                        shape = slide.shapes.add_textbox(left, top, width, height)
                    else:
                        rows = item_shape["rows"]
                        shape = slide.shapes.add_table(len(rows), len(rows[0]), left, top, width, height)
                prior = known.get(item["id"], {}).get(item_shape["id"], {})
                if "box" in item_shape and item_shape["box"] != prior.get("box"):
                    box = item_shape["box"]
                    shape.left, shape.top, shape.width, shape.height = (
                        round(box[k] * EMU_PER_PT) for k in ("x", "y", "w", "h"))
                if item_shape["kind"] == "object":
                    continue
                if item_shape["kind"] == "text":
                    _set_frame(shape.text_frame, item_shape["paragraphs"])
                    if "style" in item_shape:
                        _apply_style(shape, item_shape["style"], prior.get("style", {}))
                else:
                    for r, row in enumerate(item_shape["rows"]):
                        for c, value in enumerate(row):
                            cell = shape.table.cell(r, c)
                            if cell.is_spanned and value != cell.text:
                                raise ValueError("Edit the top-left cell of a merged presentation table.")
                            _set_frame(cell.text_frame, value.split("\n"))
            if data["design"] == "new":
                slide._element.set("showMasterSp", "0")
                slide.background.fill.solid()
                slide.background.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
                for _, shape in walk_shapes(slide.shapes):
                    if shape.shape_type in (1, 17):  # auto-shape or text box
                        shape.fill.solid()
                        shape.fill.fore_color.rgb = RGBColor.from_string("FFFFFF")
                    frames = [shape.text_frame] if shape.has_text_frame else []
                    if shape.has_table:
                        for r, row in enumerate(shape.table.rows):
                            for cell in row.cells:
                                cell.fill.solid()
                                cell.fill.fore_color.rgb = RGBColor.from_string("E8EFF7" if r == 0 else "FFFFFF")
                                frames.append(cell.text_frame)
                    for frame in frames:
                        for paragraph in frame.paragraphs:
                            for run in paragraph.runs:
                                run.font.name = "Aptos"
                                run.font.size = Pt(28 if shape == slide.shapes.title else 18)
                                run.font.color.rgb = RGBColor.from_string("17365D")
            # Retained images, charts, notes, groups, and relationships stay native.
        if "theme" in data:
            _swap(deck, data["theme"])  # after the shapes' own styles, so a swap reaches them too
        slide_ids = deck.slides._sldIdLst
        entries = {entry.id: entry for entry in slide_ids}
        for entry in list(slide_ids):
            slide_ids.remove(entry)
            if entry.id not in order:
                deck.part.drop_rel(entry.rId)
        for sid in order:
            slide_ids.append(entries[sid])
        deck.save(target)



class Conflict(ValueError):
    """Both the master and another file changed the same slide, shape, sheet or cell."""


def _merge_value(base, ours, theirs, where, conflicts):
    if ours == theirs or ours == base:
        return theirs
    if theirs == base:
        return ours
    conflicts.append(where)
    return ours


def _merge_children(base, ours, theirs, where, conflicts):
    """Shapes (a list keyed by id) or cells (a dict keyed by reference), merged three ways."""
    if isinstance(theirs, dict):
        out = {}
        for key in list(theirs) + [k for k in ours if k not in theirs]:
            b, o, t = base.get(key), ours.get(key), theirs.get(key)
            value = _merge_value(b, o, t, f"{where} {key}", conflicts)
            if value is not None:
                out[key] = value
        return out
    by = lambda items: {i["id"]: i for i in items}  # noqa: E731
    b, o, t = by(base), by(ours), by(theirs)
    out = []
    for key in [i["id"] for i in theirs] + [i["id"] for i in ours if i["id"] not in t]:
        value = _merge_value(b.get(key), o.get(key), t.get(key), f"{where} shape {key}", conflicts)
        if value is not None:
            out.append(value)
    return out


def merge(base, ours, theirs):
    """Three-way merge of Office models by slide or sheet ID, then by shape or cell. An item
    or value changed on one side only takes that side; changed on both, differently, is a
    conflict. Raises Conflict naming every one, so nothing is merged by guesswork."""
    group = "sheets" if theirs["format"] == "xlsx" else "slides"
    kids = "cells" if group == "sheets" else "shapes"
    conflicts = []
    by = lambda m: {i["id"]: i for i in m[group]}  # noqa: E731
    b, o, t = by(base), by(ours), by(theirs)
    order = [i["id"] for i in theirs[group]]
    for i, item in enumerate(ours[group]):  # an item only the master added goes after its neighbour
        if item["id"] not in t and item["id"] not in b:
            after = next((ours[group][j]["id"] for j in range(i - 1, -1, -1) if ours[group][j]["id"] in order), None)
            order.insert(order.index(after) + 1 if after else 0, item["id"])
    items = []
    for key in order:
        bi, oi, ti = b.get(key), o.get(key), t.get(key)
        where = f"{group[:-1]} {key}"
        if bi is not None and (oi is None or ti is None):  # deleted on one side
            kept = ti if oi is None else oi
            if kept != bi:
                conflicts.append(f"{where} (deleted on one side, changed on the other)")
                items.append(kept)
            continue
        if oi is None or ti is None or bi is None:
            items.append(oi or ti)
            continue
        item = {}
        for field in list(ti) + [f for f in oi if f not in ti]:
            if field == kids:
                item[kids] = _merge_children(bi.get(kids, [] if kids == "shapes" else {}), oi.get(kids, [] if kids == "shapes" else {}), ti.get(kids, [] if kids == "shapes" else {}), where, conflicts)
            else:
                value = _merge_value(bi.get(field), oi.get(field), ti.get(field), f"{where} {field}", conflicts)
                if value is not None:
                    item[field] = value
        items.append(item)
    out = dict(theirs, **{group: items})
    for field in ("design", "theme"):
        if field in theirs or field in ours:
            value = _merge_value(base.get(field), ours.get(field), theirs.get(field), field, conflicts)
            if value is not None:
                out[field] = value
    if conflicts:
        raise Conflict("Changed in both the editable document and the other file: " + "; ".join(conflicts[:20]))
    return out


def rebase(run, slug, path, base_path=None):
    """Fold a deck or workbook written outside the editor into the editable master, so the
    document stays one document. The file becomes the master's source (its slides, sheets and
    design are kept exactly) and the master's own edits are merged onto it. `base_path` is the
    file it was made from: by default the newest export of the master older than it, else the
    master's source. Changes nothing, and raises Conflict, when both changed the same thing."""
    import document

    run = Path(run).resolve()
    folder, master = paths(run, slug)
    source, original, meta = baseline(run, slug)
    path = Path(path).resolve()
    if not path.is_relative_to(run):
        raise ValueError("The file must be inside this project.")
    if base_path is None:
        made = []
        for applied in (run / "deliverable" / slug).glob("v*/applied.json"):
            info = json.loads(applied.read_text())
            native = applied.parent / str(info.get("file") or "")
            if info.get("from") == master.name and native.is_file() and native.stat().st_mtime <= path.stat().st_mtime and native != path:
                made.append(native)
        base_path = max(made, key=lambda f: f.stat().st_mtime) if made else source
    base = read_native(base_path)
    theirs = read_native(path)
    ours = decode(master.read_text())
    if theirs["format"] != original["format"]:
        raise ValueError("The file is not the same kind of document as the master.")
    merged = merge(base, ours, theirs)
    validate(merged, theirs)
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    for name in ("document.md", "office-source.json"):  # the old baseline, for the record
        shutil.copy2(folder / name, folder / f"{Path(name).stem}-before-{stamp}{Path(name).suffix}")
    shutil.copy2(path, folder / f"office-original.{theirs['format']}")
    (folder / "document.md").write_text(encode(theirs))
    (folder / "office-source.json").write_text(json.dumps({
        "format": theirs["format"], "sha256": digest(folder / f"office-original.{theirs['format']}"),
        "model_sha256": digest(folder / "document.md"), "source": str(path.relative_to(run)),
        "filename": path.name, "rebased_from": meta.get("source"), "merge_base": str(Path(base_path).resolve().relative_to(run)),
    }, indent=2) + "\n")
    master.write_text(encode(merged))
    document.register(run, slug, master)
    return master


def export(run, slug, design=None):
    import document

    folder, master = paths(run, slug)
    source, original, meta = baseline(run, slug)
    data = decode(master.read_text())
    master_digest = digest(master)
    if design is not None:
        data["design"] = design
    validate(data, original)
    if "[UNVERIFIED]" in master.read_text():
        raise ValueError("Resolve [UNVERIFIED] claims before exporting.")
    # Build and reopen before allocating a version; a failure leaves no empty round.
    with tempfile.TemporaryDirectory(prefix="office-export-") as temp:
        target = Path(temp) / f"document.{data['format']}"
        build_native(source, data, original, target)
        actual = read_native(target)
        group = "sheets" if data["format"] == "xlsx" else "slides"
        expected = json.loads(json.dumps(data))
        if group == "slides":
            _like(actual, expected)
        # Native IDs are assigned on creation; contents and order must round-trip.
        for model in (actual, expected):
            model.pop("design", None)
            for item in model[group]:
                item.pop("id", None)
                if group == "sheets":
                    for key, cell in list(item["cells"].items()):
                        if cell["kind"] == "text" and cell["value"] == "":
                            del item["cells"][key]
                        elif cell["kind"] == "date":
                            cell.update(kind="datetime", value=cell["value"] + "T00:00:00")
                for shape in item.get("shapes", []):
                    shape.pop("id", None)
        if actual != expected:
            raise ValueError("The native export did not retain all edited content; no export was published.")
        if digest(master) != master_digest:
            raise ValueError("The document changed during export. Export the saved version again.")
        version, out = document.next_version(run, slug)
        # A file named for its version ("v5 - Forecast.xlsx") is renamed for the new one.
        filename = re.sub(r"^v\d+(?=\s*-\s)", f"v{version}", Path(meta["filename"]).name)
        destination = out / filename
        shutil.copy2(target, destination)
        (out / "applied.json").write_text(json.dumps({
            "file": filename, "version": version, "format": data["format"],
            "from": master.name, "design": data["design"], "source_sha256": meta["sha256"],
            "master_sha256": master_digest,
        }, indent=2) + "\n")
        (out / "check.json").write_text(json.dumps({"ok": True, "problems": [], "check": "native content round-trip"}) + "\n")
    return str(destination.relative_to(run))
