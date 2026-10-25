#!/usr/bin/env python3
"""officedoc.py: Word, PowerPoint, Excel, PDF and Markdown files as addressable parts.

Shared by subject.py (extract), apply_changes.py (edit) and changes.py (check), so the
three agree on what `P012` or `S03.#5.p2` names. Every reader returns one shape:

    {"format": "docx", "parts": [{"id", "kind", "text", "section", ...}],
     "design": {...}, "references": {...}}

Part IDs, by format:

    docx   P012 paragraph · T2.R3.C1.P1 table-cell paragraph · T2 table · FN3.P1 footnote
    pptx   S03 slide · S03.#5 shape · S03.#5.p2 paragraph · S03.#7.R1.C2.p1 table cell
           S03.notes.p1 speaker notes
    xlsx   Sheet!C14 cell (sheet names may contain spaces; split on the last "!")
    pdf    PG03 page · PG03.b2 text block (read-only: a PDF is rebuilt, never edited)
    md     B012 block (paragraphs separated by blank lines)

IDs are positions in the file as supplied. An edit resolves every ID against the
original before changing anything, so an insertion never shifts a later target.

Visible text uses three tokens for things that are not text and cannot be split:
`[^3]` a footnote reference, `{obj}` an embedded object or image, `{fld}` a field.
"""
import re
import zipfile
from collections import Counter
from copy import deepcopy
from pathlib import Path

FORMATS = {".docx": "docx", ".pptx": "pptx", ".xlsx": "xlsx", ".xlsm": "xlsx",
           ".pdf": "pdf", ".md": "md", ".markdown": "md", ".txt": "md"}
EDITABLE = {"docx", "pptx", "xlsx", "md"}
PACKAGES = {"docx": "python-docx", "pptx": "python-pptx", "xlsx": "openpyxl", "pdf": "pymupdf"}

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
A = "http://schemas.openxmlformats.org/drawingml/2006/main"
P = "http://schemas.openxmlformats.org/presentationml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

OBJ, FLD = "{obj}", "{fld}"
TOKEN = re.compile(r"\[\^\d+\]|\{obj\}|\{fld\}")


def w(tag):
    return f"{{{W}}}{tag}"


def a(tag):
    return f"{{{A}}}{tag}"


def fmt_of(path):
    return FORMATS.get(Path(path).suffix.lower())


def need(fmt):
    raise SystemExit(f"ERROR reading .{fmt} needs {PACKAGES[fmt]}: "
                     f"python3 -m pip install {PACKAGES[fmt]}")


def read(path):
    """The model of one file. Raises SystemExit for a format it cannot read."""
    fmt = fmt_of(path)
    if fmt is None:
        raise SystemExit(f"ERROR {Path(path).name}: unsupported format "
                         f"(supported: {', '.join(sorted(set(FORMATS)))})")
    model = {"docx": read_docx, "pptx": read_pptx, "xlsx": read_xlsx,
             "pdf": read_pdf, "md": read_md}[fmt](Path(path))
    model["format"] = fmt
    model["references"] = references(model)
    if fmt == "docx":
        model["marks"] = working_marks(model)
    return model


def leaves(model):
    """The parts that hold text directly, in reading order."""
    return [p for p in model["parts"] if p["kind"] not in ("table", "slide", "shape", "sheet", "page")]


def strip_tokens(text):
    return TOKEN.sub("", text)


# The markup a change's `after` text may carry. Everything else is literal text.
#   **bold**   *italic*   ~~struck~~   [label](https://url)   ^[new footnote]   [^3] existing footnote
MARKUP = re.compile(
    r"~~(?P<s>.+?)~~"
    r"|\[\^(?P<ref>\d+)\](?!:)"
    r"|\*\*(?P<b>.+?)\*\*"
    r"|(?<![\w*])\*(?P<i>[^*\s][^*]*?)\*(?![\w*])"
    r"|\[(?P<lt>[^\]]+)\]\((?P<lu>(?:https?|mailto):[^)\s]+)\)"
    r"|\^\[(?P<fn>[^\]]+)\]")


def segments(text):
    """`after` text as [{text, bold?, italic?, link?, footnote?}], in order."""
    out, pos = [], 0
    for m in MARKUP.finditer(text):
        if m.start() > pos:
            out.append({"text": text[pos:m.start()]})
        if m.group("s") is not None:
            out += [{**seg, "strike": True} for seg in segments(m.group("s"))]
        elif m.group("ref") is not None:
            out.append({"text": "", "ref": m.group("ref")})
        elif m.group("b") is not None:
            out += [{**seg, "bold": True} for seg in segments(m.group("b"))]
        elif m.group("i") is not None:
            out.append({"text": m.group("i"), "italic": True})
        elif m.group("lt") is not None:
            out.append({"text": m.group("lt"), "link": m.group("lu")})
        else:
            out.append({"text": "", "footnote": m.group("fn")})
        pos = m.end()
    if pos < len(text):
        out.append({"text": text[pos:]})
    return out


def plain(text):
    """What a reader of the finished document sees of `after` text: footnote text sits
    elsewhere, and struck text is deleted text."""
    return "".join(seg["text"] for seg in segments(text) if not seg.get("strike"))


def _nearest(el, tag):
    el = el.getparent()
    while el is not None and el.tag != tag:
        el = el.getparent()
    return el


# ------------------------------------------------------------------ docx

class Atom:
    """One piece of a paragraph's visible text: a run's single content child."""
    __slots__ = ("run", "child", "text", "fixed")

    def __init__(self, run, child, text, fixed):
        self.run, self.child, self.text, self.fixed = run, child, text, fixed


DOCX_TEXT = {"t": None, "tab": "\t", "br": "\n", "cr": "\n", "noBreakHyphen": "-", "softHyphen": ""}


def docx_runs(p):
    """Runs whose own paragraph is `p`: not text-box runs, not tracked deletions."""
    out = []
    for r in p.iter(w("r")):
        if _nearest(r, w("p")) is not p:
            continue
        anc, deleted = r.getparent(), False
        while anc is not None and anc is not p:
            if anc.tag in (w("del"), w("moveFrom")):
                deleted = True
                break
            anc = anc.getparent()
        if not deleted:
            out.append(r)
    return out


def docx_atoms(p):
    atoms = []
    for r in docx_runs(p):
        for child in r:
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "rPr":
                continue
            if tag in DOCX_TEXT:
                text = child.text or "" if tag == "t" else DOCX_TEXT[tag]
                atoms.append(Atom(r, child, text, tag != "t"))
            elif tag in ("footnoteReference", "endnoteReference"):
                atoms.append(Atom(r, child, f"[^{child.get(w('id'))}]", True))
            elif tag in ("drawing", "pict", "object"):
                atoms.append(Atom(r, child, OBJ, True))
            else:                              # fldChar, instrText, sym, page breaks
                atoms.append(Atom(r, child, "", True))
    return atoms


def docx_text(p):
    return "".join(x.text for x in docx_atoms(p))


def docx_blocks(container):
    """Paragraphs and tables in reading order, descending into content controls."""
    for el in container:
        if el.tag in (w("p"), w("tbl")):
            yield el
        elif el.tag in (w("sdt"), w("customXml")):
            inner = el.find(w("sdtContent")) if el.tag == w("sdt") else el
            if inner is not None:
                yield from docx_blocks(inner)


def docx_footnotes_part(doc):
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/footnotes") and not rel.is_external:
            return rel.target_part
    return None


class XmlPart:
    """A package part as an lxml tree, whether python-docx parsed it or kept bytes.
    One tree per part: `XmlPart.of(part)` returns the same tree every time, so an edit
    made through one lookup is the edit that is saved."""

    def __init__(self, part):
        from lxml import etree
        self.part = part
        self.root = part.element if hasattr(part, "element") else etree.fromstring(part.blob)

    @classmethod
    def of(cls, part):
        cached = getattr(part, "_officedoc_xml", None)
        if cached is None:
            cached = cls(part)
            part._officedoc_xml = cached
        return cached

    def flush(self):
        from lxml import etree
        if not hasattr(self.part, "element"):
            self.part._blob = etree.tostring(self.root, xml_declaration=True,
                                             encoding="UTF-8", standalone=True)


def docx_footnotes(doc):
    """(XmlPart, [footnote elements that hold notes, not separators])."""
    part = docx_footnotes_part(doc)
    if part is None:
        return None, []
    xp = XmlPart.of(part)
    notes = [f for f in xp.root.findall(w("footnote"))
             if f.get(w("type")) in (None, "normal")]
    return xp, notes


def docx_index(doc):
    """Every addressable element of a python-docx Document, by ID, in reading order."""
    index = {}
    np_ = nt = 0
    for el in docx_blocks(doc.element.body):
        if el.tag == w("p"):
            np_ += 1
            index[f"P{np_:03d}"] = el
            continue
        nt += 1
        tid = f"T{nt}"
        index[tid] = el
        for ri, tr in enumerate(el.findall(w("tr")), 1):
            for ci, tc in enumerate(tr.findall(w("tc")), 1):
                for pi, p in enumerate(tc.findall(w("p")), 1):
                    index[f"{tid}.R{ri}.C{ci}.P{pi}"] = p
    _, notes = docx_footnotes(doc)
    for f in notes:
        for pi, p in enumerate(f.findall(w("p")), 1):
            index[f"FN{f.get(w('id'))}.P{pi}"] = p
    return index


def docx_style_names(doc):
    names = {}
    for s in doc.styles.element.findall(w("style")):
        n = s.find(w("name"))
        names[s.get(w("styleId"))] = n.get(w("val")) if n is not None else s.get(w("styleId"))
    return names


def docx_para_info(p, names):
    ppr = p.find(w("pPr"))
    sid = ppr.find(w("pStyle")).get(w("val")) if ppr is not None and ppr.find(w("pStyle")) is not None else None
    style = names.get(sid, sid) or "Normal"
    level = None
    m = re.match(r"(?i)heading\s*(\d)", style)
    if m:
        level = int(m.group(1))
    elif style.lower() == "title":
        level = 0
    elif ppr is not None and ppr.find(w("outlineLvl")) is not None:
        level = int(ppr.find(w("outlineLvl")).get(w("val"))) + 1
    listed = ppr is not None and ppr.find(w("numPr")) is not None
    return style, level, listed


def run_marks(run):
    """Working marks on a run: struck through, an explicit colour, a highlight."""
    rpr = run.find(w("rPr"))
    if rpr is None:
        return False, None, None
    on = lambda el: el is not None and el.get(w("val")) not in ("0", "false", "none")  # noqa: E731
    struck = on(rpr.find(w("strike"))) or on(rpr.find(w("dstrike")))
    col = rpr.find(w("color"))
    color = (col.get(w("val")) or "").upper() if col is not None else None
    hl = rpr.find(w("highlight"))
    return struck, color if color not in (None, "", "AUTO") else None, hl.get(w("val")) if on(hl) else None


def in_link(run):
    return run.getparent() is not None and run.getparent().tag == w("hyperlink")


def docx_part_marks(p):
    """Struck character ranges, draft colours (on text that is not a link) and comment anchors."""
    struck, colors, pos = [], {}, 0
    for x in docx_atoms(p):
        n = len(x.text)
        if n:
            is_struck, color, _ = run_marks(x.run)
            if is_struck:
                if struck and struck[-1][1] == pos:
                    struck[-1][1] = pos + n
                else:
                    struck.append([pos, pos + n])
            if color and not in_link(x.run) and x.text.strip():
                colors[color] = colors.get(color, 0) + n
        pos += n
    comments = sorted({el.get(w("id")) for tag in ("commentRangeStart", "commentReference")
                       for el in p.iter(w(tag))})
    return struck, colors, comments


def docx_comments(doc):
    """[{id, author, text}] from the comments part, if any."""
    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/comments") and not rel.is_external:
            root = XmlPart.of(rel.target_part).root
            return [{"id": c.get(w("id")), "author": c.get(w("author")) or "",
                     "text": " ".join("".join(t.text or "" for t in q.iter(w("t"))) for q in c.findall(w("p"))).strip()}
                    for c in root.findall(w("comment"))]
    return []


def read_docx(path):
    try:
        import docx
    except ImportError:
        need("docx")
    doc = docx.Document(str(path))
    names = docx_style_names(doc)
    index = docx_index(doc)
    parts, section = [], ""
    for pid, el in index.items():
        if el.tag == w("tbl"):
            rows = el.findall(w("tr"))
            parts.append({"id": pid, "kind": "table", "text": "", "section": section,
                          "rows": len(rows), "cols": max((len(r.findall(w("tc"))) for r in rows), default=0)})
            continue
        style, level, listed = docx_para_info(el, names)
        text = docx_text(el)
        kind = "footnote" if pid.startswith("FN") else "cell" if pid.startswith("T") else (
            "heading" if level is not None and text.strip() else "para")
        if kind == "heading" and not pid.startswith(("T", "FN")):
            section = text.strip()
        part = {"id": pid, "kind": kind, "text": text, "style": style, "section": section}
        if level is not None:
            part["level"] = level
        if listed:
            part["list"] = True
        struck, colors, comments = docx_part_marks(el)
        if struck:
            part["struck"] = struck
            part["unstruck"] = "".join(c for i, c in enumerate(text)
                                       if not any(a <= i < b for a, b in struck))
        if colors:
            part["colors"] = colors
        if comments:
            part["comments"] = comments
        parts.append(part)
    return {"parts": parts, "design": docx_design(doc, path, names, index),
            "links": docx_links(doc), "comments": docx_comments(doc)}


def working_marks(model):
    """What a draft carries that a finished document does not: struck text, draft colours,
    comments. A clean-up change removes them once the author accepts it."""
    parts = [p for p in model["parts"] if p.get("struck") or p.get("colors") or p.get("comments")]
    colors = {}
    for p in parts:
        for c, n in (p.get("colors") or {}).items():
            colors[c] = colors.get(c, 0) + n
    return {
        "struck_parts": [p["id"] for p in parts if p.get("struck")],
        "wholly_struck": [p["id"] for p in parts if p.get("struck") and not p.get("unstruck", "").strip()],
        "draft_colors": dict(sorted(colors.items(), key=lambda kv: -kv[1])),
        "colored_parts": [p["id"] for p in parts if p.get("colors")],
        "comments": [{**c, "anchors": [p["id"] for p in model["parts"] if c["id"] in (p.get("comments") or [])]}
                     for c in model.get("comments", [])],
    }


def docx_links(doc):
    """External hyperlinks in the body: text and target."""
    out = []
    for h in doc.element.body.iter(w("hyperlink")):
        rid = h.get(f"{{{R}}}id")
        if not rid or rid not in doc.part.rels:
            continue
        text = "".join(t.text or "" for t in h.iter(w("t")))
        out.append({"text": text, "url": doc.part.rels[rid].target_ref})
    return out


def rpr_values(root):
    """Explicit formatting values in a WordprocessingML tree, counted: the vocabulary a
    faithful edit may use. Values are (kind, value) pairs."""
    c = Counter()
    for rpr in root.iter(w("rPr")):
        for el in rpr:
            tag = el.tag.rsplit("}", 1)[-1]
            if tag == "rFonts":
                for k in ("ascii", "hAnsi"):
                    if el.get(w(k)):
                        c[("font", el.get(w(k)))] += 1
            elif tag == "sz":
                c[("size", el.get(w("val")))] += 1
            elif tag == "color" and el.get(w("val")):
                c[("color", el.get(w("val")).upper())] += 1
            elif tag == "highlight":
                c[("highlight", el.get(w("val")))] += 1
            elif tag == "shd" and el.get(w("fill")):
                c[("shading", el.get(w("fill")).upper())] += 1
    for tag, kind in (("pStyle", "style"), ("rStyle", "style"), ("tblStyle", "style")):
        for el in root.iter(w(tag)):
            c[(kind, el.get(w("val")))] += 1
    return c


def _zip_xml(path, name):
    from lxml import etree
    try:
        with zipfile.ZipFile(path) as z:
            return etree.fromstring(z.read(name))
    except (KeyError, zipfile.BadZipFile):
        return None


def theme_info(root):
    """Colour scheme and font scheme of a DrawingML theme."""
    if root is None:
        return {}
    colors, fonts = {}, {}
    scheme = root.find(f".//{a('clrScheme')}")
    if scheme is not None:
        for slot in scheme:
            val = slot.find(a("srgbClr"))
            sysc = slot.find(a("sysClr"))
            colors[slot.tag.rsplit("}", 1)[-1]] = (val.get("val") if val is not None else
                                                    sysc.get("lastClr") if sysc is not None else None)
    for which in ("majorFont", "minorFont"):
        el = root.find(f".//{a(which)}/{a('latin')}")
        if el is not None:
            fonts[which.replace("Font", "")] = el.get("typeface")
    return {"colors": colors, "fonts": fonts}


def docx_design(doc, path, names, index):
    body = doc.element.body
    sect = body.find(w("sectPr"))
    page = {}
    if sect is not None:
        sz, mar = sect.find(w("pgSz")), sect.find(w("pgMar"))
        if sz is not None:
            page = {"width_pt": int(sz.get(w("w"), 0)) / 20, "height_pt": int(sz.get(w("h"), 0)) / 20}
        if mar is not None:
            page["margins_pt"] = {k: int(mar.get(w(k), 0)) / 20 for k in ("top", "bottom", "left", "right")}
    styles = {}
    st_root = doc.styles.element
    defaults = {}
    dfl = st_root.find(f"{w('docDefaults')}/{w('rPrDefault')}/{w('rPr')}")
    if dfl is not None:
        defaults = _rpr_summary(dfl)
    used = Counter()
    for pid, el in index.items():
        if el.tag == w("p"):
            used[docx_para_info(el, names)[0]] += len(docx_text(el))
    for s in st_root.findall(w("style")):
        name = names.get(s.get(w("styleId")))
        if name not in used and s.get(w("styleId")) not in ("Hyperlink", "FootnoteReference", "FootnoteText"):
            continue
        rpr, ppr = s.find(w("rPr")), s.find(w("pPr"))
        entry = {"id": s.get(w("styleId")), "type": s.get(w("type")), "chars": used.get(name, 0)}
        based = s.find(w("basedOn"))
        if based is not None:
            entry["based_on"] = based.get(w("val"))
        if rpr is not None:
            entry.update(_rpr_summary(rpr))
        if ppr is not None:
            sp = ppr.find(w("spacing"))
            if sp is not None:
                entry["spacing"] = {k: sp.get(w(k)) for k in ("before", "after", "line") if sp.get(w(k))}
        styles[name] = entry
    values = rpr_values(body)
    xp, _ = docx_footnotes(doc)
    if xp is not None:
        values += rpr_values(xp.root)
    link_rpr = None
    for h in body.iter(w("hyperlink")):
        r = h.find(w("r"))
        if h.get(f"{{{R}}}id") and r is not None and r.find(w("rPr")) is not None:
            link_rpr = _rpr_summary(r.find(w("rPr")))
            break
    return {
        "page": page,
        "theme": theme_info(_zip_xml(path, "word/theme/theme1.xml")),
        "defaults": defaults,
        "styles": styles,
        "hyperlink": link_rpr,
        "fonts": _top(values, "font"),
        "sizes_pt": {str(int(k) / 2): n for k, n in _top(values, "size").items()},
        "colors": _top(values, "color"),
        "shading": _top(values, "shading"),
        "vocabulary": sorted(f"{k}:{v}" for k, v in values),
    }


def _top(counter, kind):
    return {v: n for (k, v), n in counter.most_common() if k == kind}


def _rpr_summary(rpr):
    out = {}
    f = rpr.find(w("rFonts"))
    if f is not None and (f.get(w("ascii")) or f.get(w("asciiTheme"))):
        out["font"] = f.get(w("ascii")) or f"theme:{f.get(w('asciiTheme'))}"
    sz = rpr.find(w("sz"))
    if sz is not None:
        out["size_pt"] = int(sz.get(w("val"))) / 2
    col = rpr.find(w("color"))
    if col is not None:
        out["color"] = col.get(w("val"))
    for tag in ("b", "i", "u"):
        el = rpr.find(w(tag))
        if el is not None and el.get(w("val")) not in ("0", "false", "none"):
            out[{"b": "bold", "i": "italic", "u": "underline"}[tag]] = True
    return out


# ------------------------------------------------------------------ pptx

def pptx_atoms(p):
    atoms = []
    for child in p:
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "r":
            t = child.find(a("t"))
            atoms.append(Atom(child, t, (t.text or "") if t is not None else "", False))
        elif tag == "br":
            atoms.append(Atom(child, child, "\n", True))
        elif tag == "fld":
            atoms.append(Atom(child, child, FLD, True))
    return atoms


def pptx_text(p):
    return "".join(x.text for x in pptx_atoms(p))


def iter_shapes(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type is not None and sh.shape_type == 6:      # MSO_SHAPE_TYPE.GROUP
            yield from iter_shapes(sh.shapes)


def pptx_index(prs):
    """ID -> (element, slide, extra). Slides map to the slide object itself."""
    index = {}
    for si, slide in enumerate(prs.slides, 1):
        sid = f"S{si:02d}"
        index[sid] = (slide, slide, {})
        for sh in iter_shapes(slide.shapes):
            base = f"{sid}.#{sh.shape_id}"
            index[base] = (sh._element, slide, {"shape": sh})
            if sh.has_text_frame:
                for k, p in enumerate(sh.text_frame._txBody.findall(a("p")), 1):
                    index[f"{base}.p{k}"] = (p, slide, {"shape": sh})
            if getattr(sh, "has_table", False) and sh.has_table:
                for ri, tr in enumerate(sh.table._tbl.findall(a("tr")), 1):
                    for ci, tc in enumerate(tr.findall(a("tc")), 1):
                        body = tc.find(a("txBody"))
                        for k, p in enumerate(body.findall(a("p")) if body is not None else [], 1):
                            index[f"{base}.R{ri}.C{ci}.p{k}"] = (p, slide, {"shape": sh})
        if slide.has_notes_slide:
            tf = slide.notes_slide.notes_text_frame
            if tf is not None:
                for k, p in enumerate(tf._txBody.findall(a("p")), 1):
                    index[f"{sid}.notes.p{k}"] = (p, slide, {"notes": True})
    return index


def read_pptx(path):
    try:
        from pptx import Presentation
    except ImportError:
        need("pptx")
    prs = Presentation(str(path))
    parts = []
    section = ""
    for pid, (el, slide, extra) in pptx_index(prs).items():
        if el is slide:
            title = slide.shapes.title.text_frame.text if slide.shapes.title is not None and slide.shapes.title.has_text_frame else ""
            section = f"{pid} {title}".strip()
            parts.append({"id": pid, "kind": "slide", "text": title, "section": section,
                          "layout": slide.slide_layout.name, "slide_id": slide.slide_id})
        elif el.tag == a("p"):
            kind = "notes" if extra.get("notes") else "cell" if ".R" in pid else "para"
            ppr = el.find(a("pPr"))
            part = {"id": pid, "kind": kind, "text": pptx_text(el), "section": section}
            if ppr is not None and ppr.get("lvl"):
                part["level"] = int(ppr.get("lvl"))
            parts.append(part)
        else:
            sh = extra["shape"]
            emu = 12700
            part = {"id": pid, "kind": "shape", "section": section, "name": sh.name,
                    "text": "\n".join(pptx_text(q) for q in sh.text_frame._txBody.findall(a("p")))
                    if sh.has_text_frame else "",
                    "box_pt": [round((v or 0) / emu, 1) for v in (sh.left, sh.top, sh.width, sh.height)]}
            if sh.is_placeholder:
                part["placeholder"] = str(sh.placeholder_format.type).split(".")[-1].split(" ")[0]
            if sh.shape_type == 13:
                part["picture"] = True
            parts.append(part)
    return {"parts": parts, "design": pptx_design(prs, path)}


def drawing_values(root):
    """Explicit DrawingML formatting values: fonts, sizes, colours."""
    c = Counter()
    for tag in ("rPr", "defRPr", "endParaRPr"):
        for rpr in root.iter(a(tag)):
            if rpr.get("sz"):
                c[("size", str(int(rpr.get("sz")) / 100))] += 1
            for ft in ("latin", "ea", "cs"):
                el = rpr.find(a(ft))
                if el is not None and el.get("typeface") and not el.get("typeface").startswith("+"):
                    c[("font", el.get("typeface"))] += 1
    for clr in root.iter(a("srgbClr")):
        c[("color", (clr.get("val") or "").upper())] += 1
    for clr in root.iter(a("schemeClr")):
        c[("scheme", clr.get("val"))] += 1
    return c


def pptx_design(prs, path):
    from lxml import etree
    emu = 12700
    values = Counter()
    for slide in prs.slides:
        values += drawing_values(slide._element)
    layouts = []
    for master in prs.slide_masters:
        for lay in master.slide_layouts:
            layouts.append({"name": lay.name, "placeholders": [
                {"idx": ph.placeholder_format.idx,
                 "type": str(ph.placeholder_format.type).split(".")[-1].split(" ")[0],
                 "name": ph.name} for ph in lay.placeholders]})
    theme = None
    for rel in prs.slide_masters[0].part.rels.values():
        if rel.reltype.endswith("/theme"):
            theme = etree.fromstring(rel.target_part.blob)
    used = Counter(s.slide_layout.name for s in prs.slides)
    return {
        "page": {"width_pt": prs.slide_width / emu, "height_pt": prs.slide_height / emu},
        "theme": theme_info(theme),
        "layouts": layouts,
        "layouts_used": dict(used),
        "fonts": _top(values, "font"),
        "sizes_pt": _top(values, "size"),
        "colors": _top(values, "color"),
        "scheme_colors_used": _top(values, "scheme"),
        "vocabulary": sorted(f"{k}:{v}" for k, v in values),
    }


# ------------------------------------------------------------------ xlsx

def split_cell_id(cid):
    sheet, _, coord = cid.rpartition("!")
    return sheet, coord


def read_xlsx(path):
    try:
        import openpyxl
    except ImportError:
        need("xlsx")
    wb = openpyxl.load_workbook(str(path), data_only=False)
    cached = openpyxl.load_workbook(str(path), data_only=True)
    parts = []
    fonts, fills, numfmts = Counter(), Counter(), Counter()
    for ws in wb.worksheets:
        parts.append({"id": ws.title, "kind": "sheet", "text": ws.title, "section": ws.title,
                      "dims": ws.dimensions})
        cws = cached[ws.title]
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                text = str(cell.value)
                part = {"id": f"{ws.title}!{cell.coordinate}", "section": ws.title,
                        "kind": "formula" if text.startswith("=") else "value", "text": text}
                if part["kind"] == "formula":
                    part["cached"] = cws[cell.coordinate].value
                if cell.number_format and cell.number_format != "General":
                    part["format"] = cell.number_format
                    numfmts[cell.number_format] += 1
                f = cell.font
                fonts[(f.name, f.sz, bool(f.b), f.color.rgb if f.color is not None and isinstance(f.color.rgb, str) else None)] += 1
                if cell.fill is not None and cell.fill.fill_type == "solid" and isinstance(cell.fill.fgColor.rgb, str):
                    fills[cell.fill.fgColor.rgb] += 1
                parts.append(part)
    design = {
        "sheets": wb.sheetnames,
        "fonts": [{"name": n, "size": s, "bold": b, "color": c, "cells": k}
                  for (n, s, b, c), k in fonts.most_common(20)],
        "fills": dict(fills.most_common(20)),
        "number_formats": dict(numfmts.most_common(20)),
        "defined_names": sorted(list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else []),
    }
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
    design["charts"] = sum(1 for n in names if n.startswith("xl/charts/chart"))
    design["drawings"] = sum(1 for n in names if n.startswith("xl/drawings/drawing"))
    return {"parts": parts, "design": design}


# ------------------------------------------------------------------ pdf

def _pdf_open(path):
    try:
        import pymupdf as fitz
    except ImportError:
        try:
            import fitz
        except ImportError:
            need("pdf")
    return fitz.open(str(path))


def _family(font):
    base = font.split("+")[-1]
    fam, _, style = base.partition("-")
    fam = re.sub(r"(MT|PSMT|PS)$", "", fam)
    weight = "bold" if re.search(r"(?i)bold|black|heavy", style + fam) else (
        "light" if re.search(r"(?i)light|thin", style + fam) else "regular")
    italic = bool(re.search(r"(?i)italic|oblique", style))
    return fam, weight, italic


def read_pdf(path):
    pdf = _pdf_open(path)
    parts = []
    fonts, sizes, colors, fills = Counter(), Counter(), Counter(), Counter()
    roles = Counter()
    frames = []
    for pi, page in enumerate(pdf, 1):
        pid = f"PG{pi:02d}"
        blocks = [b for b in page.get_text("dict")["blocks"] if b.get("type") == 0]
        page_text = []
        for bi, b in enumerate(blocks, 1):
            lines = []
            for line in b["lines"]:
                spans = line["spans"]
                lines.append("".join(s["text"] for s in spans))
                for s in spans:
                    n = len(s["text"].strip())
                    if not n:
                        continue
                    fam, weight, italic = _family(s["font"])
                    fonts[fam] += n
                    size = round(s["size"], 1)
                    sizes[size] += n
                    col = "%06X" % s["color"]
                    colors[col] += n
                    roles[(fam, weight, italic, size, col)] += n
            text = "\n".join(lines).strip()
            if text:
                page_text.append(text)
                parts.append({"id": f"{pid}.b{bi}", "kind": "block", "text": text,
                              "section": pid, "box_pt": [round(v, 1) for v in b["bbox"]]})
        parts.insert(len(parts) - len(page_text), {"id": pid, "kind": "page",
                                                   "text": page_text[0][:120] if page_text else "",
                                                   "section": pid})
        frame = None
        page_area = page.rect.width * page.rect.height
        for d in page.get_drawings():
            fill = d.get("fill")
            if fill and d.get("rect") is not None:
                hexv = "%02X%02X%02X" % tuple(int(round(c * 255)) for c in fill[:3])
                fills[hexv] += int(d["rect"].width * d["rect"].height)
                if hexv != "FFFFFF" and d["rect"].width * d["rect"].height > page_area * 0.01:
                    frame = _union(frame, tuple(d["rect"]))
        for img in page.get_image_info():
            if img.get("bbox"):
                frame = _union(frame, tuple(img["bbox"]))
        if frame is not None:
            frames.append(frame)
    first = pdf[0].rect if pdf.page_count else None
    design = {
        "page": {"width_pt": round(first.width, 1), "height_pt": round(first.height, 1)} if first else {},
        "pages": pdf.page_count,
        "content_box_pt": _median_box(frames),
        "fonts": dict(fonts.most_common()),
        "sizes_pt": {str(k): v for k, v in sizes.most_common(25)},
        "colors": dict(colors.most_common(15)),
        "fills": dict(fills.most_common(15)),
        "text_styles": [{"font": f, "weight": wt, "italic": it, "size_pt": sz, "color": c, "chars": n}
                        for (f, wt, it, sz, c), n in roles.most_common(20)],
    }
    return {"parts": parts, "design": design}


def _union(box, other):
    if box is None:
        return list(other[:4])
    return [min(box[0], other[0]), min(box[1], other[1]), max(box[2], other[2]), max(box[3], other[3])]


def _median_box(frames):
    """Where the slide sits on the page, when a deck was exported onto a larger sheet."""
    if not frames:
        return None
    cols = list(zip(*frames))
    mid = [sorted(c)[len(c) // 2] for c in cols]
    return {"x": round(mid[0], 1), "y": round(mid[1], 1),
            "width": round(mid[2] - mid[0], 1), "height": round(mid[3] - mid[1], 1)}


def render_pdf(path, out_dir, dpi=80):
    """Page images and embedded pictures of a PDF, for an agent to look at."""
    pdf = _pdf_open(path)
    out_dir = Path(out_dir)
    (out_dir / "pages").mkdir(parents=True, exist_ok=True)
    (out_dir / "images").mkdir(parents=True, exist_ok=True)
    pages, images, seen = [], [], set()
    for pi, page in enumerate(pdf, 1):
        target = out_dir / "pages" / f"page-{pi:02d}.png"
        page.get_pixmap(dpi=dpi).save(str(target))
        pages.append(str(target.relative_to(out_dir)))
        for img in page.get_images(full=True):
            xref = img[0]
            if xref in seen:
                continue
            seen.add(xref)
            info = pdf.extract_image(xref)
            if not info or info.get("width", 0) < 48 or info.get("height", 0) < 48:
                continue
            target = out_dir / "images" / f"p{pi:02d}-x{xref}.{info.get('ext', 'png')}"
            target.write_bytes(info["image"])
            images.append({"file": str(target.relative_to(out_dir)), "page": pi,
                           "width": info["width"], "height": info["height"]})
    return pages, images


# ------------------------------------------------------------------ md

def md_blocks(text):
    """Blocks separated by blank lines, keeping each block's exact text."""
    blocks, cur = [], []
    for line in text.splitlines():
        if line.strip():
            cur.append(line)
        elif cur:
            blocks.append("\n".join(cur))
            cur = []
    if cur:
        blocks.append("\n".join(cur))
    return blocks


def read_md(path):
    parts, section = [], ""
    for i, block in enumerate(md_blocks(path.read_text(errors="replace")), 1):
        m = re.match(r"(#{1,6})\s+(.*)", block)
        if m:
            section = m.group(2).strip()
        part = {"id": f"B{i:03d}", "kind": "heading" if m else "para", "text": block, "section": section}
        if m:
            part["level"] = len(m.group(1))
        parts.append(part)
    return {"parts": parts, "design": {}}


# ------------------------------------------------------------------ references

NUMERIC = re.compile(r"\[(\d+(?:\s*[-–,]\s*\d+)*)\]")
AUTHOR_DATE = re.compile(r"\((?:[A-Z][^()]{1,80}?),?\s+(?:\d{1,2}\s+\w+\s+)?(?:\d{4}[a-z]?|n\.d\.)\)")
PAREN_SOURCE = re.compile(r"\((?:[A-Z][\w&.'’-]*(?:\s+[A-Z0-9][\w&.'’-]*){0,5})(?:,\s*[^()]{1,40})?\)")
ACRONYM = re.compile(r"^\([A-Z0-9&/-]{1,8}s?\)$")
REF_HEADING = re.compile(r"(?i)^\W*(?:\d+\.?\s*)?(references|sources|bibliography|works cited|notes|endnotes|citations)\b")
ENTRY_NUM = (("[n]", re.compile(r"^\s*\[(\d+)\]")), ("n.", re.compile(r"^\s*(\d+)[.)]\s")))


def references(model):
    """How the document cites, so new references can follow the same system."""
    texts = [p for p in leaves(model) if p["kind"] != "footnote"]
    counts = Counter()
    samples = {}
    numbers = []
    for p in texts:
        t = p["text"]
        for kind, rx in (("numbered", NUMERIC), ("author_date", AUTHOR_DATE), ("parenthetical_source", PAREN_SOURCE)):
            for m in rx.finditer(t):
                if kind == "parenthetical_source" and (AUTHOR_DATE.match(m.group(0))
                                                       or ACRONYM.match(m.group(0))):
                    continue
                counts[kind] += 1
                if kind == "numbered":
                    numbers += [int(n) for n in re.findall(r"\d+", m.group(1))]
                samples.setdefault(kind, [])
                if len(samples[kind]) < 5:
                    lo = max(0, m.start() - 60)
                    samples[kind].append({"id": p["id"], "context": t[lo:m.end() + 20]})
        counts["footnote_refs"] += len(re.findall(r"\[\^\d+\]", t))
    footnotes = [p for p in model["parts"] if p["kind"] == "footnote"]
    links = [x for x in model.get("links", []) if x.get("url", "").startswith(("http", "mailto"))]
    if links:
        counts["links"] = len(links)
        samples["links"] = links[:5]
        bodies = [p["text"] for p in texts]
        inside = 0
        for x in links:
            t = re.escape(x["text"].strip())
            if t and any(re.search(r"\([^()]*" + t + r"[^()]*\)", b) for b in bodies):
                inside += 1
        if inside:
            counts["linked_parenthetical"] = inside
    ref_list = None
    parts = model["parts"]
    for i, p in enumerate(parts):
        if p["kind"] in ("heading", "slide") and REF_HEADING.search(p["text"].strip()) or (
                p["kind"] == "para" and len(p["text"]) < 40 and REF_HEADING.search(p["text"].strip())
                and p["text"].strip().rstrip(":").lower() in ("references", "sources", "bibliography", "notes")):
            entries = []
            for q in parts[i + 1:]:
                if q["kind"] in ("heading", "slide"):
                    break
                if q["kind"] in ("para", "cell") and q["text"].strip():
                    entries.append(q)
            numbering, last = "none", None
            for label, rx in ENTRY_NUM:
                nums = [int(m.group(1)) for e in entries for m in [rx.match(e["text"])] if m]
                if nums and len(nums) >= len(entries) / 2:
                    numbering, last = label, max(nums)
                    break
            if numbering == "none" and entries and all(e.get("list") for e in entries):
                numbering, last = "auto-numbered list", len(entries)
            ref_list = {"heading": p["id"], "entries": [e["id"] for e in entries],
                        "numbering": numbering, "last": last,
                        "first_entry": entries[0]["text"][:200] if entries else None}
    systems = ("numbered", "author_date", "linked_parenthetical", "parenthetical_source", "footnote_refs")
    ranked = sorted((k for k in systems if counts[k]), key=lambda k: -counts[k])
    if counts["linked_parenthetical"] and counts["linked_parenthetical"] >= counts["parenthetical_source"] / 3:
        ranked.insert(0, ranked.pop(ranked.index("linked_parenthetical")))
    system = (ranked[0].replace("footnote_refs", "footnotes") if ranked
              else "inline_links" if counts["links"] else "none")
    out = {"system": system, "counts": dict(counts), "samples": samples,
           "footnotes": {"count": len({p["id"].split(".")[0] for p in footnotes}),
                         "ids": sorted({int(p["id"][2:].split(".")[0]) for p in footnotes})}}
    if numbers:
        out["numbered"] = {"max": max(numbers), "next": max(numbers) + 1}
    if ref_list:
        out["reference_list"] = ref_list
        if ref_list["last"]:
            out.setdefault("numbered", {"max": ref_list["last"], "next": ref_list["last"] + 1})
    return out


# ------------------------------------------------------------------ vocabulary

def vocabulary(path, slides_only=False):
    """Every explicit formatting value in a docx or pptx: what a faithful file may use.
    `slides_only` leaves out a deck's masters, layouts and theme."""
    from lxml import etree
    fmt = fmt_of(path)
    c = Counter()
    with zipfile.ZipFile(path) as z:
        for name in z.namelist():
            if not name.endswith(".xml"):
                continue
            if fmt == "docx" and re.match(r"word/(document|footnotes|endnotes|header\d*|footer\d*)\.xml$", name):
                c += rpr_values(etree.fromstring(z.read(name)))
            elif fmt == "pptx" and re.match(r"ppt/slides/slide\d+\.xml$" if slides_only else
                                            r"ppt/(slides/slide|slideLayouts/slideLayout|slideMasters/slideMaster|theme/theme)\d*\.xml$", name):
                c += drawing_values(etree.fromstring(z.read(name)))
    return c


# ------------------------------------------------------------------ docx as Markdown

def list_kinds(doc):
    """numId -> "bullet" or "number", from the document's numbering part."""
    try:
        numbering = doc.part.numbering_part.element
    except (KeyError, NotImplementedError, AttributeError):
        return {}
    abstract = {a_.get(w("abstractNumId")): a_ for a_ in numbering.findall(w("abstractNum"))}
    out = {}
    for num in numbering.findall(w("num")):
        ref = num.find(w("abstractNumId"))
        a_el = abstract.get(ref.get(w("val"))) if ref is not None else None
        fmt = a_el.find(f"{w('lvl')}/{w('numFmt')}") if a_el is not None else None
        if fmt is not None:
            out[num.get(w("numId"))] = "bullet" if fmt.get(w("val")) == "bullet" else "number"
    return out


def _on(rpr, tag):
    el = rpr.find(w(tag)) if rpr is not None else None
    return el is not None and el.get(w("val")) not in ("0", "false", "none")


def para_markdown(p, part):
    """A paragraph's text as Markdown: **bold**, *italic*, ~~struck~~, [links](url), [^n]."""
    groups = []                       # [flags, text]
    for r in docx_runs(p):
        text = "".join(x.text for x in docx_atoms(p) if x.run is r)
        if not text:
            continue
        rpr = r.find(w("rPr"))
        url = None
        parent = r.getparent()
        if parent.tag == w("hyperlink") and parent.get(f"{{{R}}}id") in part.rels:
            url = part.rels[parent.get(f"{{{R}}}id")].target_ref
        flags = (_on(rpr, "b") and not url, _on(rpr, "i") and not url,
                 _on(rpr, "strike") or _on(rpr, "dstrike"), url)
        if groups and groups[-1][0] == flags:
            groups[-1][1] += text
        else:
            groups.append([flags, text])
    pieces = []                       # [struck, markdown], neighbours merged below
    for (bold, italic, struck, url), text in groups:
        if text.startswith("[^") and text.endswith("]") or not text.strip():
            pieces.append([struck, text])
            continue
        lead, core, trail = re.match(r"(\s*)(.*?)(\s*)$", text, re.S).groups()
        if url:
            core = f"[{core}]({url})"
        if italic:
            core = f"*{core}*"
        if bold:
            core = f"**{core}**"
        pieces.append([struck, lead + core + trail])
    out, i = [], 0
    while i < len(pieces):
        if not pieces[i][0]:
            out.append(pieces[i][1])
            i += 1
            continue
        j = i
        while j < len(pieces) and pieces[j][0]:
            j += 1
        text = "".join(x[1] for x in pieces[i:j])
        lead, core, trail = re.match(r"(\s*)(.*?)(\s*)$", text, re.S).groups()
        out.append(f"{lead}~~{core}~~{trail}" if core else text)
        i = j
    return "".join(out).replace("\n", " ")


def docx_markdown(path):
    """The document as Markdown, one block per paragraph or table, each under a hidden
    `<!-- ID -->` anchor. Comments sit under the block they are attached to; footnotes
    come last. The anchors let an edited copy be mapped back onto the original."""
    import docx
    doc = docx.Document(str(path))
    names = docx_style_names(doc)
    index = docx_index(doc)
    kinds = list_kinds(doc)
    comments = {c["id"]: c for c in docx_comments(doc)}
    fn_part = docx_footnotes_part(doc)
    blocks = []
    for pid, el in index.items():
        if pid.startswith("FN") or (pid.startswith("T") and "." in pid):
            continue
        if el.tag == w("tbl"):
            rows = []
            for tr in el.findall(w("tr")):
                cells = []
                for tc in tr.findall(w("tc")):
                    cells.append(" <br> ".join(para_markdown(q, doc.part) for q in tc.findall(w("p"))
                                               if docx_text(q).strip()).replace("|", "\\|"))
                rows.append("| " + " | ".join(cells) + " |")
            if rows:
                rows.insert(1, "|" + " --- |" * rows[0].count(" | ") + " --- |")
            blocks.append(f"<!-- {pid} -->\n" + "\n".join(rows))
            continue
        md = para_markdown(el, doc.part)
        if not md.strip():
            continue
        style, level, _ = docx_para_info(el, names)
        nid = el.find(f"{w('pPr')}/{w('numPr')}/{w('numId')}")
        if level is not None:
            md = "#" * max(1, min(6, level or 1)) + " " + md
        elif nid is not None and kinds.get(nid.get(w("val"))):
            md = ("- " if kinds[nid.get(w("val"))] == "bullet" else "1. ") + md
        block = f"<!-- {pid} -->\n{md}"
        for cid in sorted({c.get(w("id")) for tag in ("commentRangeStart", "commentReference")
                           for c in el.iter(w(tag))}):
            if cid in comments:
                c = comments[cid]
                block += f"\n<!-- comment {cid} ({c['author']}): {c['text']} -->"
        blocks.append(block)
    _, notes = docx_footnotes(doc)
    for f in notes:
        fid = f.get(w("id"))
        text = " ".join(para_markdown(q, fn_part) for q in f.findall(w("p"))).strip()
        blocks.append(f"<!-- FN{fid} -->\n[^{fid}]: {text}")
    return "\n\n".join(blocks) + "\n"
