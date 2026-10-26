"""Draw a PowerPoint deck as pictures, exactly as PowerPoint shows it.

The app's HTML view uses `powerpoint_pdf` for its slide pictures; `/slides` uses the
command line to look at a deck before and after a redesign.

    python3 .claude/scripts/deck_render.py <deck.pptx | run/<slug>.md> <out_dir> [--width 1600]

A `<slug>.md` deck master is built into a .pptx first, with its edits applied. The
pictures are `slide-001.png` onward; a hidden slide keeps its number.
"""

import argparse
import subprocess
import sys
import tempfile
from pathlib import Path

POWERPOINT = Path("/Applications/Microsoft PowerPoint.app")

# Opens a copy of the deck, exports it to PDF and closes it, leaving any deck the
# author has open alone, and quits PowerPoint only if this started it. The paths are
# written in as literals: compiled that way they reach PowerPoint as file objects that
# carry sandbox access, where a path built at run time makes it ask the author.
EXPORT = """
set wasRunning to application "Microsoft PowerPoint" is running
tell application "Microsoft PowerPoint"
  with timeout of 600 seconds
    open POSIX file {src}
    set p to missing value
    repeat 240 times
      try
        set p to presentation {name}
        exit repeat
      end try
      delay 0.5
    end repeat
    if p is missing value then error "PowerPoint did not open the deck"
    save p in POSIX file {out} as save as PDF
    close p saving no
  end timeout
  if not wasRunning and (count of presentations) is 0 then quit
end tell
"""


def _applescript(text):
    return '"' + text.replace("\\", "\\\\").replace('"', '\\"') + '"'


def powerpoint_pdf(src, pdf):
    src, pdf = Path(src), Path(pdf)
    if sys.platform != "darwin" or not POWERPOINT.is_dir():
        raise RuntimeError("PowerPoint is not installed on this computer")
    script = EXPORT.format(src=_applescript(str(src)), out=_applescript(str(pdf)), name=_applescript(src.name))
    done = subprocess.run(["osascript", "-e", script], capture_output=True, text=True, timeout=900)
    if done.returncode or not pdf.is_file():
        detail = done.stderr.strip()[:300]
        raise RuntimeError("PowerPoint could not export the deck" + (f" ({detail})" if detail else ""))


def pictures(deck_path, out_dir, width=1600):
    """Write one PNG per slide into out_dir; return their paths in slide order."""
    import pymupdf
    from pptx import Presentation

    deck_path, out_dir = Path(deck_path).resolve(), Path(out_dir).resolve()
    out_dir.mkdir(parents=True, exist_ok=True)
    deck = Presentation(str(deck_path))
    # A visible folder beside the output: PowerPoint's sandbox asks for access to a hidden one.
    with tempfile.TemporaryDirectory(prefix="render-", dir=out_dir) as work:
        src = Path(work) / f"render-{deck_path.stem[:40]}.pptx"
        src.write_bytes(deck_path.read_bytes())
        pdf = Path(work) / "deck.pdf"
        powerpoint_pdf(src, pdf)
        shown = [i for i, s in enumerate(deck.slides) if s._element.get("show") != "0"]
        written = []
        with pymupdf.open(pdf) as doc:
            order = list(range(len(deck.slides))) if doc.page_count == len(deck.slides) else shown
            if doc.page_count != len(order):
                raise RuntimeError(f"the PDF has {doc.page_count} pages for {len(deck.slides)} slides")
            for page, i in zip(doc, order):
                zoom = width / page.rect.width
                target = out_dir / f"slide-{i + 1:03}.png"
                page.get_pixmap(matrix=pymupdf.Matrix(zoom, zoom)).save(target)
                written.append(target)
    return written


def main():
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("deck", help="a .pptx file, or a run's <slug>.md deck master")
    parser.add_argument("out_dir")
    parser.add_argument("--width", type=int, default=1600)
    args = parser.parse_args()
    deck = Path(args.deck).resolve()
    if deck.suffix == ".md":
        import office_master

        run, slug = deck.parent, deck.stem
        source, original, _ = office_master.baseline(run, slug)
        data = office_master.decode(deck.read_text())
        with tempfile.TemporaryDirectory(prefix="deck-", dir=run / "subjects" / slug) as temp:
            built = Path(temp) / f"{slug}.pptx"
            office_master.build_native(source, data, original, built)
            files = pictures(built, args.out_dir, args.width)
    else:
        files = pictures(deck, args.out_dir, args.width)
    for f in files:
        print(f"wrote {f}")


if __name__ == "__main__":
    main()
