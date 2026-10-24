#!/usr/bin/env python3
"""deck_kit.py: build a deck that looks like part of an existing one.

    python3 .claude/scripts/deck_kit.py theme   <run> <slug>          # draft theme.json from the design
    python3 .claude/scripts/deck_kit.py preview <file.pptx> <out_dir> # one PNG per slide

`theme` writes `subjects/<slug>/theme.json`: the design's fonts, colours and type
sizes assigned to roles (text, muted, heading, accent, dark background...). The
roles are a first guess from how much each value is used; the agent checks them
against the page images in `subjects/<slug>/pages/` and corrects them. Every value
must exist in design.json, which `changes.py fidelity` enforces.

A build script imports the kit and draws only through it, so nothing off-palette
or off-font can reach a slide:

    import sys; sys.path.insert(0, ".claude/scripts")
    from deck_kit import Deck
    deck = Deck.load("outputs/<run>", "<slug>")            # template=<pptx> to reuse its masters
    deck.title("Life science experts", "Perfecting biology", dark=True)
    deck.bullets("The problem", ["Growth factors are costly", "They vary batch to batch"])
    deck.stats("Unit economics", [("$31,350/mg", "Activin A list price"), ("3", "suppliers")])
    deck.table("Use of funds", [["Line", "Year 1"], ["Salaries", "£120k"]])
    deck.image("Platform", "subjects/<slug>/images/p03-x29.jpeg", caption="...")
    deck.two_column("Recombinant vs synthetic", ["..."], ["..."], "Recombinant", "Synthetic")
    deck.statement("One go/no-go gate", "Reproducible receptor activation")
    deck.sources(["Roots Analysis (2025). Cell culture media market."])
    deck.save("outputs/<run>/deliverable/<slug>/vN/<name>.pptx")

Text takes the same inline markup as a change: **bold**, *italic*, [label](url).
`save` estimates overflow for every text box and prints a warning per box that
will not fit; fix the text, not the font size.

`preview` renders each slide through macOS Quick Look (or LibreOffice when it is
installed), so the agent can compare the new slides with the original pages.
"""
import json
import math
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import officedoc  # noqa: E402

ROLES = ("background", "surface", "dark", "on_dark", "muted_on_dark", "text", "muted", "heading",
         "accent", "accent2", "rule")


def _lum(hexv):
    r, g, b = (int(hexv[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def _sat(hexv):
    r, g, b = (int(hexv[i:i + 2], 16) for i in (0, 2, 4))
    return max(r, g, b) - min(r, g, b)


def draft_theme(run, slug):
    d = Path(run) / "subjects" / slug
    design = json.loads((d / "design.json").read_text())
    text_cols = [c.upper() for c in design.get("colors", {})]
    fills = [c.upper() for c in design.get("fills", design.get("shading", {}))]
    everything = list(dict.fromkeys(text_cols + fills))
    greys = [c for c in text_cols if _sat(c) < 40]
    vivid = [c for c in everything if _sat(c) >= 60]
    darks = [c for c in fills if _lum(c) < 0.15] or [c for c in everything if _lum(c) < 0.15]
    lights = [c for c in fills if _lum(c) > 0.85]
    by_use = sorted(greys, key=lambda c: -design["colors"].get(c, 0)) if greys else text_cols
    darkest_text = sorted(by_use[:4], key=_lum)
    fonts = list(design.get("fonts", {})) or list((design.get("theme", {}).get("fonts") or {}).values())
    sizes = sorted({float(s) for s in design.get("sizes_pt", {})})
    page = dict(design.get("page", {}))
    box = design.get("content_box_pt")
    if box and page and box["width"] * box["height"] < 0.85 * page["width_pt"] * page["height_pt"]:
        page = {"width_pt": box["width"], "height_pt": box["height"]}     # slide inset on a larger sheet
    styles = design.get("text_styles") or []
    biggest = max(styles, key=lambda st: st["size_pt"], default={})

    def pick(seq, i=0, fallback="000000"):
        return seq[i] if len(seq) > i else fallback

    def near(target):
        return min(sizes, key=lambda s: abs(s - target)) if sizes else target

    theme = {
        "source": slug,
        "fonts": {"heading": pick(fonts, 0, "Calibri"), "body": pick(fonts, 0, "Calibri")},
        "colors": {
            "background": pick(lights, 0, "FFFFFF"),
            "surface": pick(lights, 1, pick(lights, 0, "FFFFFF")),
            "dark": pick(darks, 0, "000000"),
            "on_dark": "FFFFFF" if "FFFFFF" in everything else pick(lights, 0, "FFFFFF"),
            "text": pick(by_use, 0),
            "muted": pick(by_use, 1, pick(by_use, 0)),
            "heading": pick(darkest_text, 0, pick(by_use, 0)),
            "accent": pick(vivid, 0, pick(by_use, 0)),
            "accent2": pick(vivid, 1, pick(vivid, 0, pick(by_use, 0))),
            "rule": pick([c for c in fills if 0.75 < _lum(c) <= 0.95], 0, pick(lights, 0, "FFFFFF")),
            "muted_on_dark": pick([c for c in text_cols if 0.3 < _lum(c) < 0.75], 0, "FFFFFF"),
        },
        "sizes": {"title": near(26), "subtitle": near(14), "body": near(11), "small": near(8),
                  "stat": near(35), "statement": near(32)},
        "slide": {"width_pt": page.get("width_pt", 960), "height_pt": page.get("height_pt", 540)},
        "margin_pt": 36,
        "style": {
            "title_bold": biggest.get("weight") == "bold" if biggest else True,
            "kicker": {"case": "upper", "color": "accent", "bold": True, "rule": False},
            "header_band": None,
            "panel": None,
            "list": "bullets",
            "number_color": "accent",
            "dividers": False,
        },
        "footer": "",
        "logo": None,
        "page_numbers": True,
        "note": "Drafted from design.json by usage. Check every role against the page images and correct it.",
    }
    out = d / "theme.json"
    out.write_text(json.dumps(theme, indent=2) + "\n")
    print(f"wrote {out}")
    for role in ROLES:
        print(f"  {role:<13} {theme['colors'][role]}")
    print(f"  fonts         {theme['fonts']}\n  sizes         {theme['sizes']}\n  slide         {theme['slide']}")
    print("Set style (kicker, header_band, panel, list, dividers, title_bold) from the page images.")
    return 0


class Deck:
    def __init__(self, theme, base, template=None):
        try:
            from pptx import Presentation
            from pptx.util import Pt
        except ImportError:
            officedoc.need("pptx")
        self.t, self.base, self.Pt = theme, Path(base), Pt
        if template:
            self.prs = Presentation(str(template))
            lst = self.prs.slides._sldIdLst
            for s in list(lst):
                self.prs.part.drop_rel(s.rId)
                lst.remove(s)
        else:
            self.prs = Presentation()
            self.prs.slide_width = Pt(theme["slide"]["width_pt"])
            self.prs.slide_height = Pt(theme["slide"]["height_pt"])
            self._retheme()
        self.blank = min(self.prs.slide_layouts, key=lambda lay: len(lay.placeholders))
        self.W, self.H = theme["slide"]["width_pt"], theme["slide"]["height_pt"]
        self.M = theme.get("margin_pt", 36)
        self.warnings = []
        self.count = 0

    @classmethod
    def load(cls, run, slug, template=None):
        base = Path(run) / "subjects" / slug
        theme = json.loads((base / "theme.json").read_text())
        return cls(theme, base, template)

    # -- theme part: the defaults Office falls back on must be the design's too

    def _retheme(self):
        from lxml import etree
        A = officedoc.A
        part = None
        for rel in self.prs.slide_masters[0].part.rels.values():
            if rel.reltype.endswith("/theme"):
                part = rel.target_part
        if part is None:
            return
        root = etree.fromstring(part.blob)
        c = self.t["colors"]
        slots = {"dk1": c["text"], "lt1": c["background"], "dk2": c["dark"], "lt2": c["surface"],
                 "accent1": c["accent"], "accent2": c["accent2"], "accent3": c["muted"],
                 "accent4": c["heading"], "accent5": c["accent"], "accent6": c["accent2"],
                 "hlink": c["accent"], "folHlink": c["accent2"]}
        scheme = root.find(f".//{{{A}}}clrScheme")
        for slot in list(scheme):
            name = slot.tag.rsplit("}", 1)[-1]
            if name in slots:
                for child in list(slot):
                    slot.remove(child)
                etree.SubElement(slot, f"{{{A}}}srgbClr", val=slots[name].upper())
        for which, role in (("majorFont", "heading"), ("minorFont", "body")):
            latin = root.find(f".//{{{A}}}{which}/{{{A}}}latin")
            if latin is not None:
                latin.set("typeface", self.t["fonts"][role])
        part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    # -- drawing

    def _rgb(self, role):
        from pptx.dml.color import RGBColor
        return RGBColor.from_string(self.t["colors"].get(role, role).upper())

    def _new(self, dark=False):
        slide = self.prs.slides.add_slide(self.blank)
        for ph in list(slide.placeholders):
            ph._element.getparent().remove(ph._element)
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = self._rgb("dark" if dark else "background")
        self.count += 1
        return slide

    def _box(self, slide, x, y, w, h, fill=None):
        from pptx.enum.shapes import MSO_SHAPE
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, self.Pt(x), self.Pt(y), self.Pt(w), self.Pt(h))
        shp.fill.solid()
        shp.fill.fore_color.rgb = self._rgb(fill or "surface")
        shp.line.fill.background()
        shp.shadow.inherit = False
        return shp

    def text(self, slide, x, y, w, h, paras, size="body", color="text", bold=False,
             font="body", align=None, bullet=False, spacing=1.15):
        """A text box. `paras` is a string or a list; each item is one paragraph."""
        from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
        if isinstance(paras, str):
            paras = [paras]
        pt = self.t["sizes"].get(size, size) if isinstance(size, str) else size
        box = slide.shapes.add_textbox(self.Pt(x), self.Pt(y), self.Pt(w), self.Pt(h))
        tf = box.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        for side in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
            setattr(tf, side, 0)
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.line_spacing = spacing
            if align:
                p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(align, PP_ALIGN.LEFT)
            if bullet and i:
                p.space_before = self.Pt(pt * 0.5)
            prefix = "• " if bullet else ""
            for seg in officedoc.segments(prefix + str(para)):
                if not seg["text"]:
                    continue
                r = p.add_run()
                r.text = seg["text"]
                f = r.font
                f.name = self.t["fonts"][font]
                f.size = self.Pt(pt)
                f.bold = bool(bold or seg.get("bold"))
                f.italic = bool(seg.get("italic"))
                f.color.rgb = self._rgb("accent" if seg.get("link") else color)
                if seg.get("link"):
                    r.hyperlink.address = seg["link"]
        self._fit(paras, w, h, pt, spacing, f"slide {self.count}")
        return box

    def _fit(self, paras, w, h, pt, spacing, where):
        per_line = max(1, int(w / (pt * 0.52)))
        lines = sum(max(1, math.ceil(len(officedoc.plain(str(p))) / per_line)) for p in paras)
        need = lines * pt * spacing * 1.2
        if need > h * 1.02:
            first = officedoc.plain(str(paras[0]))[:40]
            self.warnings.append(f"{where}: text starting {first!r} needs ~{need:.0f}pt, box is {h:.0f}pt")

    def _chrome(self, slide, dark=False):
        col = "muted_on_dark" if dark else "muted"
        if self.t.get("footer"):
            self.text(slide, self.M, self.H - self.M * 0.75, self.W / 2, 14, self.t["footer"], "small", col)
        if self.t.get("page_numbers"):
            self.text(slide, self.W - self.M - 40, self.H - self.M * 0.75, 40, 14, str(self.count),
                      "small", col, align="right")
        logo = self.t.get("logo")
        if logo and (self.base / logo).is_file() and not dark:
            slide.shapes.add_picture(str(self.base / logo), self.Pt(self.W - self.M - 60),
                                     self.Pt(self.M * 0.5), height=self.Pt(18))

    def _style(self):
        base = {"title_bold": True, "kicker": {"case": "upper", "color": "accent", "bold": True, "rule": False},
                "header_band": None, "panel": None, "list": "bullets", "number_color": "accent",
                "dividers": False}
        base.update(self.t.get("style") or {})
        return base

    def _rule(self, slide, x, y, w, role):
        self._box(slide, x, y, w, 1.25, role)

    def _heading(self, slide, title, kicker=None, dark=False, subtitle=None):
        st = self._style()
        k = st["kicker"]
        title_h = self.t["sizes"]["title"] * 2.6
        band_h = self.M + (28 if kicker else 0) + title_h + (22 if subtitle else 0) + 6
        if st["header_band"] and not dark:
            self._box(slide, 0, 0, self.W, band_h, st["header_band"])
        y = self.M
        if kicker:
            label = kicker.upper() if k.get("case") == "upper" else kicker[:1].upper() + kicker[1:]
            size = "small" if k.get("case") == "upper" else "subtitle"
            self.text(slide, self.M, y, self.W - 2 * self.M, self.t["sizes"][size] * 1.5, label, size,
                      "on_dark" if dark else k.get("color", "accent"), bold=k.get("bold", False))
            y += self.t["sizes"][size] * 1.7
            if k.get("rule"):
                self._rule(slide, self.M, y, 28, "muted")
                y += 8
        self.text(slide, self.M, y, self.W - 2 * self.M, title_h, title, "title",
                  "on_dark" if dark else "heading", bold=st["title_bold"], font="heading")
        y += title_h
        if subtitle:
            self.text(slide, self.M, y, self.W - 2 * self.M, 18, subtitle, "small", "muted")
            y += 22
        return max(y, band_h if st["header_band"] else y) + 12

    def _body(self, slide, y):
        """The area under the heading; with a panel style it is filled and its text turns light."""
        panel = self._style()["panel"]
        if panel:
            self._box(slide, 0, y - 12, self.W, self.H - y + 12, panel)
        return bool(panel)

    def _list(self, slide, x, y, w, h, items, on_dark=False):
        st = self._style()
        main, sub = ("on_dark", "muted_on_dark") if on_dark else ("text", "muted")
        if st["list"] == "bullets":
            return self.text(slide, x, y, w, h, items, "body", main, bullet=True)
        if st["list"] == "plain":
            return self.text(slide, x, y, w, h, items, "body", main)
        body = self.t["sizes"]["body"]
        small = self.t["sizes"]["small"]
        step = h / max(1, len(items))
        for i, item in enumerate(items):
            head, _, rest = str(item).partition("\n")
            yy = y + i * step
            self.text(slide, x, yy + 2, 22, body * 1.4, f"{i + 1:02d}", "small", st["number_color"], bold=True)
            self.text(slide, x + 24, yy, w - 24, body * 1.6, head, body + 2, main)
            if rest:
                self.text(slide, x + 24, yy + body * 1.8, w - 24, step - body * 2.2, rest, small + 1, sub)
            if st["dividers"]:
                self._rule(slide, x, yy + step - 6, w, "muted_on_dark" if on_dark else "rule")

    # -- slide types

    def title(self, title, subtitle="", dark=True, image=None):
        s = self._new(dark)
        if image and (self.base / image).is_file():
            s.shapes.add_picture(str(self.base / image), self.Pt(self.W * 0.55), 0, height=self.Pt(self.H))
        width = self.W * (0.48 if image else 0.8)
        size = self.t["sizes"]["statement"]
        self.text(s, self.M, self.H * 0.32, width, size * 3.6, title, size,
                  "on_dark" if dark else "heading", bold=True, font="heading")
        if subtitle:
            self.text(s, self.M, self.H * 0.32 + size * 3.8, width, 60, subtitle, "subtitle",
                      "on_dark" if dark else "muted")
        return s

    def statement(self, text, support="", dark=True):
        s = self._new(dark)
        size = self.t["sizes"]["statement"]
        self.text(s, self.M * 1.5, self.H * 0.3, self.W - 3 * self.M, size * 3.8, text, size,
                  "on_dark" if dark else "heading", bold=True, font="heading")
        if support:
            self.text(s, self.M * 1.5, self.H * 0.3 + size * 4, self.W - 3 * self.M, 80, support,
                      "subtitle", "on_dark" if dark else "muted")
        self._chrome(s, dark)
        return s

    def bullets(self, title, items, kicker=None, notes=None, subtitle=None, image=None, closing=None):
        """A list slide. An item "Head\nDetail" shows a heading line and a quieter detail line.
        `image` puts a picture on the right half of the body; `closing` is one line under the list."""
        s = self._new()
        y = self._heading(s, title, kicker, subtitle=subtitle)
        dark = self._body(s, y)
        width = (self.W - 2 * self.M) * (0.5 if image else 1)
        bottom = self.H - self.M * 1.5 - (40 if closing else 0)
        self._list(s, self.M + (self.M if dark else 0), y + 8, width - (self.M if dark else 0), bottom - y - 8,
                   items, dark)
        if closing:
            self.text(s, self.M + (self.M if dark else 0), bottom + 6, width, 30, closing, "subtitle",
                      "on_dark" if dark else "heading")
        if image:
            p = self.base / image if not Path(image).is_file() else Path(image)
            x = self.M + width + 24
            pic = s.shapes.add_picture(str(p), self.Pt(x), self.Pt(y + 8), width=self.Pt(self.W - x - self.M))
            if pic.height > self.Pt(self.H - y - self.M * 1.5):
                pic.height = self.Pt(self.H - y - self.M * 1.5)
        self._notes(s, notes)
        self._chrome(s, dark)
        return s

    def two_column(self, title, left, right, left_head="", right_head="", kicker=None, notes=None):
        s = self._new()
        y = self._heading(s, title, kicker)
        dark = self._body(s, y)
        col = (self.W - 2 * self.M - 24) / 2
        for i, (head, items) in enumerate(((left_head, left), (right_head, right))):
            x = self.M + i * (col + 24)
            if not dark:
                self._box(s, x, y, col, self.H - y - self.M * 1.5, "surface")
            yy = y + 14
            if head:
                self.text(s, x + 14, yy, col - 28, 22, head, "subtitle", "on_dark" if dark else "heading",
                          font="heading", bold=self._style()["title_bold"])
                yy += 30
            self._list(s, x + 14, yy, col - 28, self.H - yy - self.M * 1.5 - 14, items, dark)
        self._notes(s, notes)
        self._chrome(s, dark)
        return s

    def stats(self, title, figures, kicker=None, notes=None):
        s = self._new()
        y = self._heading(s, title, kicker)
        n = max(1, len(figures))
        gap = 18
        col = (self.W - 2 * self.M - gap * (n - 1)) / n
        stat = self.t["sizes"]["stat"]
        for i, (value, label) in enumerate(figures):
            x = self.M + i * (col + gap)
            self._box(s, x, y, col, stat * 2 + 80, "surface")
            self.text(s, x + 14, y + 16, col - 28, stat * 1.4, value, stat, "accent", bold=True, font="heading")
            self.text(s, x + 14, y + 22 + stat * 1.4, col - 28, 60, label, "body", "muted")
        self._notes(s, notes)
        self._chrome(s)
        return s

    def table(self, title, rows, kicker=None, col_widths=None, notes=None):
        s = self._new()
        y = self._heading(s, title, kicker)
        n_rows, n_cols = len(rows), max(len(r) for r in rows)
        width = self.W - 2 * self.M
        row_h = self.t["sizes"]["body"] * 2.2
        shape = s.shapes.add_table(n_rows, n_cols, self.Pt(self.M), self.Pt(y), self.Pt(width),
                                   self.Pt(row_h * n_rows))
        tbl = shape.table
        tbl.first_row = True
        for j in range(n_cols):
            frac = (col_widths[j] if col_widths else 1 / n_cols)
            tbl.columns[j].width = self.Pt(width * frac)
        for i, row in enumerate(rows):
            for j in range(n_cols):
                cell = tbl.cell(i, j)
                cell.fill.solid()
                cell.fill.fore_color.rgb = self._rgb("dark" if i == 0 else ("surface" if i % 2 else "background"))
                tf = cell.text_frame
                tf.paragraphs[0].text = ""
                r = tf.paragraphs[0].add_run()
                r.text = officedoc.plain(str(row[j])) if j < len(row) else ""
                r.font.name = self.t["fonts"]["body"]
                r.font.size = self.Pt(self.t["sizes"]["body"])
                r.font.bold = i == 0
                r.font.color.rgb = self._rgb("on_dark" if i == 0 else "text")
        if y + row_h * n_rows > self.H - self.M:
            self.warnings.append(f"slide {self.count}: table of {n_rows} rows runs past the slide")
        self._notes(s, notes)
        self._chrome(s)
        return s

    def image(self, title, path, caption="", kicker=None, notes=None):
        s = self._new()
        y = self._heading(s, title, kicker)
        p = Path(path)
        if not p.is_absolute() and not p.is_file():
            p = self.base / path
        room = self.H - y - self.M * 1.5 - (24 if caption else 0)
        pic = s.shapes.add_picture(str(p), self.Pt(self.M), self.Pt(y), height=self.Pt(room))
        if pic.width > self.Pt(self.W - 2 * self.M):
            ratio = (self.W - 2 * self.M) / (pic.width / 12700)
            pic.width = self.Pt(self.W - 2 * self.M)
            pic.height = int(pic.height * ratio)
        if caption:
            self.text(s, self.M, self.H - self.M * 1.5 - 18, self.W - 2 * self.M, 18, caption, "small", "muted")
        self._notes(s, notes)
        self._chrome(s)
        return s

    def sources(self, entries, title="Sources"):
        s = self._new()
        y = self._heading(s, title)
        half = math.ceil(len(entries) / 2) if len(entries) > 8 else len(entries)
        cols = [entries[:half], entries[half:]] if half < len(entries) else [entries]
        width = (self.W - 2 * self.M - 24) / len(cols)
        for i, chunk in enumerate(cols):
            self.text(s, self.M + i * (width + 24), y, width, self.H - y - self.M * 1.5, chunk, "small", "muted")
        self._chrome(s)
        return s

    def _notes(self, slide, notes):
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    def save(self, dest):
        dest = Path(dest)
        dest.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(dest))
        for wmsg in self.warnings:
            print(f"warn    overflow {wmsg}")
        print(f"wrote {dest} · {self.count} slide(s) · {len(self.warnings)} overflow warning(s)")
        return dest


def preview(pptx, out_dir):
    """One PNG per slide: each slide alone in a copy, rendered by Quick Look or LibreOffice."""
    from pptx import Presentation
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(out), str(pptx)],
                       capture_output=True, timeout=300)
        pdf = out / (Path(pptx).stem + ".pdf")
        if pdf.is_file():
            doc = officedoc._pdf_open(pdf)
            for i, page in enumerate(doc, 1):
                page.get_pixmap(dpi=80).save(str(out / f"slide-{i:02d}.png"))
            print(f"rendered {doc.page_count} slide(s) into {out}")
            return 0
    if not shutil.which("qlmanage"):
        sys.exit("ERROR no renderer: install LibreOffice (brew install --cask libreoffice)")
    n = len(Presentation(str(pptx)).slides)
    with tempfile.TemporaryDirectory() as tmp:
        for i in range(n):
            prs = Presentation(str(pptx))
            lst = prs.slides._sldIdLst
            for j, s in enumerate(list(lst)):
                if j != i:
                    prs.part.drop_rel(s.rId)
                    lst.remove(s)
            one = Path(tmp) / f"slide-{i + 1:02d}.pptx"
            prs.save(str(one))
            subprocess.run(["qlmanage", "-t", "-s", "1200", "-o", tmp, str(one)],
                           capture_output=True, timeout=60)
            png = Path(tmp) / f"{one.name}.png"
            if png.is_file():
                shutil.move(str(png), str(out / f"slide-{i + 1:02d}.png"))
    print(f"rendered {n} slide(s) into {out} (Quick Look: close to PowerPoint, not identical)")
    return 0


def main():
    args = [x for x in sys.argv[1:] if not x.startswith("--")]
    if len(args) == 3 and args[0] == "theme":
        return draft_theme(args[1], args[2])
    if len(args) == 3 and args[0] == "preview":
        return preview(args[1], args[2])
    sys.exit(__doc__)


if __name__ == "__main__":
    sys.exit(main())
