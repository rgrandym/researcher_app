#!/usr/bin/env python3
"""apply_changes.py: apply a subject's accepted changes to a copy of the original file.

    python3 .claude/scripts/apply_changes.py <run> <slug> [--track] [--author NAME] [--partial]

Reads `subjects/<slug>/changes.json` with the author's decisions, applies only the
accepted changes (an author-edited text replaces the proposal), and writes
`deliverable/<slug>/vN/<original filename>` with `applied.json` beside it. The file in
`external_data/` is never touched; every version starts again from it.

The edit happens inside the original's own structure, so its formatting survives:

- Word: text is replaced inside the existing runs, new paragraphs copy the pPr and
  run formatting of a donor paragraph (`like`, or the target), new tables copy an
  existing table, links copy the document's link formatting, and `^[...]` adds a
  footnote shaped like the document's first one. `--track` writes every change as a
  Word tracked change, so the author can review it in Word.
- PowerPoint: text is replaced in place; new slides clone an existing slide (`like`)
  with its layout, shapes, pictures and background, then `fill` sets shape text.
- Excel: cells are set in the sheet XML directly, keeping styles, charts and every
  other part, and the workbook is marked to recalculate when it opens.
- Markdown: blocks are replaced, inserted or deleted.

If any accepted change cannot be applied, nothing is written and each refusal is
listed (exit 1); `--partial` writes the rest and records the refusals.
"""
import json
import re
import sys
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import changes as plans  # noqa: E402
import officedoc  # noqa: E402
import subject as subjects  # noqa: E402
from officedoc import A, P, R, S, W, XML_SPACE, a, segments, w  # noqa: E402

RPR_ORDER = ["rStyle", "rFonts", "b", "bCs", "i", "iCs", "caps", "smallCaps", "strike", "dstrike",
             "outline", "shadow", "emboss", "imprint", "noProof", "snapToGrid", "vanish", "webHidden",
             "color", "spacing", "w", "kern", "position", "sz", "szCs", "highlight", "u", "effect",
             "bdr", "shd", "fitText", "vertAlign", "rtl", "cs", "em", "lang", "eastAsianLayout",
             "specVanish", "oMath"]
TRACK_MARKS = ("ins", "del", "rPrChange", "moveFrom", "moveTo")
HYPERLINK = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink"


class Refused(Exception):
    pass


class TextEditor:
    """What Word and PowerPoint share: find text in a paragraph, cut it on run boundaries."""

    def atoms(self, p):
        raise NotImplementedError

    def split(self, atom, k):
        raise NotImplementedError

    def locate(self, p, before, target):
        if not before:
            raise Refused("replace needs `before`: the exact original text it changes")
        if officedoc.TOKEN.search(before):
            raise Refused("`before` contains a footnote marker or object; replace the text around it")
        text = "".join(x.text for x in self.atoms(p))
        n = text.count(before)
        if n == 0:
            hint = " (it matches once whitespace is ignored: copy it exactly)" if (
                " ".join(before.split()) in " ".join(text.split())) else ""
            raise Refused(f"`before` is not in {target}{hint}")
        if n > 1:
            raise Refused(f"`before` occurs {n} times in {target}; quote more of the sentence")
        s = text.index(before)
        return s, s + len(before)

    def cut(self, p, k):
        pos = 0
        for x in self.atoms(p):
            n = len(x.text)
            if pos < k < pos + n:
                if x.fixed:
                    raise Refused("the change would split a footnote marker, field or object")
                self.split(x, k - pos)
                return
            pos += n

    def span(self, p, s, e):
        """Atoms wholly inside [s, e) once the paragraph is cut at s and e."""
        self.cut(p, s)
        self.cut(p, e)
        out, pos = [], 0
        for x in self.atoms(p):
            n = len(x.text)
            if n and pos >= s and pos + n <= e:
                if x.fixed and x.text not in ("\n", "\t"):
                    raise Refused("the replaced text includes a footnote marker, field or object")
                out.append(x)
            pos += n
        return out


# ------------------------------------------------------------------ docx

class DocxEditor(TextEditor):
    def __init__(self, src, track=False, author="Revision"):
        try:
            import docx
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
        except ImportError:
            officedoc.need("docx")
        self.E, self.qn = OxmlElement, qn
        self.doc = docx.Document(str(src))
        self.index = officedoc.docx_index(self.doc)
        self.fn, self.notes = officedoc.docx_footnotes(self.doc)
        self.track, self.author = track, author
        self.rev = 9000
        self.when = datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")
        self.cursor, self.deleted = {}, set()
        self.notes_made = 0
        self.notes_out = []
        body = self.doc.element.body
        self.names = officedoc.docx_style_names(self.doc)
        self.body_color = self._body_color(body)
        self.link_rpr = self._link_rpr(body)
        sect = body.find(w("sectPr"))
        size, mar = (sect.find(w("pgSz")), sect.find(w("pgMar"))) if sect is not None else (None, None)
        self.text_width = 9360
        if size is not None and mar is not None:
            self.text_width = int(float(size.get(w("w"), 12240)) - float(mar.get(w("left"), 1440))
                                  - float(mar.get(w("right"), 1440)))
        self.comments = None
        for rel in self.doc.part.rels.values():
            if rel.reltype.endswith("/comments") and not rel.is_external:
                self.comments = officedoc.XmlPart.of(rel.target_part)
        self.late = []                      # the clean-up runs after every other change
        self.fnref_rpr = None
        for ref in body.iter(w("footnoteReference")):
            rpr = ref.getparent().find(w("rPr"))
            if rpr is not None:
                self.fnref_rpr = self.clean(rpr)
                break

    # -- elements

    def clean(self, rpr):
        if rpr is None:
            return None
        rpr = deepcopy(rpr)
        for mark in TRACK_MARKS:
            for el in rpr.findall(w(mark)):
                rpr.remove(el)
        return rpr

    def _body_color(self, body):
        """The explicit colour of ordinary body text, or None when it carries none."""
        counts = {}
        for r in body.iter(w("r")):
            if officedoc.in_link(r):
                continue
            n = sum(len(t.text or "") for t in r.findall(w("t")))
            struck, color, _ = officedoc.run_marks(r)
            if n and not struck:
                counts[color] = counts.get(color, 0) + n
        if not counts:
            return None
        best = max(counts, key=counts.get)
        return best if best in ("000000", "222222", "333333") or best is None else None

    def _link_rpr(self, body):
        """The formatting most of the document's links carry, free of working marks."""
        found = {}
        for h in body.iter(w("hyperlink")):
            if not h.get(f"{{{R}}}id"):
                continue
            for r in h.findall(w("r")):
                struck, color, _ = officedoc.run_marks(r)
                if not struck and r.find(w("rPr")) is not None:
                    found.setdefault(color, [0, r.find(w("rPr"))])[0] += 1
        if not found:
            return None
        rpr = self.clean(max(found.values(), key=lambda v: v[0])[1])
        for tag in ("strike", "dstrike", "highlight"):
            for el in rpr.findall(w(tag)):
                rpr.remove(el)
        return rpr

    def plain(self, rpr):
        """A donor's formatting for new text: its font, size and weight, without the
        working marks of a draft (strikethrough, draft colour, highlight)."""
        rpr = self.clean(rpr)
        if rpr is None:
            return None
        for tag in ("strike", "dstrike", "highlight", "color", "shd"):
            for el in rpr.findall(w(tag)):
                rpr.remove(el)
        if self.body_color:
            col = self.E("w:color")
            col.set(w("val"), self.body_color)
            rank = RPR_ORDER.index("color")
            for child in rpr:
                ctag = child.tag.rsplit("}", 1)[-1]
                if ctag in RPR_ORDER and RPR_ORDER.index(ctag) > rank:
                    child.addprevious(col)
                    break
            else:
                rpr.append(col)
        return rpr

    def rpr_set(self, rpr, tag):
        el = rpr.find(w(tag))
        if el is not None:
            el.attrib.pop(w("val"), None)
            return
        new = self.E(f"w:{tag}")
        rank = RPR_ORDER.index(tag)
        for child in rpr:
            ctag = child.tag.rsplit("}", 1)[-1]
            if ctag in RPR_ORDER and RPR_ORDER.index(ctag) > rank:
                child.addprevious(new)
                return
        rpr.append(new)

    def revision(self, tag):
        self.rev += 1
        el = self.E(f"w:{tag}")
        el.set(w("id"), str(self.rev))
        el.set(w("author"), self.author)
        el.set(w("date"), self.when)
        return el

    def text_run(self, text, rpr, bold=False, italic=False, strike=False):
        r = self.E("w:r")
        if rpr is not None:
            r.append(deepcopy(rpr))
        if bold or italic or strike:
            rp = r.find(w("rPr"))
            if rp is None:
                rp = self.E("w:rPr")
                r.insert(0, rp)
            for flag, tag in ((bold, "b"), (italic, "i"), (strike, "strike")):
                if flag:
                    self.rpr_set(rp, tag)
        for i, line in enumerate(text.split("\n")):
            if i:
                r.append(self.E("w:br"))
            for j, chunk in enumerate(line.split("\t")):
                if j:
                    r.append(self.E("w:tab"))
                if chunk:
                    t = self.E("w:t")
                    t.text = chunk
                    t.set(XML_SPACE, "preserve")
                    r.append(t)
        return r

    def hyperlink(self, text, url, donor, part):
        rid = part.relate_to(url, HYPERLINK, is_external=True)
        h = self.E("w:hyperlink")
        h.set(f"{{{R}}}id", rid)
        h.set(w("history"), "1")
        h.append(self.text_run(text, self.link_rpr if self.link_rpr is not None else donor))
        return h

    def footnote_ref(self, text):
        if self.fn is None or not self.notes:
            raise Refused("this document has no footnotes to copy; cite in its own style instead")
        tpl = self.notes[0]
        new_id = max(int(f.get(w("id"))) for f in self.fn.root.findall(w("footnote"))) + 1
        note = deepcopy(tpl)
        note.set(w("id"), str(new_id))
        paras = note.findall(w("p"))
        for extra in paras[1:]:
            note.remove(extra)
        p = paras[0]
        text_rpr = None
        for r in [r for r in p.iter(w("r"))]:
            t = r.find(w("t"))
            if r.find(w("footnoteRef")) is not None:
                continue
            if text_rpr is None and t is not None and (t.text or "").strip():
                text_rpr = self.clean(r.find(w("rPr")))
            parent = r.getparent()
            parent.remove(r)
            if parent.tag == w("hyperlink") and parent.find(w("r")) is None:
                parent.getparent().remove(parent)
        for el in self.make(segments(" " + text), text_rpr, self.fn.part, in_note=True):
            p.append(el)
        self.fn.root.append(note)
        self.notes_made += 1
        r = self.E("w:r")
        if self.fnref_rpr is not None:
            r.append(deepcopy(self.fnref_rpr))
        else:
            rp = self.E("w:rPr")
            va = self.E("w:vertAlign")
            va.set(w("val"), "superscript")
            rp.append(va)
            r.append(rp)
        ref = self.E("w:footnoteReference")
        ref.set(w("id"), str(new_id))
        r.append(ref)
        return r

    def existing_ref(self, fid):
        """A reference to a footnote the document already has: `[^3]` in Markdown text."""
        if self.fn is None or not any(f.get(w("id")) == fid for f in self.notes):
            raise Refused(f"[^{fid}] refers to a footnote the document does not have")
        r = self.E("w:r")
        if self.fnref_rpr is not None:
            r.append(deepcopy(self.fnref_rpr))
        ref = self.E("w:footnoteReference")
        ref.set(w("id"), str(fid))
        r.append(ref)
        return r

    def make(self, segs, donor, part, inside_link=False, in_note=False):
        out = []
        for seg in segs:
            if seg.get("footnote") is not None:
                if in_note:
                    raise Refused("a footnote cannot contain another footnote")
                out.append(self.footnote_ref(seg["footnote"]))
            elif seg.get("ref") is not None:
                out.append(self.existing_ref(seg["ref"]))
            elif not seg["text"]:
                continue
            elif seg.get("link"):
                if inside_link:
                    raise Refused("a link cannot sit inside another link")
                out.append(self.hyperlink(seg["text"], seg["link"], donor, part))
            else:
                out.append(self.text_run(seg["text"], donor, seg.get("bold"), seg.get("italic"), seg.get("strike")))
        if self.track:
            out = [self.wrap_ins(el) for el in out]
        return out

    def wrap_ins(self, el):
        if el.tag == w("hyperlink"):
            for r in list(el):
                ins = self.revision("ins")
                r.addprevious(ins)
                ins.append(r)
            return el
        ins = self.revision("ins")
        ins.append(el)
        return ins

    def remove_runs(self, runs):
        for r in runs:
            parent = r.getparent()
            if self.track:
                for t in r.findall(w("t")):
                    t.tag = w("delText")
                for t in r.findall(w("instrText")):
                    t.tag = w("delInstrText")
                d = self.revision("del")
                r.addprevious(d)
                d.append(r)
            else:
                parent.remove(r)
                if parent.tag == w("hyperlink") and parent.find(f".//{w('r')}") is None:
                    parent.getparent().remove(parent)

    def mark_paragraph(self, p, tag):
        ppr = p.find(w("pPr"))
        if ppr is None:
            ppr = self.E("w:pPr")
            p.insert(0, ppr)
        rpr = ppr.find(w("rPr"))
        if rpr is None:
            rpr = self.E("w:rPr")
            change = ppr.find(w("pPrChange"))
            if change is not None:
                change.addprevious(rpr)
            else:
                ppr.append(rpr)
        rpr.insert(0, self.revision(tag))

    # -- text atoms

    def normalize(self, p):
        """One content child per run, so any atom can be cut or removed alone."""
        for r in officedoc.docx_runs(p):
            content = [c for c in r if c.tag != w("rPr")]
            prev = r
            for child in content[1:]:
                nr = self.E("w:r")
                if r.find(w("rPr")) is not None:
                    nr.append(deepcopy(r.find(w("rPr"))))
                nr.append(child)
                prev.addnext(nr)
                prev = nr

    def atoms(self, p):
        return officedoc.docx_atoms(p)

    def split(self, atom, k):
        clone = deepcopy(atom.run)
        t = atom.child
        text = t.text or ""
        t.text = text[:k]
        t.set(XML_SPACE, "preserve")
        ct = clone.find(w("t"))
        ct.text = text[k:]
        ct.set(XML_SPACE, "preserve")
        atom.run.addnext(clone)

    def donor(self, p):
        """Run formatting for new text in or like paragraph `p`: an unmarked run first."""
        candidates = []
        for r in officedoc.docx_runs(p):
            if r.getparent().tag == w("hyperlink"):
                continue
            t = r.find(w("t"))
            rpr = r.find(w("rPr"))
            if t is None or not (t.text or "").strip():
                continue
            if rpr is not None and (rpr.find(w("vertAlign")) is not None or rpr.find(w("rStyle")) is not None
                                    and "ootnote" in (rpr.find(w("rStyle")).get(w("val")) or "")):
                continue
            struck, color, hl = officedoc.run_marks(r)
            bold = rpr is not None and any(
                rpr.find(w(tag)) is not None and rpr.find(w(tag)).get(w("val")) not in ("0", "false")
                for tag in ("b", "i"))
            candidates.append((bool(struck or color or hl), bold, -len(t.text or ""), rpr))
        if candidates:
            # The paragraph's ordinary text: unmarked, not a bold or italic label, the longest run.
            return self.plain(sorted(candidates, key=lambda c: c[:3])[0][3])
        ppr = p.find(w("pPr"))
        return self.plain(ppr.find(w("rPr")) if ppr is not None else None)

    def part_of(self, el):
        if self.fn is not None and el.getroottree().getroot() is self.fn.root:
            return self.fn.part
        return self.doc.part

    # -- operations

    def el(self, target):
        if target in self.deleted:
            raise Refused(f"{target} was deleted by an earlier change")
        el = self.index.get(target)
        if el is None:
            raise Refused(f"{target} is not a part of the original")
        return el

    def para(self, target):
        el = self.el(target)
        if el.tag != w("p"):
            raise Refused(f"{target} is a table; target one of its cell paragraphs (T1.R1.C1.P1)")
        return el

    def apply(self, ch):
        op = ch["op"]
        if op == "replace":
            return self.replace(ch)
        if op in ("insert_after", "insert_before"):
            return self.insert(ch, op == "insert_after")
        if op == "delete":
            return self.delete(ch)
        if op == "insert_table":
            return self.insert_table(ch)
        if op == "rewrite_section":
            return self.rewrite_section(ch)
        if op == "rewrite":
            return self.rewrite(ch)
        if op == "clean":
            if ch.get("target") != "document":
                raise Refused("clean applies to the whole document: target \"document\"")
            self.late.append(ch)
            return None
        raise Refused(f"op {op} is not available for Word documents")

    def replace(self, ch):
        p = self.para(ch["target"])
        self.normalize(p)
        s, e = self.locate(p, ch["before"], ch["target"])
        span = self.span(p, s, e)
        first = span[0].run
        parent = first.getparent()
        in_link = parent.tag == w("hyperlink")
        if in_link and all(x.run.getparent() is parent for x in span):
            anchor, donor, inside = first, self.link_rpr if self.link_rpr is not None else self.plain(first.find(w("rPr"))), True
        elif in_link:
            anchor, donor, inside = parent, self.donor(p), False
        else:
            rpr = first.find(w("rPr"))
            plain_run = rpr is None or rpr.find(w("vertAlign")) is None
            anchor, donor, inside = first, self.plain(rpr) if plain_run else self.donor(p), False
        for el in self.make(segments(ch["after"] or ""), donor, self.part_of(p), inside):
            anchor.addprevious(el)
        self.remove_runs([x.run for x in span])

    def new_paragraph(self, donor_p, text, part):
        p = self.E("w:p")
        ppr = donor_p.find(w("pPr"))
        if ppr is not None:
            ppr = deepcopy(ppr)
            for tag in ("sectPr", "pPrChange"):
                for el in ppr.findall(w(tag)):
                    ppr.remove(el)
            mark = ppr.find(w("rPr"))
            if mark is not None:
                for tag in TRACK_MARKS:
                    for el in mark.findall(w(tag)):
                        mark.remove(el)
            p.append(ppr)
        for el in self.make(segments(text), self.donor(donor_p), part):
            p.append(el)
        if self.track:
            self.mark_paragraph(p, "ins")
        return p

    def default_donor(self, el):
        if el.tag == w("p"):
            return el
        prev = el.getprevious()
        while prev is not None and prev.tag != w("p"):
            prev = prev.getprevious()
        if prev is None:
            raise Refused("give `like`: the paragraph whose formatting the new text copies")
        return prev

    def place(self, key, el, new, after):
        if after:
            ref = self.cursor.get(key, el)
            ref.addnext(new)
            self.cursor[key] = new
        else:
            el.addprevious(new)

    def insert(self, ch, after):
        el = self.el(ch["target"])
        donor = self.el(ch["like"]) if ch.get("like") else self.default_donor(el)
        if donor.tag != w("p"):
            raise Refused("`like` must be a paragraph")
        made = self.build(ch.get("after") or "", donor, self.part_of(el), ch.get("like_table"))
        if not made:
            raise Refused("insert needs `after`: the new text")
        for new in made:
            self.place(ch["target"], el, new, after)

    def delete(self, ch):
        target = ch["target"]
        el = self.el(target)
        if el.tag == w("tbl"):
            if self.track:
                raise Refused("tracked mode cannot delete a whole table; delete its cells' text instead")
            el.getparent().remove(el)
            self.deleted.update(k for k in self.index if k == target or k.startswith(target + "."))
            return
        if ch.get("before") and plans.norm(ch["before"]) != plans.norm(officedoc.docx_text(el)):
            raise Refused(f"`before` does not match the text of {target}")
        self.normalize(el)
        runs = list(dict.fromkeys(x.run for x in self.atoms(el)))
        parent = el.getparent()
        only = parent.tag == w("tc") and len(parent.findall(w("p"))) == 1
        if self.track:
            self.remove_runs(runs)
            if not only:
                self.mark_paragraph(el, "del")
        elif only:
            self.remove_runs(runs)
        else:
            parent.remove(el)
        self.deleted.add(target)

    def set_text(self, p, text, part):
        donor = self.donor(p)
        for child in list(p):
            if child.tag != w("pPr"):
                p.remove(child)
        for el in self.make(segments(text), donor, part):
            p.append(el)

    def rewrite(self, ch):
        """A paragraph's whole text, from Markdown, in the paragraph's own formatting. Its
        heading, list or footnote form stays; the Markdown prefix only says which it is."""
        p = self.para(ch["target"])
        text = (ch.get("after") or "").strip()
        text = re.sub(r"^(#{1,6}|[-*•]|\d+[.)])\s+", "", text)
        text = re.sub(r"^\[\^\d+\]:\s*", "", text).replace("\n", " ")
        part = self.part_of(p)
        in_note = part is not self.doc.part
        donor = self.donor(p)
        self.normalize(p)
        mark = [r for r in officedoc.docx_runs(p) if r.find(w("footnoteRef")) is not None]
        old = [r for r in dict.fromkeys(x.run for x in self.atoms(p)) if r not in mark]
        if self.track:
            self.remove_runs(old)
        else:
            for child in list(p):
                if child.tag != w("pPr") and child not in mark:
                    p.remove(child)
        lead = (" " if in_note and mark else "")
        for el in self.make(segments(lead + text), donor, part, in_note=in_note):
            p.append(el)

    def insert_table(self, ch):
        el = self.el(ch["target"])
        like = self.el(ch.get("like") or "") if ch.get("like") else None
        if like is not None and like.tag != w("tbl"):
            raise Refused("`like` must be a table (T1)")
        rows = [[str(v) for v in r] for r in ch.get("rows") or []]
        self.place(ch["target"], el, self.table(rows, like), True)

    # -- blocks: the light Markdown a change's `after` may use

    def build(self, text, donor, part, like_table=None):
        """Elements for `after` text: paragraphs, `#` headings, `- ` bullets, `|` tables,
        `{P012}` to take a paragraph's formatting, `{T2}` alone to keep an original table."""
        out = []
        like = self.el(like_table) if like_table else None
        for block in [b.strip("\n") for b in text.split("\n\n") if b.strip()]:
            block = block.strip()
            keep = re.fullmatch(r"\{([PT]\d+)\}", block)
            if keep:
                out.append(self.el(keep.group(1)))
                continue
            d = donor
            m = plans.DONOR.match(block)
            if m:
                d = self.el(m.group(1))
                block = block[m.end():]
            lines = block.split("\n")
            if all(x.strip().startswith("|") for x in lines):
                rows = [[c.strip() for c in x.strip().strip("|").split("|")] for x in lines
                        if not re.fullmatch(r"[\s|:-]+", x)]
                out.append(self.table(rows, like))
                spacer = self.E("w:p")             # the empty line the document leaves after a table
                if d.find(w("pPr")) is not None:
                    spacer.append(deepcopy(d.find(w("pPr"))))
                out.append(spacer)
                continue
            head = re.match(r"(#{1,6})\s+(.*)", block)
            if head and "\n" not in block:
                hd = self.heading_donor(len(head.group(1)))
                text_ = head.group(2) if hd is not None else f"**{head.group(2).strip('*')}**"
                out.append(self.new_paragraph(hd if hd is not None else d, text_, part))
                continue
            if all(re.match(r"\s*[-*•]\s+", x) for x in lines):
                bullet = self.bullet_donor()
                bullet = d if bullet is None else bullet
                for x in lines:
                    out.append(self.new_paragraph(bullet, re.sub(r"^\s*[-*•]\s+", "", x), part))
                continue
            if all(re.match(r"\s*\d+[.)]\s+", x) for x in lines):
                numbered = self.bullet_donor("number")
                numbered = d if numbered is None else numbered
                for x in lines:
                    out.append(self.new_paragraph(numbered, re.sub(r"^\s*\d+[.)]\s+", "", x), part))
                continue
            out.append(self.new_paragraph(d, block, part))
        return out

    def heading_donor(self, level):
        """The document's own heading of that level; None when it has none, and the
        heading becomes a bold line in body formatting, as a run-in label."""
        for el in self.index.values():
            if el.tag == w("p"):
                _, lvl, _ = officedoc.docx_para_info(el, self.names)
                if lvl == level and officedoc.docx_text(el).strip():
                    return el
        return None

    def bullet_donor(self, kind="bullet"):
        """The document's own bulleted (or numbered) paragraph, if it has one."""
        kinds = officedoc.list_kinds(self.doc)
        for el in self.index.values():
            nid = el.find(f"{w('pPr')}/{w('numPr')}/{w('numId')}") if el.tag == w("p") else None
            if nid is not None and kinds.get(nid.get(w("val"))) == kind and officedoc.docx_text(el).strip():
                return el
        return None

    def table(self, rows, like=None):
        """A table in the document's own table formatting, as wide as the text column."""
        if not rows:
            raise Refused("a table needs rows")
        n = max(len(r) for r in rows)
        rows = [r + [""] * (n - len(r)) for r in rows]
        if like is None:
            like = next((e for e in self.index.values() if e.tag == w("tbl")), None)
        weights = [max(4, min(40, max(len(officedoc.plain(r[j])) for r in rows))) for j in range(n)]
        widths = [int(self.text_width * x / sum(weights)) for x in weights]
        tbl = self.E("w:tbl")
        src_rows = like.findall(w("tr")) if like is not None else []
        tblpr = deepcopy(like.find(w("tblPr"))) if like is not None and like.find(w("tblPr")) is not None else self.E("w:tblPr")
        for tag in ("tblW", "tblLayout"):
            for e_ in tblpr.findall(w(tag)):
                tblpr.remove(e_)
        tw = self.E("w:tblW")
        tw.set(w("w"), str(self.text_width))
        tw.set(w("type"), "dxa")
        tblpr.append(tw)
        layout = self.E("w:tblLayout")
        layout.set(w("type"), "fixed")
        tblpr.append(layout)
        tbl.append(tblpr)
        grid = self.E("w:tblGrid")
        for wd in widths:
            col = self.E("w:gridCol")
            col.set(w("w"), str(wd))
            grid.append(col)
        tbl.append(grid)
        head_tpl = src_rows[0] if src_rows else None
        body_tpls = src_rows[1:3] or src_rows[:1]
        for i, values in enumerate(rows):
            tpl = head_tpl if i == 0 else (body_tpls[(i - 1) % len(body_tpls)] if body_tpls else None)
            cell_tpl = tpl.find(w("tc")) if tpl is not None else None
            tr = self.E("w:tr")
            if tpl is not None and tpl.find(w("trPr")) is not None:
                tr.append(deepcopy(tpl.find(w("trPr"))))
            for wd, value in zip(widths, values):
                tc = self.E("w:tc")
                tcpr = deepcopy(cell_tpl.find(w("tcPr"))) if cell_tpl is not None and cell_tpl.find(w("tcPr")) is not None else self.E("w:tcPr")
                for tag in ("tcW", "gridSpan", "vMerge", "hMerge"):
                    for e_ in tcpr.findall(w(tag)):
                        tcpr.remove(e_)
                tcw = self.E("w:tcW")
                tcw.set(w("w"), str(wd))
                tcw.set(w("type"), "dxa")
                tcpr.insert(0, tcw)
                tc.append(tcpr)
                donor_p = cell_tpl.find(w("p")) if cell_tpl is not None else None
                para = self.E("w:p")
                if donor_p is not None and donor_p.find(w("pPr")) is not None:
                    para.append(deepcopy(donor_p.find(w("pPr"))))
                rpr = self.donor(donor_p) if donor_p is not None else None
                for el in self.make(segments(value), rpr, self.doc.part):
                    para.append(el)
                if i == 0 and not any(el.tag == w("r") and el.find(f"{w('rPr')}/{w('b')}") is not None
                                      for el in para.iter(w("r"))):
                    for r in para.iter(w("r")):
                        rp = r.find(w("rPr"))
                        if rp is None:
                            rp = self.E("w:rPr")
                            r.insert(0, rp)
                        self.rpr_set(rp, "b")
                tc.append(para)
                tr.append(tc)
            if self.track:
                trpr = tr.find(w("trPr"))
                if trpr is None:
                    trpr = self.E("w:trPr")
                    tr.insert(0, trpr)
                trpr.append(self.revision("ins"))
            tbl.append(tr)
        return tbl

    # -- sections

    def section(self, heading):
        """The body elements under a heading, up to the next heading of its level or above."""
        _, level, _ = officedoc.docx_para_info(heading, self.names)
        if level is None:
            raise Refused("rewrite_section targets a heading paragraph")
        out, el = [], heading.getnext()
        while el is not None and el.tag in (w("p"), w("tbl"), w("sdt"), w("bookmarkStart"), w("bookmarkEnd")):
            if el.tag == w("p"):
                _, lvl, _ = officedoc.docx_para_info(el, self.names)
                if lvl is not None and lvl <= level and officedoc.docx_text(el).strip():
                    break
            out.append(el)
            el = el.getnext()
        return out

    def rewrite_section(self, ch):
        heading = self.el(ch["target"])
        if heading.tag != w("p"):
            raise Refused("rewrite_section targets a heading paragraph")
        old = self.section(heading)
        first_body = next((e for e in old if e.tag == w("p") and officedoc.docx_text(e).strip()), heading)
        donor = self.el(ch["like"]) if ch.get("like") else first_body
        made = self.build(ch.get("after") or "", donor, self.doc.part, ch.get("like_table"))
        kept = {id(e) for e in made}
        ids = {id(e): k for k, e in self.index.items()}
        for e_ in old:
            if id(e_) in kept:
                continue
            key = ids.get(id(e_))
            if self.track and e_.tag == w("p"):
                self.normalize(e_)
                self.remove_runs(list(dict.fromkeys(x.run for x in self.atoms(e_))))
                self.mark_paragraph(e_, "del")
            elif self.track and e_.tag == w("tbl"):
                raise Refused("tracked mode cannot delete a whole table; keep it with {Tn}")
            elif e_.tag in (w("p"), w("tbl"), w("sdt")):
                e_.getparent().remove(e_)
            if key:
                self.deleted.update(k for k in self.index if k == key or k.startswith(key + "."))
        ref = heading
        for new in made:
            ref.addnext(new)
            ref = new
        self.cursor[ch["target"]] = ref

    # -- the clean-up

    def clean_up(self, ch):
        """Remove the working marks of a draft: struck text, draft colours, highlights, comments."""
        opts = {"strike": True, "colors": True, "highlight": True, "comments": True, **(ch.get("options") or {})}
        roots = [self.doc.element.body] + ([self.fn.root] if self.fn is not None else [])
        link_color = self.link_rpr.find(w("color")) if self.link_rpr is not None else None
        for root in roots:
            if opts["strike"]:
                for p in list(root.iter(w("p"))):
                    had = officedoc.docx_text(p).strip()
                    struck = [r for r in officedoc.docx_runs(p) if officedoc.run_marks(r)[0]]
                    if not struck:
                        continue
                    self.normalize(p)
                    struck = [r for r in officedoc.docx_runs(p) if officedoc.run_marks(r)[0]]
                    self.remove_runs(struck)
                    left = "".join(x.text for x in self.atoms(p) if x.run.getparent().tag != w("del"))
                    if had and not left.strip() and p.getparent() is not None and p.getparent().tag != w("tc"):
                        if self.track:
                            self.mark_paragraph(p, "del")
                        else:
                            p.getparent().remove(p)
            for r in root.iter(w("r")):
                rpr = r.find(w("rPr"))
                if rpr is None:
                    continue
                if opts["highlight"]:
                    for e_ in rpr.findall(w("highlight")):
                        rpr.remove(e_)
                if opts["colors"]:
                    col = rpr.find(w("color"))
                    if officedoc.in_link(r):
                        if col is not None and link_color is not None:
                            col.set(w("val"), link_color.get(w("val")))
                    elif col is not None:
                        if self.body_color:
                            col.set(w("val"), self.body_color)
                        else:
                            rpr.remove(col)
        if opts["comments"]:
            body = self.doc.element.body
            for tag in ("commentRangeStart", "commentRangeEnd"):
                for e_ in list(body.iter(w(tag))):
                    e_.getparent().remove(e_)
            for ref in list(body.iter(w("commentReference"))):
                run = ref.getparent()
                run.getparent().remove(run)
            if self.comments is not None:
                for c in list(self.comments.root):
                    self.comments.root.remove(c)

    def drop_orphan_comments(self):
        """A comment whose anchor an edit removed is removed too; Word mishandles orphans."""
        if self.comments is None:
            return
        body = self.doc.element.body
        refs = {e.get(w("id")) for e in body.iter(w("commentReference"))}
        for c in list(self.comments.root.findall(w("comment"))):
            if c.get(w("id")) not in refs:
                self.comments.root.remove(c)
                self.notes_out.append(f"comment {c.get(w('id'))} removed: the text it was attached to is gone")
        for tag in ("commentRangeStart", "commentRangeEnd"):
            for e_ in list(body.iter(w(tag))):
                if e_.get(w("id")) not in refs:
                    e_.getparent().remove(e_)

    def save(self, dest):
        for ch in self.late:
            self.clean_up(ch)
        self.drop_orphan_comments()
        if self.comments is not None:
            self.comments.flush()
        if self.fn is not None:
            self.fn.flush()
        self.doc.save(str(dest))


# ------------------------------------------------------------------ pptx

class PptxEditor(TextEditor):
    def __init__(self, src, track=False, author=""):
        try:
            from pptx import Presentation
            from pptx.oxml.xmlchemy import OxmlElement
        except ImportError:
            officedoc.need("pptx")
        self.E = OxmlElement
        self.prs = Presentation(str(src))
        self.index = officedoc.pptx_index(self.prs)
        self.cursor, self.deleted = {}, set()
        self.notes_out = []

    def atoms(self, p):
        return officedoc.pptx_atoms(p)

    def split(self, atom, k):
        clone = deepcopy(atom.run)
        text = atom.child.text or ""
        atom.child.text = text[:k]
        clone.find(a("t")).text = text[k:]
        atom.run.addnext(clone)

    def donor(self, p):
        for r in p.findall(a("r")):
            if (r.findtext(a("t")) or "").strip():
                rpr = r.find(a("rPr"))
                if rpr is None:
                    return None
                rpr = deepcopy(rpr)
                for h in rpr.findall(a("hlinkClick")):
                    rpr.remove(h)
                return rpr
        end = p.find(a("endParaRPr"))
        if end is not None:
            rpr = deepcopy(end)
            rpr.tag = a("rPr")
            return rpr
        return None

    def make(self, segs, donor, slide):
        out = []
        for seg in segs:
            if seg.get("footnote") is not None:
                raise Refused("slides have no footnotes; cite in the deck's own style")
            for i, line in enumerate(seg["text"].split("\n")):
                if i:
                    br = self.E("a:br")
                    if donor is not None:
                        br.append(deepcopy(donor))
                    out.append(br)
                if not line:
                    continue
                r = self.E("a:r")
                rpr = deepcopy(donor) if donor is not None else self.E("a:rPr")
                if seg.get("bold"):
                    rpr.set("b", "1")
                if seg.get("italic"):
                    rpr.set("i", "1")
                if seg.get("link"):
                    rid = slide.part.relate_to(seg["link"], HYPERLINK, is_external=True)
                    h = self.E("a:hlinkClick")
                    h.set(f"{{{R}}}id", rid)
                    tail = rpr.find(a("extLst"))
                    if tail is not None:
                        tail.addprevious(h)
                    else:
                        rpr.append(h)
                r.append(rpr)
                t = self.E("a:t")
                t.text = line
                r.append(t)
                out.append(r)
        return out

    def new_paragraph(self, donor_p, text, slide):
        p = self.E("a:p")
        ppr = donor_p.find(a("pPr"))
        if ppr is not None:
            p.append(deepcopy(ppr))
        for el in self.make(segments(text), self.donor(donor_p), slide):
            p.append(el)
        end = donor_p.find(a("endParaRPr"))
        if end is not None:
            p.append(deepcopy(end))
        return p

    def entry(self, target):
        if target in self.deleted:
            raise Refused(f"{target} was deleted by an earlier change")
        if target not in self.index:
            raise Refused(f"{target} is not a part of the original")
        return self.index[target]

    def apply(self, ch):
        op = ch["op"]
        if op == "replace":
            return self.replace(ch)
        if op in ("insert_after", "insert_before"):
            return self.insert(ch, op == "insert_after")
        if op == "delete":
            return self.delete(ch)
        if op == "add_slide":
            return self.add_slide(ch)
        raise Refused(f"op {op} is not available for PowerPoint")

    def replace(self, ch):
        el, slide, extra = self.entry(ch["target"])
        before, after = ch["before"], ch.get("after") or ""
        if el is slide:
            raise Refused("target a shape or a paragraph, not a whole slide")
        if el.tag != a("p"):
            body = self.txbody(el)
            paras = body.findall(a("p"))
            whole = "\n".join(officedoc.pptx_text(q) for q in paras)
            if plans.norm(before) == plans.norm(whole):
                return self.set_shape_text(el, after, slide)
            hits = [q for q in paras if before in officedoc.pptx_text(q)]
            if len(hits) != 1:
                raise Refused(f"`before` must be the whole text of {ch['target']} or sit in one paragraph")
            el = hits[0]
        s, e = self.locate(el, before, ch["target"])
        span = self.span(el, s, e)
        first = span[0].run
        donor = self.donor_of_run(first) if first.tag == a("r") else self.donor(el)
        for new in self.make(segments(after), donor, slide):
            first.addprevious(new)
        for x in span:
            x.run.getparent().remove(x.run)

    def donor_of_run(self, r):
        rpr = r.find(a("rPr"))
        if rpr is None:
            return None
        rpr = deepcopy(rpr)
        for h in rpr.findall(a("hlinkClick")):
            rpr.remove(h)
        return rpr

    def txbody(self, el):
        body = el.find(f"{{{P}}}txBody")
        if body is None:
            raise Refused("that shape holds no text")
        return body

    def set_shape_text(self, sp, text, slide):
        body = self.txbody(sp)
        donors = body.findall(a("p"))
        for q in donors:
            body.remove(q)
        for i, line in enumerate(text.split("\n")):
            d = donors[min(i, len(donors) - 1)]
            m = re.match(r"^\{p(\d+)\}\s*", line)
            if m:
                d = donors[min(int(m.group(1)) - 1, len(donors) - 1)]
                line = line[m.end():]
            body.append(self.new_paragraph(d, line, slide))

    def insert(self, ch, after):
        el, slide, extra = self.entry(ch["target"])
        if el.tag != a("p"):
            raise Refused("insert after a paragraph (S03.#5.p2); to add a slide use add_slide")
        donor = self.entry(ch["like"])[0] if ch.get("like") else el
        for line in [x for x in (ch.get("after") or "").split("\n") if x.strip()]:
            d = donor
            m = plans.DONOR.match(line)
            if m:
                d = self.entry(m.group(1))[0]
                line = line[m.end():]
            new = self.new_paragraph(d, line, slide)
            if after:
                ref = self.cursor.get(ch["target"], el)
                ref.addnext(new)
                self.cursor[ch["target"]] = new
            else:
                el.addprevious(new)

    def sld_id(self, slide):
        for s in self.prs.slides._sldIdLst:
            if self.prs.part.related_part(s.rId) is slide.part:
                return s
        raise Refused("slide not found in the deck")

    def delete(self, ch):
        target = ch["target"]
        el, slide, extra = self.entry(target)
        if el is slide:
            s = self.sld_id(slide)
            self.prs.slides._sldIdLst.remove(s)
            self.prs.part.drop_rel(s.rId)
        elif el.tag == a("p"):
            body = el.getparent()
            if len(body.findall(a("p"))) == 1:
                for child in list(el):
                    if child.tag in (a("r"), a("br"), a("fld")):
                        el.remove(child)
            else:
                body.remove(el)
        else:
            el.getparent().remove(el)
        self.deleted.update(k for k in self.index if k == target or k.startswith(target + "."))

    def add_slide(self, ch):
        anchor = self.entry(ch["target"])[1]
        src = self.entry(ch.get("like") or ch["target"])[1]
        new = self.prs.slides.add_slide(src.slide_layout)
        tree = new.shapes._spTree
        keep = (f"{{{P}}}nvGrpSpPr", f"{{{P}}}grpSpPr")
        for el in list(tree):
            if el.tag not in keep:
                tree.remove(el)
        rids = {}
        for rel in src.part.rels.values():
            if rel.reltype.endswith(("/slideLayout", "/notesSlide")):
                continue
            if rel.is_external:
                rids[rel.rId] = new.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
            else:
                rids[rel.rId] = new.part.relate_to(rel.target_part, rel.reltype)
                if rel.reltype.endswith("/chart"):
                    self.notes_out.append(f"{ch['id']}: the cloned chart shares its data with {ch.get('like')}")
        for el in src.shapes._spTree:
            if el.tag in keep:
                continue
            c = deepcopy(el)
            for node in c.iter():
                for attr in (f"{{{R}}}embed", f"{{{R}}}link", f"{{{R}}}id", f"{{{R}}}pict"):
                    if node.get(attr) in rids:
                        node.set(attr, rids[node.get(attr)])
            tree.append(c)
        bg = src._element.find(f"{{{P}}}cSld/{{{P}}}bg")
        if bg is not None:
            new._element.find(f"{{{P}}}cSld").insert(0, deepcopy(bg))
        for sid, text in (ch.get("fill") or {}).items():
            self.set_shape_text(self.shape_el(new, sid), text, new)
        for sid in ch.get("drop") or []:
            sp = self.shape_el(new, sid)
            sp.getparent().remove(sp)
        lst = self.prs.slides._sldIdLst
        made = lst[-1]
        lst.remove(made)
        self.cursor.get(ch["target"], self.sld_id(anchor)).addnext(made)
        self.cursor[ch["target"]] = made

    def shape_el(self, slide, sid):
        want = str(sid).lstrip("#")
        for cnv in slide.shapes._spTree.iter(f"{{{P}}}cNvPr"):
            if cnv.get("id") == want:
                return cnv.getparent().getparent()
        raise Refused(f"shape #{want} is not on the cloned slide")

    def save(self, dest):
        self.prs.save(str(dest))


# ------------------------------------------------------------------ xlsx

def col_index(letters):
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


class XlsxEditor:
    def __init__(self, src, track=False, author=""):
        from lxml import etree
        self.etree = etree
        with zipfile.ZipFile(src) as z:
            self.infos = z.infolist()
            self.files = {i.filename: z.read(i.filename) for i in self.infos}
        self.trees = {}
        wb = etree.fromstring(self.files["xl/workbook.xml"])
        rels = etree.fromstring(self.files["xl/_rels/workbook.xml.rels"])
        targets = {r.get("Id"): r.get("Target") for r in rels}
        self.sheets = {}
        for sh in wb.iter(f"{{{S}}}sheet"):
            t = targets.get(sh.get(f"{{{R}}}id"), "")
            self.sheets[sh.get("name")] = t.lstrip("/") if t.startswith("/") else "xl/" + t
        self.notes_out = []

    def tree(self, name):
        if name not in self.sheets:
            raise Refused(f"no sheet named {name!r}")
        path = self.sheets[name]
        if path not in self.trees:
            self.trees[path] = self.etree.fromstring(self.files[path])
        return self.trees[path]

    def apply(self, ch):
        if ch["op"] != "set_cell":
            raise Refused("spreadsheets take set_cell only: inserting rows would break the formulas")
        sheet, coord = officedoc.split_cell_id(ch["target"])
        m = re.fullmatch(r"([A-Z]{1,3})([1-9]\d*)", coord)
        if not m:
            raise Refused(f"{coord} is not a cell reference")
        col, row_n = m.group(1), int(m.group(2))
        tree = self.tree(sheet)
        data = tree.find(f"{{{S}}}sheetData")
        row = None
        for r in data.findall(f"{{{S}}}row"):
            n = int(r.get("r"))
            if n == row_n:
                row = r
                break
            if n > row_n:
                row = self.etree.Element(f"{{{S}}}row", r=str(row_n))
                r.addprevious(row)
                break
        if row is None:
            row = self.etree.SubElement(data, f"{{{S}}}row", r=str(row_n))
        cell = None
        for c in row.findall(f"{{{S}}}c"):
            letters = re.match(r"[A-Z]+", c.get("r")).group(0)
            if c.get("r") == coord:
                cell = c
                break
            if col_index(letters) > col_index(col):
                cell = self.etree.Element(f"{{{S}}}c", r=coord)
                c.addprevious(cell)
                break
        if cell is None:
            cell = self.etree.SubElement(row, f"{{{S}}}c", r=coord)
        if cell.get("s") is None:
            style = self.neighbour_style(data, col, row_n, row)
            if style:
                cell.set("s", style)
        f = cell.find(f"{{{S}}}f")
        if f is not None and f.get("t") == "shared" and f.get("ref"):
            raise Refused(f"{coord} is the master of a shared formula ({f.get('ref')}); "
                          f"changing it would change every cell that shares it")
        for child in list(cell):
            if child.tag in (f"{{{S}}}f", f"{{{S}}}v", f"{{{S}}}is"):
                cell.remove(child)
        cell.attrib.pop("t", None)
        value = str(ch.get("after", ""))
        if value.startswith("="):
            self.etree.SubElement(cell, f"{{{S}}}f").text = value[1:]
        elif re.fullmatch(r"-?\d+(\.\d+)?([eE][-+]?\d+)?", value):
            self.etree.SubElement(cell, f"{{{S}}}v").text = value
        elif value:
            cell.set("t", "inlineStr")
            t = self.etree.SubElement(self.etree.SubElement(cell, f"{{{S}}}is"), f"{{{S}}}t")
            t.text = value
            if value != value.strip():
                t.set(XML_SPACE, "preserve")

    def neighbour_style(self, data, col, row_n, row):
        for r in data.findall(f"{{{S}}}row"):
            if int(r.get("r")) == row_n - 1:
                for c in r.findall(f"{{{S}}}c"):
                    if c.get("r") == f"{col}{row_n - 1}":
                        return c.get("s")
        prev = None
        for c in row.findall(f"{{{S}}}c"):
            if c.get("r") == f"{col}{row_n}":
                break
            prev = c
        return prev.get("s") if prev is not None else None

    def save(self, dest):
        et = self.etree
        for path, tree in self.trees.items():
            self.files[path] = et.tostring(tree, xml_declaration=True, encoding="UTF-8", standalone=True)
        wb = et.fromstring(self.files["xl/workbook.xml"])
        calc = wb.find(f"{{{S}}}calcPr")
        if calc is None:
            calc = et.Element(f"{{{S}}}calcPr")
            anchor = wb.find(f"{{{S}}}definedNames")
            if anchor is None:
                anchor = wb.find(f"{{{S}}}sheets")
            anchor.addnext(calc)
        calc.set("fullCalcOnLoad", "1")
        self.files["xl/workbook.xml"] = et.tostring(wb, xml_declaration=True, encoding="UTF-8", standalone=True)
        drop = "xl/calcChain.xml"
        if drop in self.files:
            rels = et.fromstring(self.files["xl/_rels/workbook.xml.rels"])
            for r in list(rels):
                if r.get("Target", "").endswith("calcChain.xml"):
                    rels.remove(r)
            self.files["xl/_rels/workbook.xml.rels"] = et.tostring(rels, xml_declaration=True, encoding="UTF-8", standalone=True)
            ct = et.fromstring(self.files["[Content_Types].xml"])
            for o in list(ct):
                if (o.get("PartName") or "").endswith("calcChain.xml"):
                    ct.remove(o)
            self.files["[Content_Types].xml"] = et.tostring(ct, xml_declaration=True, encoding="UTF-8", standalone=True)
        with zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as z:
            for info in self.infos:
                if info.filename == drop:
                    continue
                z.writestr(info, self.files[info.filename])


# ------------------------------------------------------------------ md

class MdEditor:
    def __init__(self, src, track=False, author=""):
        self.blocks = [[f"B{i:03d}", t] for i, t in enumerate(officedoc.md_blocks(Path(src).read_text()), 1)]
        self.notes_out = []

    def find(self, target):
        for i, (bid, _) in enumerate(self.blocks):
            if bid == target:
                return i
        raise Refused(f"{target} is not a block of the original, or was deleted")

    def apply(self, ch):
        op, i = ch["op"], self.find(ch["target"])
        after = ch.get("after") or ""
        if op == "replace":
            text = self.blocks[i][1]
            n = text.count(ch["before"])
            if n != 1:
                raise Refused(f"`before` occurs {n} times in {ch['target']}")
            self.blocks[i][1] = text.replace(ch["before"], after)
        elif op == "delete":
            del self.blocks[i]
        elif op in ("insert_after", "insert_before"):
            new = [[None, b.strip()] for b in after.split("\n\n") if b.strip()]
            at = i + 1 if op == "insert_after" else i
            while op == "insert_after" and at < len(self.blocks) and self.blocks[at][0] is None:
                at += 1
            self.blocks[at:at] = new
        else:
            raise Refused(f"op {op} is not available for Markdown")

    def save(self, dest):
        Path(dest).write_text("\n\n".join(t for _, t in self.blocks) + "\n")


EDITORS = {"docx": DocxEditor, "pptx": PptxEditor, "xlsx": XlsxEditor, "md": MdEditor}


def main():
    argv = sys.argv[1:]
    author = "Revision"
    if "--author" in argv:
        i = argv.index("--author")
        author = argv[i + 1]
        del argv[i:i + 2]
    args = [x for x in argv if not x.startswith("--")]
    if len(args) < 2:
        sys.exit(__doc__)
    run = Path(args[0])
    item = subjects.find(run, args[1])
    slug = item["slug"]
    if item["role"] != "subject":
        sys.exit(f"ERROR {item['path']} is a template: its design is copied, the file itself is not edited")
    src = run / "external_data" / item["path"]
    fmt = officedoc.fmt_of(src)
    if fmt not in EDITORS:
        sys.exit(f"ERROR {src.name} is {fmt}: it cannot be edited in place. Rebuild it with /rebuild, "
                 f"or supply its source file (.docx, .pptx, .xlsx).")
    track = "--track" in sys.argv
    if track and fmt != "docx":
        print("note: tracked changes are a Word feature; writing a clean file")
        track = False
    problems, _ = plans.validate(run, slug, quiet=True)
    todo = plans.accepted(run, slug)
    stale = [c["id"] for c, st, _ in plans.states(run, slug) if st == "stale"]
    if stale:
        print(f"note: {', '.join(stale)} changed after the author decided; they are not applied")
    if not todo:
        sys.exit("ERROR no accepted changes to apply. The author accepts changes in the app's "
                 "Changes view or with changes.py decide.")
    blocking = [p for p in problems if p.split(":")[0] in {c["id"] for c in todo}]
    if blocking:
        for p in blocking:
            print(f"PROBLEM {p}")
        sys.exit("ERROR accepted changes fail validation; fix the plan (the author then decides again)")
    editor = EDITORS[fmt](src, track, author)
    results = []
    for ch in todo:
        try:
            editor.apply(ch)
            results.append({"id": ch["id"], "status": "applied"})
        except Refused as exc:
            results.append({"id": ch["id"], "status": "refused", "reason": str(exc)})
    refused = [r for r in results if r["status"] == "refused"]
    for r in refused:
        print(f"REFUSED {r['id']}: {r['reason']}")
    if refused and "--partial" not in sys.argv:
        sys.exit(f"ERROR {len(refused)} accepted change(s) could not be applied; nothing was written. "
                 f"Fix them, or rerun with --partial to write the rest.")
    base = run / "deliverable" / slug
    n = 1 + max([int(p.name[1:]) for p in base.glob("v*") if p.name[1:].isdigit()] or [0])
    vdir = base / f"v{n}"
    vdir.mkdir(parents=True)
    dest = vdir / src.name
    editor.save(dest)
    (vdir / "applied.json").write_text(json.dumps({
        "subject": item["path"], "file": src.name, "version": n, "format": fmt, "track": track,
        "source_sha256": subjects.sha256(src),
        "applied": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "results": results, "notes": editor.notes_out, "changes": todo,
    }, indent=2, ensure_ascii=False) + "\n")
    done = len(results) - len(refused)
    print(f"wrote {dest.relative_to(run)} · {done} change(s) applied" + (f", {len(refused)} refused" if refused else "")
          + (" · tracked" if track else ""))
    for note in editor.notes_out:
        print(f"note: {note}")
    print(f"next: python3 .claude/scripts/changes.py check {run} {slug}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
