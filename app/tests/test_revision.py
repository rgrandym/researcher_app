"""Revision runs: the change plan the app shows, the author's decisions, and the apply action."""

import asyncio
import json
import subprocess
import sys
from pathlib import Path

import pytest
from conftest import RUN
from fastapi import HTTPException

from backend import main, revision


def _subject(root: Path) -> Path:
    run = root / "outputs" / RUN
    (run / "external_data").mkdir(exist_ok=True)
    (run / "external_data" / "plan.md").write_text(
        "# Plan\n\nThe market is worth USD 9 billion.\n\nWe sell peptides.\n"
    )
    scripts = root / ".claude/scripts"
    for args in (["add", str(run), "plan.md"], ["extract", str(run)]):
        subprocess.run([sys.executable, str(scripts / "subject.py"), *args], check=True,
                       capture_output=True)
    (run / "subjects" / "plan" / "changes.json").write_text(json.dumps({"changes": [
        {"id": "C01", "op": "replace", "target": "B002", "before": "USD 9 billion",
         "after": "USD 9.7 billion", "kind": "update", "basis": "author", "rationale": "Newer figure."},
        {"id": "C02", "op": "insert_after", "target": "B003", "after": "We start with Activin A.",
         "kind": "add", "basis": "author", "rationale": "Author's instruction.", "pairs_with": ["C01"]},
    ], "open": [{"where": "Costs", "what": "GMP quote", "why": "no public price"}]}))
    revision._modules.clear()          # the fixture workspace has its own copy of the scripts
    return run


class Sender:
    locks: dict[str, str] = {}

    async def send(self, session: str, prompt: str) -> None:
        self.session, self.prompt = session, prompt


def test_plan_decisions_and_apply(root: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    run = _subject(root)
    listed = main.run_subjects(RUN)["subjects"]
    assert [(s["slug"], s["role"], s["counts"]) for s in listed] == [("plan", "subject", {"proposed": 2})]

    plan = main.run_changes(RUN, "plan")
    assert [c["status"] for c in plan["changes"]] == ["proposed", "proposed"]
    assert plan["changes"][0]["where"] == "Plan"
    assert plan["open"][0]["what"] == "GMP quote"

    sender = Sender()
    monkeypatch.setattr(main, "manager", sender)
    with pytest.raises(HTTPException, match="Accept at least one"):
        asyncio.run(main.apply_changes(RUN, main.ApplyChangesBody(slug="plan")))

    body = main.DecisionsBody(decisions={
        "C01": {"status": "accepted", "after": "USD 9.7 billion in 2026"},
        "C02": {"status": "rejected"},
    })
    plan = main.save_decisions(RUN, "plan", body)
    assert [c["status"] for c in plan["changes"]] == ["accepted", "rejected"]
    stored = json.loads((run / "subjects/plan/decisions.json").read_text())["decisions"]
    assert stored["C01"]["after"] == "USD 9.7 billion in 2026"
    assert stored["C01"]["by"] == "author"

    # The agent edits an accepted change: it goes back to the author.
    data = json.loads((run / "subjects/plan/changes.json").read_text())
    data["changes"][0]["after"] = "USD 9.8 billion"
    (run / "subjects/plan/changes.json").write_text(json.dumps(data))
    assert main.run_changes(RUN, "plan")["changes"][0]["status"] == "stale"
    main.save_decisions(RUN, "plan", main.DecisionsBody(decisions={"C01": {"status": "accepted"}}))

    result = asyncio.run(main.apply_changes(RUN, main.ApplyChangesBody(slug="plan", mode="clean", notes="Keep it short.")))
    assert result["session"] == sender.session
    assert sender.prompt.startswith(f"/revise outputs/{RUN} plan")
    request = next((run / "deliverable/plan").glob("request-*.json"))
    assert request.name in sender.prompt
    assert json.loads(request.read_text())["accepted"] == ["C01"]

    with pytest.raises(HTTPException):
        main.save_decisions(RUN, "plan", main.DecisionsBody(decisions={"C99": {"status": "accepted"}}))
    with pytest.raises(HTTPException):
        main.save_decisions(RUN, "plan", main.DecisionsBody(decisions={"C01": {"status": "maybe"}}))
    with pytest.raises(HTTPException):
        main.run_changes(RUN, "nope")


def test_pdf_subject_is_rebuilt(root: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    run = _subject(root)
    source = run / "subjects/plan/source.json"
    meta = json.loads(source.read_text())
    meta.update({"editable": False, "format": "pdf"})
    source.write_text(json.dumps(meta))
    main.save_decisions(RUN, "plan", main.DecisionsBody(decisions={"C01": {"status": "accepted"}}))
    sender = Sender()
    monkeypatch.setattr(main, "manager", sender)
    asyncio.run(main.apply_changes(RUN, main.ApplyChangesBody(slug="plan")))
    assert sender.prompt.startswith(f"/rebuild outputs/{RUN} plan")
    request = next((run / "deliverable/plan").glob("request-*.json"))
    assert json.loads(request.read_text())["mode"] == "clean"


def test_apply_end_to_end_with_scripts(root: Path) -> None:
    """The workspace scripts apply an accepted Markdown change and the check passes."""
    run = _subject(root)
    main.save_decisions(RUN, "plan", main.DecisionsBody(decisions={
        "C01": {"status": "accepted"}, "C02": {"status": "accepted"}}))
    scripts = root / ".claude/scripts"
    subprocess.run([sys.executable, str(scripts / "apply_changes.py"), str(run), "plan"],
                   check=True, capture_output=True)
    out = (run / "deliverable/plan/v1/plan.md").read_text()
    assert "USD 9.7 billion" in out and "We start with Activin A." in out
    check = subprocess.run([sys.executable, str(scripts / "changes.py"), "check", str(run), "plan"],
                           capture_output=True, text=True)
    assert check.returncode == 0, check.stdout
    assert main.run_subjects(RUN)["subjects"][0]["deliverable"]["ok"] is True


def _word_subject(root: Path) -> Path:
    import docx

    run = root / "outputs" / RUN
    (run / "external_data").mkdir(exist_ok=True)
    d = docx.Document()
    d.add_heading("Plan", level=1)
    d.add_heading("1. Market", level=2)
    d.add_paragraph("The market is worth USD 9 billion.")
    struck = d.add_paragraph().add_run("An old claim the author struck out.")
    struck.font.strike = True
    d.add_heading("2. Team", level=2)
    d.add_paragraph("Two founders run the company.")
    d.save(str(run / "external_data" / "plan.docx"))
    scripts = root / ".claude/scripts"
    for args in (["add", str(run), "plan.docx"], ["extract", str(run)]):
        subprocess.run([sys.executable, str(scripts / "subject.py"), *args], check=True, capture_output=True)
    revision._modules.clear()
    return run


def test_word_plan_from_markdown(root: Path) -> None:
    run = _word_subject(root)
    with (run / "evidence.jsonl").open("a") as ev:
        ev.write('\n{"id": "E1", "claim": "c", "url": "https://example.org/a"}\n')
    folder = run / "subjects" / "plan"
    text = (folder / "document.md").read_text()
    assert "~~An old claim the author struck out.~~" in text
    text = text.replace("Two founders run the company.",
                        "<!-- why: the team grew [E1] -->\nThree founders run the company.")
    text = text.replace("The market is worth USD 9 billion.",
                        "The market is worth USD 9.7 billion.\n\n| Layer | Value |\n| --- | --- |\n| Media | USD 9.7 billion |")
    (folder / "revised.md").write_text(text)
    plan = main.plan_from_markdown(RUN, "plan")
    by_id = {c["id"]: c for c in plan["changes"]}
    assert plan["revised"] == "subjects/plan/revised.md"
    assert by_id["C-P006"]["op"] == "rewrite" and by_id["C-P006"]["rationale"] == "the team grew"
    assert by_id["C-P002"]["op"] == "rewrite_section"
    assert "C-clean" in by_id

    # The author edits revised.md: decisions on untouched changes stand.
    main.save_decisions(RUN, "plan", main.DecisionsBody(decisions={c: {"status": "accepted"} for c in by_id}))
    (folder / "revised.md").write_text(text.replace("Three founders", "Three founders and an adviser"))
    status = {c["id"]: c["status"] for c in main.plan_from_markdown(RUN, "plan")["changes"]}
    assert status["C-P006"] == "stale" and status["C-P002"] == "accepted"

    main.save_decisions(RUN, "plan", main.DecisionsBody(decisions={"C-P006": {"status": "accepted"}}))
    scripts = root / ".claude/scripts"
    subprocess.run([sys.executable, str(scripts / "apply_changes.py"), str(run), "plan"], check=True,
                   capture_output=True)
    check = subprocess.run([sys.executable, str(scripts / "changes.py"), "check", str(run), "plan"],
                           capture_output=True, text=True)
    assert check.returncode == 0, check.stdout
    out = (run / "deliverable/plan/v1/text.md").read_text()
    assert "Three founders and an adviser" in out and "USD 9.7 billion" in out
    assert "struck out" not in out


def test_word_document_master_exports_without_a_plan(root: Path) -> None:
    run = _word_subject(root)
    with (run / "evidence.jsonl").open("a") as ev:
        ev.write('\n{"id": "E1", "claim": "c", "url": "https://example.org/a"}\n')
    scripts = root / ".claude/scripts"
    subprocess.run([sys.executable, str(scripts / "document.py"), "start", str(run), "plan"],
                   check=True, capture_output=True)
    master = run / "plan.md"
    assert main.run_subjects(RUN)["subjects"][0]["master"] == "plan.md"
    meta = json.loads((run / ".run.json").read_text())
    assert meta["subjects"] and meta["documents"][0]["master"] == "plan.md"

    # The master is edited as a report is, with evidence IDs in the text.
    master.write_text(master.read_text().replace(
        "Two founders run the company.", "Three founders run the company [E1]."))
    diff = main.document_diff(RUN, "plan")
    assert "1 changed" in diff["summary"]
    assert "Three founders" in (run / diff["path"]).read_text()

    result = asyncio.run(main.document_export(RUN, "plan", main.DocumentExportBody(formats=["docx"])))
    assert result["ok"], result["output"]
    assert result["files"] == ["deliverable/plan/v1/plan.docx"]
    out = (run / "deliverable/plan/v1/text.md").read_text()
    assert "Three founders run the company." in out
    assert "[E1]" not in out and "struck out" not in out
    assert not (run / "subjects/plan/changes.json").exists()

    # An unsourced claim does not reach the finished file.
    master.write_text(master.read_text().replace("Three founders", "Three founders [UNVERIFIED]"))
    result = asyncio.run(main.document_export(RUN, "plan", main.DocumentExportBody(formats=["docx"])))
    assert not result["ok"] and "UNVERIFIED" in result["output"]
