"""Editable Office documents: saved content, native fidelity, and export contracts."""

import asyncio
import json
import re
from datetime import datetime
from pathlib import Path

import pytest
from conftest import RUN
from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Font, PatternFill
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from backend import main, revision, workspace


def declare(root: Path, fmt: str) -> tuple[Path, Path]:
    run = root / "outputs" / RUN
    ext = run / "external_data"
    ext.mkdir(exist_ok=True)
    source = ext / f"sample.{fmt}"
    meta = json.loads((run / ".run.json").read_text()) if (run / ".run.json").exists() else {}
    meta["subjects"] = [{"slug": "sample", "path": source.name, "role": "subject"}]
    (run / ".run.json").write_text(json.dumps(meta))
    revision._modules.clear()
    return run, source


def workbook(root: Path) -> tuple[Path, Path]:
    run, source = declare(root, "xlsx")
    book = Workbook()
    sheet = book.active
    sheet.title = "Budget"
    sheet.append(["Item", "Cost", "Total"])
    sheet.append(["Salary", 70000, "=B2*2"])
    sheet["A1"].font = Font(name="Georgia", bold=True, color="FF123456")
    sheet["A1"].fill = PatternFill("solid", fgColor="FFAA6600")
    sheet["B2"].number_format = '£#,##0.00'
    sheet["A4"] = "Merged"
    sheet.merge_cells("A4:B4")
    sheet["D1"] = datetime(2026, 10, 26)
    sheet["E1"] = False
    sheet["F1"] = "=literal"
    sheet["F1"].data_type = "s"
    sheet.freeze_panes = "B2"
    sheet.column_dimensions["A"].width = 34
    sheet.print_area = "A1:F10"
    chart = BarChart()
    chart.add_data(Reference(sheet, min_col=2, min_row=1, max_row=2), titles_from_data=True)
    sheet.add_chart(chart, "H2")
    book.create_sheet("Hidden").sheet_state = "hidden"
    book.save(source)
    return run, source


def presentation(root: Path) -> tuple[Path, Path]:
    run, source = declare(root, "pptx")
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    shape = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    p = shape.text_frame.paragraphs[0]
    a = p.add_run()
    a.text = "Our "
    a.font.bold = True
    a.font.name = "Georgia"
    a.font.size = Pt(24)
    a.font.color.rgb = RGBColor(10, 20, 30)
    b = p.add_run()
    b.text = "plan"
    b.font.italic = True
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(3), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Budget"
    table.cell(1, 1).text = "£500k"
    slide.notes_slide.notes_text_frame.text = "Private speaker note"
    deck.slides.add_slide(deck.slide_layouts[6])
    deck.save(source)
    return run, source


def start() -> dict:
    asyncio.run(main.office_start(RUN, "sample", main.OfficeStartBody()))
    return main.office_content(RUN, "sample")


def save(content: dict) -> None:
    office = revision.script("office_master")
    workspace.save_run_file(RUN, "sample.md", office.encode(content["data"]), content["mtime"])


def export(fmt: str, design: str | None = None) -> dict:
    return asyncio.run(main.document_export(RUN, "sample", main.DocumentExportBody(formats=[fmt], design=design)))


def test_workbook_edit_export_preserves_native_features(root: Path) -> None:
    run, source = workbook(root)
    original = source.read_bytes()
    content = start()
    content["data"]["sheets"][0]["cells"]["B2"] = {"kind": "number", "value": 80000}
    content["data"]["sheets"][0]["cells"]["G1"] = {"kind": "text", "value": "=do not calculate"}
    save(content)
    result = export("xlsx")
    assert result["ok"]
    book = load_workbook(run / result["files"][0])
    sheet = book["Budget"]
    assert sheet["B2"].value == 80000
    assert sheet["C2"].value == "=B2*2"
    assert sheet["G1"].data_type == "s"
    assert sheet["A1"].font.name == "Georgia"
    assert sheet["A1"].fill.fgColor.rgb == "FFAA6600"
    assert sheet["B2"].number_format == '£#,##0.00'
    assert str(sheet.merged_cells) == "A4:B4"
    assert sheet.freeze_panes == "B2"
    assert sheet.column_dimensions["A"].width == 34
    assert len(sheet._charts) == 1
    assert book["Hidden"].sheet_state == "hidden"
    assert book.calculation.fullCalcOnLoad
    assert source.read_bytes() == original
    assert main.run_subjects(RUN)["subjects"][0]["master"] == "sample.md"
    # Retry reopens the same saved master, including the author's change.
    retry = start()
    assert retry["data"]["sheets"][0]["cells"]["B2"]["value"] == 80000
    assert len(list((run / "deliverable/sample").glob("v*"))) == 1


def test_new_workbook_theme_and_new_sheet(root: Path) -> None:
    run, _ = workbook(root)
    content = start()
    content["data"]["sheets"].append({"id": "new-sheet", "name": "Scenario", "cells": {"A1": {"kind": "text", "value": "Alternative"}}})
    content["data"]["design"] = "new"
    save(content)
    result = export("xlsx")
    book = load_workbook(run / result["files"][0])
    assert book["Budget"]["A1"].font.name == "Aptos"
    assert book["Budget"]["A1"].fill.fgColor.rgb == "0017365D"
    assert book["Budget"]["C2"].value == "=B2*2"
    assert book["Scenario"]["A1"].value == "Alternative"


@pytest.mark.parametrize("design", ["preserve", "new"])
def test_presentation_text_table_notes_and_added_slide(root: Path, design: str) -> None:
    run, source = presentation(root)
    original = source.read_bytes()
    content = start()
    slide = content["data"]["slides"][0]
    slide["shapes"][0]["paragraphs"] = ["Our revised plan"]
    slide["shapes"][1]["rows"][1][1] = "£600k"
    content["data"]["slides"].append({"id": "new-slide", "shapes": [{"id": "new-title", "kind": "text", "paragraphs": ["Next steps", "Validate"]}]})
    save(content)
    result = export("pptx", design)
    deck = Presentation(run / result["files"][0])
    assert len(deck.slides) == 3
    shape = deck.slides[0].shapes[0]
    assert shape.text == "Our revised plan"
    assert shape.left == Inches(1) and shape.top == Inches(1)
    assert shape.text_frame.paragraphs[0].runs[0].font.name == ("Georgia" if design == "preserve" else "Aptos")
    assert shape.text_frame.paragraphs[0].runs[-1].font.italic
    assert deck.slides[0].shapes[1].table.cell(1, 1).text == "£600k"
    assert "Private speaker note" in deck.slides[0].notes_slide.notes_text_frame.text
    assert deck.slides[2].shapes[0].text == "Next steps\nValidate"
    assert source.read_bytes() == original


def test_failed_export_does_not_publish_version(root: Path) -> None:
    run, _ = workbook(root)
    content = start()
    content["data"]["sheets"][0]["cells"]["B4"] = {"kind": "text", "value": "merge collision"}
    save(content)
    with pytest.raises(HTTPException, match="merged cell"):
        export("xlsx")
    assert not (run / "deliverable/sample/v1").exists()


def test_conflicts_validation_and_wrong_formats(root: Path) -> None:
    run, _ = workbook(root)
    content = start()
    stale = dict(content)
    content["data"]["design"] = "new"
    save(content)
    with pytest.raises(workspace.Conflict):
        save(stale)
    content = main.office_content(RUN, "sample")
    content["data"]["sheets"][0]["name"] = "Renamed"
    with pytest.raises(ValueError, match="Existing names"):
        save(content)
    with pytest.raises(HTTPException, match="as .xlsx"):
        export("pptx")
    with pytest.raises(HTTPException, match="preserve or new"):
        export("xlsx", "unknown")
    assert not (run / "deliverable/sample").exists()


def test_start_from_existing_deliverable_and_reject_other_files(root: Path) -> None:
    run, source = workbook(root)
    delivered = run / "deliverable/sample/v5"
    delivered.mkdir(parents=True)
    book = load_workbook(source)
    book["Budget"]["B2"] = 90000
    book.save(delivered / "updated.xlsx")
    asyncio.run(main.office_start(RUN, "sample", main.OfficeStartBody(source="deliverable/sample/v5/updated.xlsx")))
    assert main.office_content(RUN, "sample")["data"]["sheets"][0]["cells"]["B2"]["value"] == 90000
    other = run / "other.xlsx"
    other.write_bytes(source.read_bytes())
    with pytest.raises(HTTPException, match="supplied file"):
        asyncio.run(main.office_start(RUN, "sample", main.OfficeStartBody(source="other.xlsx")))


def test_busy_project_refuses_start_and_export(root: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    workbook(root)
    monkeypatch.setattr(main, "_run_busy", lambda _: True)
    with pytest.raises(HTTPException) as exc:
        start()
    assert exc.value.status_code == 409
    with pytest.raises(HTTPException) as exc:
        export("xlsx")
    assert exc.value.status_code == 409


def test_uploaded_file_opens_without_a_change_plan(root: Path) -> None:
    run, source = workbook(root)
    (run / ".run.json").write_text(json.dumps({"subjects": []}))
    result = asyncio.run(main.office_open(RUN, main.OfficeOpenBody(path=f"external_data/{source.name}")))
    assert result["master"] == "sample.md"
    assert main.office_content(RUN, "sample")["data"]["format"] == "xlsx"
    again = asyncio.run(main.office_open(RUN, main.OfficeOpenBody(path=f"external_data/{source.name}")))
    assert again == result
    assert len(json.loads((run / ".run.json").read_text())["documents"]) == 1


def test_date_blank_and_protected_cells(root: Path) -> None:
    run, source = workbook(root)
    content = start()
    cells = content["data"]["sheets"][0]["cells"]
    cells["A8"] = {"kind": "date", "value": "2026-10-27"}
    cells["A9"] = {"kind": "text", "value": ""}
    save(content)
    result = export("xlsx")
    book = load_workbook(run / result["files"][0])
    assert book["Budget"]["A8"].value == datetime(2026, 10, 27)
    assert book["Budget"]["A9"].value is None
    # A changed original source is refused; no stale coordinate writes.
    original = run / "subjects/sample/office-original.xlsx"
    original.write_bytes(source.read_bytes() + b"changed")
    with pytest.raises(HTTPException, match="source changed"):
        export("xlsx")


def test_protected_workbook_is_not_silently_unlocked(root: Path) -> None:
    run, source = workbook(root)
    book = load_workbook(source)
    book["Budget"].protection.sheet = True
    book.save(source)
    content = start()
    content["data"]["sheets"][0]["cells"]["B2"]["value"] = 42
    save(content)
    with pytest.raises(HTTPException, match="protected"):
        export("xlsx")
    assert not (run / "deliverable/sample/v1").exists()


def test_reorder_and_remove_slides(root: Path) -> None:
    run, _ = presentation(root)
    content = start()
    content["data"]["slides"].reverse()
    save(content)
    result = export("pptx")
    deck = Presentation(run / result["files"][0])
    assert len(deck.slides[0].shapes) == 0
    assert deck.slides[1].shapes[0].text == "Our plan"
    content = main.office_content(RUN, "sample")
    content["data"]["slides"].pop(0)
    save(content)
    result = export("pptx")
    deck = Presentation(run / result["files"][0])
    assert len(deck.slides) == 1
    assert deck.slides[0].shapes[0].text == "Our plan"


def delivered_forecast(root: Path) -> Path:
    """A workbook an agent built into deliverable/ from a supplied file nobody declared."""
    run, source = workbook(root)
    (run / ".run.json").write_text(json.dumps({"subjects": []}))
    original = source.rename(source.with_name("v4 - Cash flow forecast-2.xlsx"))
    out = run / "deliverable/cash-flow-forecast/v5"
    out.mkdir(parents=True)
    book = load_workbook(original)
    book["Budget"]["B3"] = 0.1545
    book["Budget"]["B3"].number_format = "0.0%"
    book.save(out / "v5 - Cash flow forecast.xlsx")
    return run


def test_delivered_workbook_is_listed_rendered_and_opened_for_editing(root: Path) -> None:
    run = delivered_forecast(root)
    path = "deliverable/cash-flow-forecast/v5/v5 - Cash flow forecast.xlsx"
    assert [o["path"] for o in main.office_outputs(RUN)["outputs"]] == [path]
    page = asyncio.run(main.office_html(RUN, path)).body.decode()
    assert 'data-sheet="Budget"' in page and "£70,000.00" in page and "15.5%" in page
    # Written by a script, so no value is stored: the app calculates it, and says so.
    assert 'title="=B2*2  (calculated by the app)"' in page and ">140000</td>" in page
    assert "1 formula value was not stored" in page and 'class="formula"' not in page
    assert 'colspan="2"' in page and "Hidden (hidden)" in page
    result = asyncio.run(main.office_open(RUN, main.OfficeOpenBody(path=path)))
    assert result["master"] == "cash-flow-forecast.md"
    declared = json.loads((run / ".run.json").read_text())["subjects"]
    assert declared == [{"path": "v4 - Cash flow forecast-2.xlsx", "role": "subject", "slug": "cash-flow-forecast"}]
    assert main.office_outputs(RUN)["outputs"] == []  # the editable master stands for it now
    assert main.office_content(RUN, "cash-flow-forecast")["data"]["sheets"][0]["cells"]["B3"]["value"] == 0.1545
    assert "15.5%" in asyncio.run(main.office_html(RUN, "cash-flow-forecast.md")).body.decode()
    exported = asyncio.run(main.document_export(RUN, "cash-flow-forecast", main.DocumentExportBody(formats=["xlsx"])))
    assert exported["files"] == ["deliverable/cash-flow-forecast/v6/v6 - Cash flow forecast.xlsx"]
    assert (run / path).is_file()  # the delivered file is never changed


def test_delivered_file_without_its_original_is_viewable_not_editable(root: Path) -> None:
    run = delivered_forecast(root)
    (run / "external_data/v4 - Cash flow forecast-2.xlsx").unlink()
    path = "deliverable/cash-flow-forecast/v5/v5 - Cash flow forecast.xlsx"
    assert "Budget" in asyncio.run(main.office_html(RUN, path)).body.decode()
    with pytest.raises(HTTPException, match="supplied original"):
        asyncio.run(main.office_open(RUN, main.OfficeOpenBody(path=path)))
    with pytest.raises(HTTPException) as exc:
        asyncio.run(main.office_html(RUN, "prompt.md"))
    assert exc.value.status_code in (400, 404)


def test_delivered_presentation_renders_slides_to_click(root: Path) -> None:
    run, source = presentation(root)
    out = run / "deliverable/sample/v1"
    out.mkdir(parents=True)
    (out / "sample.pptx").write_bytes(source.read_bytes())
    page = asyncio.run(main.office_html(RUN, "deliverable/sample/v1/sample.pptx")).body.decode()
    assert 'data-slide="0"' in page and "data-shape=" in page and "<table>" in page
    assert "Approximate preview: PowerPoint is not installed" in page


def test_presentation_shows_powerpoints_pictures_with_playing_video(root: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import io

    import pymupdf
    from backend import office_view

    run, source = presentation(root)
    deck = Presentation(str(source))
    movie = deck.slides[1].shapes.add_movie(io.BytesIO(b"\x00\x00\x00\x18ftypmp42"), Inches(1), Inches(1), Inches(4), Inches(3), mime_type="video/mp4")
    deck.save(source)
    exports: list[Path] = []

    def export(src: Path, pdf: Path) -> None:  # stands in for PowerPoint: one page per slide
        exports.append(src)
        with pymupdf.open() as doc:
            for _ in range(2):
                doc.new_page(width=720, height=405)
            doc.save(pdf)

    monkeypatch.setattr(office_view, "_powerpoint_pdf", export)
    out = run / "deliverable/sample/v1"
    out.mkdir(parents=True)
    (out / "sample.pptx").write_bytes(source.read_bytes())
    page = asyncio.run(main.office_html(RUN, "deliverable/sample/v1/sample.pptx")).body.decode()
    assert "Approximate" not in page and page.count('class="page"') == 2
    assert 'class="shape hit edit"' in page and f'data-shape="{movie.shape_id}"' in page
    video = re.search(r'<video[^>]*src="/api/office/render/(\w+)/([\w.-]+)\?token=', page)
    assert video and video.group(2).endswith(".mp4")
    served = main.office_render_file(video.group(1), video.group(2))
    assert Path(served.path).read_bytes().endswith(b"ftypmp42")
    picture = re.search(r'class="page" src="/api/office/render/\w+/(slide-001\.png)', page)
    assert picture and main.office_render_file(video.group(1), picture.group(1))
    asyncio.run(main.office_html(RUN, "deliverable/sample/v1/sample.pptx"))
    assert len(exports) == 1  # the second view reads the cache
    for bad in ("../slides.json", ".hidden", "slides.json/.."):
        with pytest.raises(HTTPException):
            main.office_render_file(video.group(1), bad)


def test_let_is_expanded_for_the_calculation_engine() -> None:
    from backend.office_view import _without_let

    assert _without_let("=LET(a, 1, b, a+1, a*b) + 2 ") == "=((1)*((1)+1))+2"
    nested = "=_xlfn.LET(_xlpm.x, C$1, _xlfn.LET(_xlpm.y, _xlpm.x*2, _xlpm.y+_xlpm.x))"
    assert _without_let(nested) == "=((((C$1)*2)+(C$1)))"
    assert _without_let('=SUM(A1:A3) + N("LET(x)")') == '=SUM(A1:A3)+N("LET(x)")'


def test_slide_design_moves_restyles_and_rethemes_in_export(root: Path) -> None:
    run, source = presentation(root)
    content = start()
    assert "Georgia" in content["palette"]["fonts"] and "0A141E" in content["palette"]["colors"]
    data = content["data"]
    title = data["slides"][0]["shapes"][0]
    assert title["box"] == {"x": 72.0, "y": 72.0, "w": 432.0, "h": 72.0}
    assert title["style"] == {"font": "Georgia", "size": 24.0, "bold": True, "color": "0A141E"}
    title["box"]["x"] = 144
    title["style"].update(size=30, color="AA0000")
    data["theme"]["colors"]["accent1"] = "123456"
    data["theme"]["fonts"]["major"] = "Georgia"
    data["theme"]["swap_fonts"] = {"Georgia": "Helvetica"}
    save(content)
    result = export("pptx")
    deck = Presentation(run / result["files"][0])
    shape = deck.slides[0].shapes[0]
    run_ = shape.text_frame.paragraphs[0].runs[0]
    assert shape.left == Inches(2) and shape.top == Inches(1)
    assert run_.font.size == Pt(30) and str(run_.font.color.rgb) == "AA0000"
    assert run_.font.name == "Helvetica"  # the swap reaches the box's own font
    theme = revision.script("office_master")._read_theme(deck)
    assert theme["colors"]["accent1"] == "123456" and theme["fonts"]["major"] == "Georgia"
    assert shape.text == "Our plan" and "Private speaker note" in deck.slides[0].notes_slide.notes_text_frame.text


def test_slide_design_edits_that_would_lose_something_are_refused(root: Path) -> None:
    presentation(root)
    content = start()
    title = content["data"]["slides"][0]["shapes"][0]
    for broken in ({**title, "style": {}}, {**title, "box": {"x": 1, "y": 1, "w": -5, "h": 1}},
                   {**title, "style": {**title["style"], "color": "red"}}):
        content["data"]["slides"][0]["shapes"][0] = broken
        with pytest.raises(ValueError):
            save(content)


def test_slides_script_edits_and_measures_the_master(root: Path) -> None:
    run, _ = presentation(root)
    start()
    slides = revision.script("slides")
    office = revision.script("office_master")
    master = run / "sample.md"
    data = office.decode(master.read_text())
    data = slides.apply(data, [
        {"op": "box", "slide": 1, "shape": data["slides"][0]["shapes"][0]["id"], "box": {"x": 40}},
        {"op": "style", "slide": 1, "shape": data["slides"][0]["shapes"][0]["id"], "style": {"size": 12, "fill": "FFFFFF", "color": "EEEEEE"}},
        {"op": "swap_color", "from": "0A141E", "to": "111111"},
        {"op": "order", "slides": [2, 1]},
    ])
    assert data["slides"][1]["shapes"][0]["box"]["x"] == 40 and data["theme"]["swap_colors"] == {"0A141E": "111111"}
    report = slides.inspect(data)
    issues = " ".join(report["slides"][1]["issues"])
    assert "12 pt" in issues and "contrast" in issues
    with pytest.raises(ValueError):
        slides.apply(data, [{"op": "order", "slides": [1, 1]}])


def test_powerpoint_draws_a_deck_once_until_it_is_edited(root: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    import pymupdf

    from backend import config, office_view

    calls = []

    def fake_powerpoint(src: Path, pdf: Path) -> None:
        assert src.parent == config.data_dir(office_view.RENDERS) / "powerpoint"  # one folder, asked about once
        calls.append(src)
        doc = pymupdf.open()
        for _ in Presentation(str(src)).slides:
            doc.new_page()
        doc.save(pdf)

    monkeypatch.setattr(office_view, "_powerpoint_pdf", fake_powerpoint)
    presentation(root)
    content = start()
    office_view.html(RUN, "sample.md")
    office_view.html(RUN, "sample.md")  # a reload: the rebuilt file differs, the deck does not
    assert len(calls) == 1
    content["data"]["slides"][0]["shapes"][0]["box"]["x"] = 100
    save(content)
    office_view.html(RUN, "sample.md")
    assert len(calls) == 2
