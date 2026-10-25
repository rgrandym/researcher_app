"""Individual table appearance in a Word export: the same settings the HTML carries.

pandoc gives every table the reference document's "Table" style. A table the author
styled in the Palette and layout panel is then dressed here, cell by cell: header and
body fills, banding, text colour, font and size, rule weight, colour, line type and
layout, and width. A table is found by its caption, "Table n.", above it or below it
as the Markdown places it, and n is matched to its id in document order, as figures.py
numbers it.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from docx import Document  # type: ignore[import-untyped,import-not-found,unused-ignore]
from docx.oxml import OxmlElement  # type: ignore[import-untyped,import-not-found,unused-ignore]
from docx.oxml.ns import qn  # type: ignore[import-untyped,import-not-found,unused-ignore]
from docx.shared import Pt, RGBColor  # type: ignore[import-untyped,import-not-found,unused-ignore]

CAPTION = re.compile(r"^\s*Table (\d+)\.")
LINE = {"solid": "single", "dashed": "dashed", "dotted": "dotted", "double": "double"}
EDGES = ("top", "left", "bottom", "right", "insideH", "insideV")


def _hex(value: str) -> str:
    return value.lstrip("#").upper()


def _child(parent: Any, tag: str) -> Any:
    found = parent.find(qn(tag))
    if found is None:
        found = OxmlElement(tag)
        parent.append(found)
    return found


def _border(parent: Any, edge: str, item: dict[str, Any] | None, heavy: float = 1.0) -> None:
    el = _child(parent, f"w:{edge}")
    if item is None:
        el.set(qn("w:val"), "nil")
        return
    el.set(qn("w:val"), LINE.get(item.get("rule_style", "solid"), "single"))
    el.set(qn("w:sz"), str(max(2, round(item["rule_weight_pt"] * heavy * 8))))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), _hex(item["border"]))


def _shade(cell: Any, fill: str) -> None:
    pr = cell._tc.get_or_add_tcPr()
    shd = _child(pr, "w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(fill))


def _cell_bottom(cell: Any, item: dict[str, Any] | None, heavy: float = 1.0) -> None:
    pr = cell._tc.get_or_add_tcPr()
    _border(_child(pr, "w:tcBorders"), "bottom", item, heavy)


def _text(cell: Any, item: dict[str, Any], colour: str, font: str) -> None:
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = font
            run.font.size = Pt(item["font_size_pt"])
            run.font.color.rgb = RGBColor.from_string(_hex(colour))
            fonts = run._element.get_or_add_rPr().get_or_add_rFonts()
            fonts.set(qn("w:eastAsia"), font)


def style_table(table: Any, item: dict[str, Any], font: str) -> None:
    tbl_pr = table._tbl.tblPr
    width = _child(tbl_pr, "w:tblW")
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), str(round(item["width_percent"] * 50)))
    look = item.get("borders", "horizontal")
    borders = _child(tbl_pr, "w:tblBorders")
    for edge in EDGES:
        on = (
            look == "grid"
            or (look == "horizontal" and edge in ("bottom", "insideH"))
            or (look == "minimal" and edge in ("top", "bottom"))
            or (look == "boxed" and edge in ("top", "left", "bottom", "right"))
        )
        _border(borders, edge, item if on else None, 1.6 if look == "minimal" else 1.0)
    margins = _child(tbl_pr, "w:tblCellMar")
    pad = str(round(item["cell_padding_em"] * item["font_size_pt"] * 20))
    for side in ("top", "bottom", "left", "right"):
        m = _child(margins, f"w:{side}")
        m.set(qn("w:w"), pad)
        m.set(qn("w:type"), "dxa")
    for r, row in enumerate(table.rows):
        header = r == 0
        for cell in row.cells:
            if header:
                _shade(cell, item["header_bg"])
                _text(cell, item, item["header_fg"], font)
                if look in ("minimal", "boxed", "header"):
                    _cell_bottom(cell, item, 1.2)
            else:
                fill = item["stripe_bg"] if r % 2 == 0 else item["body_bg"]
                _shade(cell, fill)
                _text(cell, item, item["body_fg"], font)


def apply(
    path: Path,
    ids_in_order: list[str],
    sides: dict[str, str],
    items: dict[str, dict[str, Any]],
    fonts: dict[str, list[str]],
) -> int:
    """Style the tables of the .docx at `path` that have individual settings; how many."""
    if not items:
        return 0
    doc = Document(str(path))
    by_number = {str(i): table_id for i, table_id in enumerate(ids_in_order, 1)}
    tables = {id(t._tbl): t for t in doc.tables}
    found: dict[str, Any] = {}
    pending: str | None = None  # a caption above, waiting for its table
    previous: Any = None  # the table just passed, for a caption below it
    for block in doc.element.body.iterchildren():
        if block.tag == qn("w:p"):
            text = "".join(t.text or "" for t in block.iter(qn("w:t")))
            m = CAPTION.match(text)
            table_id = by_number.get(m.group(1)) if m else None
            if table_id and sides.get(table_id) == "below" and previous is not None:
                found[table_id] = previous
                pending = None
            elif table_id:
                pending = table_id
            elif text.strip():
                pending = None
            if text.strip():
                previous = None
        elif block.tag == qn("w:tbl"):
            table = tables.get(id(block))
            if pending and table is not None:
                found[pending] = table
            pending, previous = None, table
    for table_id, table in found.items():
        if table_id in items:
            style_table(table, items[table_id], fonts[items[table_id]["font_family"]][0])
    styled = sum(1 for table_id in found if table_id in items)
    if styled:
        doc.save(str(path))
    return styled
