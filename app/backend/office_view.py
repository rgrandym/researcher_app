"""Office files as the app shows them: every workbook and deck a run produced, as HTML.

`outputs` lists the latest .xlsx/.pptx in each `deliverable/<slug>/` that has no
editable master yet, so a file a revision or an agent wrote is selectable even
when nothing declared it. `html` renders a native file, or an Office master
(rebuilt into its native file first), as a read-only page whose cells and slides
tell the app what was clicked, so a click can open the editor at that place.

Formula cells show the value Excel stored with the file. A workbook written by a
script stores none until Excel recalculates it; those cells are calculated here
(the `formulas` engine, after LET is expanded, which it does not know), and a
cell it cannot calculate shows its formula, marked, rather than a guessed number.
"""

from __future__ import annotations

import base64
import html as h
import logging
import os
import re
import tempfile
import threading
from datetime import date, datetime, time
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path
from typing import Any

from backend import revision, workspace

OFFICE = (".xlsx", ".pptx")
MAX_ROWS, MAX_COLS = 2000, 80
MAX_FORMULAS = 20_000  # beyond this, calculating would keep the page waiting too long
log = logging.getLogger(__name__)


def _versions(folder: Path) -> list[Path]:
    found = [v for v in folder.iterdir() if v.is_dir() and re.fullmatch(r"v\d+", v.name)]
    return sorted(found, key=lambda v: int(v.name[1:]))


def _native(version: Path) -> Path | None:
    """The deliverable file in a version folder: applied.json's, else the only Office file."""
    applied = revision._json(version / "applied.json", {})
    named = version / str(applied.get("file") or "")
    if applied.get("file") and named.is_file() and named.suffix.lower() in OFFICE:
        return named
    files = sorted(f for f in version.iterdir() if f.is_file() and f.suffix.lower() in OFFICE and not f.name.startswith(("~$", ".")))
    return files[0] if len(files) == 1 else None


def outputs(name: str) -> list[dict[str, Any]]:
    run = workspace.run_dir(name)
    root = run / "deliverable"
    if not root.is_dir():
        return []
    masters = {d.get("slug") for d in workspace.run_meta(run).get("documents", []) if (run / str(d.get("master"))).is_file()}
    out = []
    for folder in sorted(p for p in root.iterdir() if p.is_dir()):
        if folder.name in masters:
            continue  # the editable master stands for it
        for version in reversed(_versions(folder)):
            native = _native(version)
            if native:
                out.append({
                    "path": str(native.relative_to(run)), "slug": folder.name, "version": version.name,
                    "format": native.suffix[1:].lower(), "label": f"{native.name} ({version.name})",
                })
                break
    return out


def _is_master(run: Path, path: Path) -> bool:
    return path.parent == run and path.suffix == ".md" and (run / "subjects" / path.stem / "office-source.json").is_file()


def html(name: str, rel: str) -> str:
    run = workspace.run_dir(name).resolve()
    path = workspace.run_file(name, rel)
    if _is_master(run, path):
        office = revision.script("office_master")
        source, original, _ = office.baseline(run, path.stem)
        data = office.decode(path.read_text())
        office.validate(data, original)
        with tempfile.TemporaryDirectory(prefix="office-view-") as temp:
            target = Path(temp) / f"document.{data['format']}"
            with revision._office_lock:
                office.build_native(source, data, original, target)
            return _render(target, path.name)
    if path.suffix.lower() not in OFFICE:
        raise ValueError("Choose a .xlsx workbook or .pptx presentation.")
    return _render(path, path.name)


def _render(path: Path, title: str) -> str:
    body = _workbook(path) if path.suffix.lower() == ".xlsx" else _deck(path)
    return PAGE.format(title=h.escape(title), body=body)


# ---------------------------------------------------------------- workbooks

def _sections(fmt: str) -> list[str]:
    parts, cur, quoted = [], "", False
    for ch in fmt:
        if ch == '"':
            quoted = not quoted
        if ch == ";" and not quoted:
            parts.append(cur)
            cur = ""
        else:
            cur += ch
    return [*parts, cur]


def _number(value: float, fmt: str) -> str:
    if not fmt or fmt == "General":
        if isinstance(value, float) and value.is_integer() and abs(value) < 1e15:
            return str(int(value))
        return f"{value:.10g}" if isinstance(value, float) else str(value)
    sections = _sections(fmt)
    section, number = sections[0], value
    if value < 0 and len(sections) > 1:
        section, number = sections[1], -value
    elif value == 0 and len(sections) > 2:
        section = sections[2]
    section = re.sub(r"\[\$([^\]-]*)[^\]]*\]", lambda m: f'"{m[1]}"', section)
    section = re.sub(r"\[[^\]]*\]|_.|\*.|\\", "", section)
    pieces = re.findall(r'"[^"]*"|[#0?,.%]+|[^"#0?,.%]', section)
    code = next((p for p in pieces if re.search(r"[#0?]", p)), None)
    if code is None:
        return "".join(p.strip('"') for p in pieces) or _number(value, "General")
    exact = Decimal(str(number)) * (100 if "%" in section else 1)
    decimals = len(code.split(".")[1].rstrip(",%")) if "." in code else 0
    # Excel rounds the decimal value half away from zero, not the binary float.
    exact = exact.quantize(Decimal(1).scaleb(-decimals), rounding=ROUND_HALF_UP)
    text = f"{exact:,.{decimals}f}" if "," in code else f"{exact:.{decimals}f}"
    shown = []
    for p in pieces:
        if p == code:
            shown.append(text + ("%" if "%" in code else ""))
        elif not re.search(r"[#0?,.]", p) or p.startswith('"'):
            shown.append(p.strip('"'))
        elif "%" in p:
            shown.append("%")
    return "".join(shown)


def _when(value: date | time, fmt: str) -> str:
    f = fmt.lower()
    if isinstance(value, datetime) and value.time() == time(0) or isinstance(value, date) and not isinstance(value, datetime):
        if "mmm" in f:
            return value.strftime("%b %Y" if "yyyy" in f else "%b-%y" if "y" in f else "%d %b")
        return value.strftime("%d/%m/%Y")
    return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()


def _colour(c: Any) -> str | None:
    rgb = getattr(c, "rgb", None)
    if isinstance(rgb, str) and len(rgb) == 8 and rgb not in ("00000000",):
        return f"#{rgb[2:]}"
    return None


def _letters(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _without_let(formula: str) -> str:
    """=LET(a, 1, b, a+1, a*b) as =((1)*((1)+1)): each name replaced by its bracketed value.

    Whitespace tokens are dropped, since the engine rejects trailing spaces; that loses
    only Excel's space-as-intersection operator, which a formula here would then not use.
    """
    from openpyxl.formula.tokenizer import Token, Tokenizer  # type: ignore[import-untyped,import-not-found,unused-ignore]

    def text(tokens: list[Any], names: dict[str, str]) -> str:
        return "".join(f"({names[t.value.upper()]})" if t.type == Token.OPERAND and t.value.upper() in names else t.value
                       for t in tokens if t.type != Token.WSPACE)

    opens = ("LET(", "_XLFN.LET(")
    for _ in range(100):
        tokens = Tokenizer(formula).items
        start = next((i for i, t in reversed(list(enumerate(tokens))) if t.type == Token.FUNC and t.value.upper() in opens), None)
        if start is None:
            return "=" + text(tokens, {})
        args: list[list[Any]] = [[]]
        depth, end = 0, len(tokens)
        for i in range(start + 1, len(tokens)):
            t = tokens[i]
            if t.subtype == Token.OPEN and t.type in (Token.FUNC, Token.PAREN):
                depth += 1
            elif t.subtype == Token.CLOSE and t.type in (Token.FUNC, Token.PAREN):
                if depth == 0:
                    end = i
                    break
                depth -= 1
            elif t.type == Token.SEP and t.subtype == Token.ARG and depth == 0:
                args.append([])
                continue
            args[-1].append(t)
        names: dict[str, str] = {}
        for k in range(0, len(args) - 1, 2):
            names["".join(t.value for t in args[k] if t.type != Token.WSPACE).upper()] = text(args[k + 1], names)
        formula = "=" + text(tokens[:start], {}) + f"({text(args[-1], names)})" + text(tokens[end + 1:], {})
    raise ValueError("LET nested too deeply")


_calculated: dict[tuple[str, float, int], dict[tuple[str, str], Any]] = {}
_calc_lock = threading.Lock()


def _calculate(path: Path) -> dict[tuple[str, str], Any]:
    """Every formula cell's value by (SHEET, A1), or {} when the engine cannot run this workbook."""
    stat = path.stat()
    key = (str(path.resolve()), stat.st_mtime, stat.st_size)
    with _calc_lock:
        if key in _calculated:
            return _calculated[key]
        try:
            values = _engine(path)
        except Exception as exc:  # noqa: BLE001 - any engine failure leaves the formulas shown
            log.warning("could not calculate %s: %s", path.name, exc)
            values = {}
        _calculated.clear()  # one workbook at a time is what the page shows
        _calculated[key] = values
        return values


def _engine(path: Path) -> dict[tuple[str, str], Any]:
    os.environ.setdefault("TQDM_DISABLE", "1")  # the engine's progress bars would fill the log
    import formulas  # type: ignore[import-untyped,import-not-found,unused-ignore]
    from openpyxl import load_workbook  # type: ignore[import-untyped,import-not-found,unused-ignore]

    book = load_workbook(path)
    count = 0
    for sheet in book.worksheets:
        for cell in sheet._cells.values():
            if cell.data_type == "f" and isinstance(cell.value, str):
                cell.value = _without_let(cell.value)
                count += 1
    if not count or count > MAX_FORMULAS:
        return {}
    with tempfile.TemporaryDirectory(prefix="office-calc-") as temp:
        target = Path(temp) / "book.xlsx"
        book.save(target)
        solution = formulas.ExcelModel().loads(str(target)).finish().calculate()
    out: dict[tuple[str, str], Any] = {}
    for name, result in solution.items():
        ref = name.split("]", 1)[-1]
        if "!" not in ref or ":" in ref:
            continue
        sheet, cell = ref.rsplit("!", 1)
        value = getattr(result, "value", result)
        value = value[0, 0] if hasattr(value, "shape") else value
        if type(value).__name__ == "XlError":
            value = str(value)
        elif hasattr(value, "item"):
            value = value.item()
        out[(sheet.strip("'").upper(), cell.upper())] = value
    return out


def _workbook(path: Path) -> str:
    from openpyxl import load_workbook  # type: ignore[import-untyped,import-not-found,unused-ignore]

    book = load_workbook(path)
    values = load_workbook(path, data_only=True)
    nav, sheets, missing, worked = [], [], 0, 0
    calculated: dict[tuple[str, str], Any] | None = None
    for i, sheet in enumerate(book.worksheets):
        cached = values[sheet.title]
        cells = {(c.row, c.column): c for c in sheet._cells.values() if c.value is not None}
        spans: dict[tuple[int, int], tuple[int, int]] = {}
        covered: set[tuple[int, int]] = set()
        for m in sheet.merged_cells.ranges:
            spans[(m.min_row, m.min_col)] = (m.max_row - m.min_row + 1, m.max_col - m.min_col + 1)
            covered |= {(r, c) for r in range(m.min_row, m.max_row + 1) for c in range(m.min_col, m.max_col + 1)} - {(m.min_row, m.min_col)}
        keys = [*cells, *spans]
        rows = min(max((r for r, _ in keys), default=0), MAX_ROWS)
        cols = min(max((c for _, c in keys), default=0), MAX_COLS)
        hidden = " (hidden)" if sheet.sheet_state != "visible" else ""
        nav.append(f'<a href="#s{i}">{h.escape(sheet.title)}{hidden}</a>')
        # One column dimension can cover a range (C:Z); hidden rows and columns stay hidden, as in Excel.
        width_of: dict[int, float] = {}
        shut: set[int] = set()
        for dim in sheet.column_dimensions.values():
            for c in range(dim.min or 1, min(dim.max or dim.min or 1, cols) + 1):
                if dim.width:
                    width_of[c] = dim.width
                if dim.hidden:
                    shut.add(c)
        shown_cols = [c for c in range(1, cols + 1) if c not in shut]
        shown_rows = [r for r in range(1, rows + 1) if not (sheet.row_dimensions.get(r) and sheet.row_dimensions[r].hidden)]
        visible = set(shown_rows)
        pixels = [round(width_of.get(c, 8.43) * 7 + 5) for c in shown_cols]
        widths = [f'<col style="width:{px}px">' for px in pixels]
        out = [f'<section id="s{i}"><h2>{h.escape(sheet.title)}{hidden}</h2>']
        if not cells:
            out.append('<p class="note">This sheet is empty.</p></section>')
            sheets.append("".join(out))
            continue
        out.append(f'<div class="grid"><table data-sheet="{h.escape(sheet.title)}" style="width:{44 + sum(pixels)}px"><colgroup><col class="rh">{"".join(widths)}</colgroup>')
        out.append("<thead><tr><th></th>" + "".join(f"<th>{_letters(c)}</th>" for c in shown_cols) + "</tr></thead><tbody>")
        for r in shown_rows:
            dim = sheet.row_dimensions.get(r)
            style = f' style="height:{round(dim.height * 4 / 3)}px"' if dim and dim.height else ""
            tr = [f"<tr{style}><th>{r}</th>"]
            for c in shown_cols:
                if (r, c) in covered:
                    continue
                ref = f"{_letters(c)}{r}"
                span = spans.get((r, c))
                attrs = f' data-cell="{ref}"'
                if span:
                    down = sum(1 for k in range(r, r + span[0]) if k in visible)
                    across = sum(1 for k in range(c, c + span[1]) if k not in shut)
                    attrs += (f' rowspan="{down}"' if down > 1 else "") + (f' colspan="{across}"' if across > 1 else "")
                cell = cells.get((r, c))
                if not cell:
                    tr.append(f"<td{attrs}></td>")
                    continue
                css, cls, text = [], [], ""
                value = cell.value
                if cell.data_type == "f" or not isinstance(value, (str, int, float, bool, date, time)):
                    stored = cached[ref].value
                    formula = value if isinstance(value, str) else str(getattr(value, "text", value))
                    if stored is None:
                        if calculated is None:
                            calculated = _calculate(path)
                        stored = calculated.get((sheet.title.upper(), ref))
                        if stored is not None:
                            worked += 1
                            formula += "  (calculated by the app)"
                    attrs += f' title="{h.escape(formula)}"'
                    if stored is None:
                        missing += 1
                        cls.append("formula")
                        text = h.escape(formula)
                    else:
                        value = stored
                if not text:
                    if isinstance(value, bool):
                        text, align = ("TRUE" if value else "FALSE"), "center"
                    elif isinstance(value, (int, float)):
                        text, align = h.escape(_number(value, cell.number_format)), "right"
                    elif isinstance(value, (date, time)):
                        text, align = h.escape(_when(value, cell.number_format)), "right"
                    else:
                        text, align = h.escape(str(value)), "left"
                    if not cell.alignment.horizontal:
                        css.append(f"text-align:{align}")
                font = cell.font
                if font.b:
                    css.append("font-weight:600")
                if font.i:
                    css.append("font-style:italic")
                if font.u:
                    css.append("text-decoration:underline")
                if colour := _colour(font.color):
                    css.append(f"color:{colour}")
                if cell.fill.fill_type == "solid" and (colour := _colour(cell.fill.fgColor)):
                    css.append(f"background:{colour}")
                if cell.alignment.horizontal in ("left", "right", "center", "justify"):
                    css.append(f"text-align:{cell.alignment.horizontal}")
                if cell.alignment.wrap_text:
                    cls.append("wrap")
                style = f' style="{";".join(css)}"' if css else ""
                klass = f' class="{" ".join(cls)}"' if cls else ""
                tr.append(f"<td{attrs}{klass}{style}>{text}</td>")
            out.append("".join(tr) + "</tr>")
        out.append("</tbody></table></div>")
        bounds = [*keys]
        if max((r for r, _ in bounds), default=0) > MAX_ROWS or max((c for _, c in bounds), default=0) > MAX_COLS:
            out.append(f'<p class="note">Showing the first {MAX_ROWS:,} rows and {MAX_COLS} columns. The editor and the file hold the rest.</p>')
        out.append("</section>")
        sheets.append("".join(out))
    book.close()
    values.close()
    notes = []
    if worked:
        notes.append(f"{worked:,} formula value{'s were' if worked != 1 else ' was'} not stored in this file, so the app calculated "
                     f"{'them' if worked != 1 else 'it'}; hover a cell for its formula. Excel recalculates every value when the file is opened there.")
    if missing:
        notes.append(f"{missing:,} formula cell{'s' if missing != 1 else ''} could not be calculated here and "
                     f"{'show their formulas' if missing != 1 else 'shows its formula'}; "
                     "Excel calculates them when the file is opened there.")
    banner = "".join(f'<p class="banner">{n}</p>' for n in notes)
    return f'<nav>{"".join(nav)}</nav>{banner}{"".join(sheets)}'


# ---------------------------------------------------------------- decks

def _emu_box(shape: Any, scale: float, origin: tuple[float, float, float, float]) -> str:
    ox, oy, sx, sy = origin
    left = (ox + (shape.left or 0) * sx) * scale
    top = (oy + (shape.top or 0) * sy) * scale
    width = (shape.width or 0) * sx * scale
    height = (shape.height or 0) * sy * scale
    return f"left:{left:.1f}px;top:{top:.1f}px;width:{width:.1f}px;height:{height:.1f}px"


def _frame(frame: Any, px: float) -> str:
    out = []
    for p in frame.paragraphs:
        css = []
        align = {1: "left", 2: "center", 3: "right", 4: "justify"}.get(int(p.alignment)) if p.alignment is not None else None
        if align:
            css.append(f"text-align:{align}")
        if p.level:
            css.append(f"margin-left:{p.level * 1.2}em")
        runs = []
        for run in p.runs:
            rcss = []
            if run.font.size:
                rcss.append(f"font-size:{run.font.size.pt * px:.1f}px")
            if run.font.bold:
                rcss.append("font-weight:600")
            if run.font.italic:
                rcss.append("font-style:italic")
            try:
                rgb = run.font.color.rgb if run.font.color and run.font.color.type is not None else None
            except AttributeError:
                rgb = None
            if rgb is not None:
                rcss.append(f"color:#{rgb}")
            style = f' style="{";".join(rcss)}"' if rcss else ""
            runs.append(f"<span{style}>{h.escape(run.text)}</span>")
        style = f' style="{";".join(css)}"' if css else ""
        out.append(f"<p{style}>{''.join(runs) or '&nbsp;'}</p>")
    return "".join(out)


def _shapes(shapes: Any, scale: float, px: float, origin: tuple[float, float, float, float], prefix: str = "") -> list[str]:
    out = []
    for shape in shapes:
        key = f"{prefix}{shape.shape_id}"
        box = _emu_box(shape, scale, origin)
        if shape.shape_type == 6:  # group: children use the group's own coordinates
            xfrm = shape._element.grpSpPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}xfrm")
            ox, oy, sx, sy = origin
            try:
                off, ext = xfrm.find("{*}off"), xfrm.find("{*}ext")
                ch_off, ch_ext = xfrm.find("{*}chOff"), xfrm.find("{*}chExt")
                kx = int(ext.get("cx")) / max(int(ch_ext.get("cx")), 1)
                ky = int(ext.get("cy")) / max(int(ch_ext.get("cy")), 1)
                inner = (ox + (int(off.get("x")) - int(ch_off.get("x")) * kx) * sx, oy + (int(off.get("y")) - int(ch_off.get("y")) * ky) * sy, sx * kx, sy * ky)
            except (AttributeError, TypeError, ValueError):
                inner = origin
            out += _shapes(shape.shapes, scale, px, inner, key + "/")
            continue
        css = [box]
        try:
            if shape.fill.type == 1 and shape.fill.fore_color.type is not None:
                css.append(f"background:#{shape.fill.fore_color.rgb}")
        except (AttributeError, TypeError, ValueError, NotImplementedError):
            pass
        style = ";".join(css)
        if shape.shape_type == 13:  # picture
            image = shape.image
            data = base64.b64encode(image.blob).decode()
            out.append(f'<img class="shape" data-shape="{key}" style="{style}" src="data:{image.content_type};base64,{data}" alt="">')
        elif getattr(shape, "has_table", False) and shape.has_table:
            rows = "".join("<tr>" + "".join(f"<td>{_frame(c.text_frame, px)}</td>" for c in row.cells) + "</tr>" for row in shape.table.rows)
            out.append(f'<div class="shape edit" data-shape="{key}" style="{style}"><table>{rows}</table></div>')
        elif getattr(shape, "has_chart", False) and shape.has_chart:
            title = shape.chart.chart_title.text_frame.text if shape.chart.has_title else ""
            out.append(f'<div class="shape chart" data-shape="{key}" style="{style}">Chart{": " + h.escape(title) if title else ""} (shown in PowerPoint)</div>')
        elif shape.has_text_frame:
            out.append(f'<div class="shape text edit" data-shape="{key}" style="{style};font-size:{18 * px:.1f}px">{_frame(shape.text_frame, px)}</div>')
    return out


def _deck(path: Path) -> str:
    from pptx import Presentation  # type: ignore[import-untyped,import-not-found,unused-ignore]

    deck = Presentation(str(path))
    width, height = deck.slide_width or 12192000, deck.slide_height or 6858000
    scale = 960 / width
    px = 960 / (width / 12700)  # px per point at this width
    slides = []
    for i, slide in enumerate(deck.slides):
        shapes = _shapes(slide.shapes, scale, px, (0, 0, 1, 1))
        slides.append(f'<section class="slide-wrap"><h2>Slide {i + 1}</h2><div class="slide" data-slide="{i}" style="height:{height * scale:.0f}px">{"".join(shapes)}</div></section>')
    return "".join(slides) or '<p class="note">This presentation has no slides.</p>'


PAGE = """<!doctype html><html><head><meta charset="utf-8"><title>{title}</title><style>
:root {{ color-scheme: light; }}
body {{ margin: 0; padding: 16px 20px 40px; background: #f6f6f4; color: #1d1d1b; font: 13px/1.4 -apple-system, "Segoe UI", Arial, sans-serif; }}
nav {{ position: sticky; top: 0; z-index: 3; display: flex; flex-wrap: wrap; gap: 4px; padding: 6px 0 10px; background: #f6f6f4; }}
nav a {{ padding: 3px 9px; border: 1px solid #d6d6d0; border-radius: 6px; background: #fff; color: inherit; text-decoration: none; }}
nav a:hover {{ border-color: #2f6feb; }}
h2 {{ margin: 22px 0 8px; font-size: 14px; font-weight: 600; }}
.banner {{ margin: 4px 0 8px; padding: 8px 10px; border: 1px solid #e5c56b; border-radius: 6px; background: #fff8e1; }}
.note {{ color: #6b6b66; }}
.grid {{ overflow: auto; max-height: 80vh; border: 1px solid #d6d6d0; background: #fff; }}
table {{ border-collapse: collapse; table-layout: fixed; }}
.grid th {{ position: sticky; top: 0; z-index: 1; background: #ecece8; color: #6b6b66; font-weight: 500; border: 1px solid #d6d6d0; padding: 1px 4px; }}
.grid tbody th {{ position: sticky; left: 0; top: auto; text-align: right; }}
col.rh {{ width: 44px; }}
.grid td {{ border: 1px solid #e6e6e2; padding: 1px 4px; white-space: nowrap; overflow: hidden; text-overflow: ellipsis; cursor: cell; }}
.grid td.wrap {{ white-space: normal; }}
.grid td.formula {{ color: #8a6d00; font-family: ui-monospace, Menlo, monospace; font-size: 11px; background: #fffbea; }}
.grid td:hover {{ box-shadow: inset 0 0 0 2px rgba(47,111,235,.45); }}
.slide {{ position: relative; width: 960px; overflow: hidden; background: #fff; box-shadow: 0 1px 4px rgba(0,0,0,.18); }}
.shape {{ position: absolute; box-sizing: border-box; overflow: hidden; }}
.shape p {{ margin: 0 0 .25em; }}
.shape table {{ width: 100%; height: 100%; }}
.shape td {{ border: 1px solid #c9c9c4; padding: 2px 5px; vertical-align: top; }}
.chart {{ display: flex; align-items: center; justify-content: center; border: 1px dashed #b9b9b3; color: #6b6b66; }}
.edit {{ cursor: pointer; }}
.edit:hover {{ outline: 2px dashed #2f6feb; outline-offset: 2px; }}
</style></head><body>{body}
<script>
document.addEventListener("click", (e) => {{
  // In a srcdoc frame a #fragment link would load a blank page, so the sheet links scroll instead.
  const link = e.target.closest("a[href^='#']");
  if (link) {{ e.preventDefault(); document.getElementById(link.getAttribute("href").slice(1))?.scrollIntoView(); return; }}
  const cell = e.target.closest("td[data-cell]");
  if (cell) {{ parent.postMessage({{ type: "office-pick", sheet: cell.closest("table").dataset.sheet, cell: cell.dataset.cell }}, "*"); return; }}
  const slide = e.target.closest("[data-slide]");
  const shape = e.target.closest("[data-shape]");
  if (slide) parent.postMessage({{ type: "office-pick", slide: Number(slide.dataset.slide), shape: shape ? shape.dataset.shape : null }}, "*");
}});
</script></body></html>"""
